# -*- coding: utf-8 -*-
"""Отобранные сообщения — в Excel.

Аналитик отбирает в «Сообщениях» ленту (тег, тип, поиск, инфоповод, срез) и
отправляет её заказчику таблицей. Выгружаются все отобранные сообщения, а не
только страница на экране. Второй лист — какой отбор выгружен: без него через
неделю не восстановить, откуда взялась таблица.

Метрики — числами, чтобы в Excel их можно было сложить и отсортировать; если
метрики не было в выгрузке периода — прочерк, как везде в продукте.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
import pandas as pd

from .formatting import ascii_filename
from .message_compute import message_link_column, message_text_column
from .message_kinds import NO_MESSAGE_TYPE_LABEL, message_type_labels
from .metrics_compute import NO_METRIC_VALUE, PERIOD_METRIC_COLUMNS, numeric_series
from .source_stats import NO_SOURCE_LABEL, platform_labels
from .story_recovery import is_residual_title

SHEET_MESSAGES = "Сообщения"
SHEET_PARAMS = "Параметры"
# Предел ячейки Excel — 32 767 символов; длиннее текст файл не откроет.
_CELL_LIMIT = 32000
_EMPTY = {"", "nan", "none", "null", "nat"}
_METRICS = (
    ("Аудитория", "audience", ["audience", "Аудитория"]),
    ("Охват", "reach", ["views", "Просмотры", "Просмотров", "reach", "Охват"]),
    ("Вовлеченность", "engagement", ["engagement", "Вовлечённость", "Вовлеченность", "engagement_count"]),
)


def _text(messages: pd.DataFrame, *columns: str) -> pd.Series:
    for column in columns:
        if column in messages.columns:
            values = messages[column].fillna("").astype(str).str.strip()
            return values.where(~values.str.lower().isin(_EMPTY), "")
    return pd.Series([""] * len(messages), index=messages.index, dtype="object")


def _metric(messages: pd.DataFrame, key: str, raw: list[str]) -> pd.Series:
    """Число или прочерк: метрики не было в выгрузке периода этого сообщения."""
    prepared = f"_{key}"
    values = (
        pd.to_numeric(messages[prepared], errors="coerce").fillna(0)
        if prepared in messages.columns
        else numeric_series(messages, raw)
    ).astype(int).astype(object)
    flag = PERIOD_METRIC_COLUMNS[key]
    if flag in messages.columns:
        known = messages[flag].map(lambda v: True if v is None or pd.isna(v) else bool(v))
        values = values.where(known.astype(bool), NO_METRIC_VALUE)
    return values


def messages_export_frame(messages: pd.DataFrame) -> pd.DataFrame:
    """Таблица для Excel: колонки по-русски, свежие сообщения сверху."""
    if not isinstance(messages, pd.DataFrame) or messages.empty:
        return pd.DataFrame()
    work = messages
    if "datetime" in work.columns:
        order = pd.to_datetime(work["datetime"], errors="coerce")
        work = work.assign(_order=order).sort_values("_order", ascending=False, kind="stable")
    text_col = message_text_column(work)
    link_col = message_link_column(work)
    platforms = platform_labels(work)
    types = message_type_labels(work)
    events = _text(work, "event_title", "source_main_topic")
    columns: dict[str, pd.Series] = {}
    if "datetime" in work.columns:
        dates = pd.to_datetime(work["datetime"], errors="coerce")
        if dates.notna().any():
            # Excel не хранит часовой пояс: время остаётся как в выгрузке.
            columns["Дата"] = dates.dt.tz_localize(None) if dates.dt.tz is not None else dates
    frame = pd.DataFrame(
        {
            **columns,
            "Площадка": platforms.where(platforms != NO_SOURCE_LABEL, ""),
            "Сообщество / канал": _text(work, "chat_title"),
            "Автор": _text(work, "author", "Автор"),
            "Тип сообщения": types.where(types != NO_MESSAGE_TYPE_LABEL, ""),
            "Тональность": _text(work, "sentiment", "Тональность"),
            "Теги": _text(work, "tags", "Теги").str.replace("|", ", ", regex=False),
            # «Без сюжета» — служебная корзина платформы, а не инфоповод.
            "Инфоповод": events.map(lambda title: "" if is_residual_title(title) else title),
            "Текст": (_text(work, text_col) if text_col else _text(work)).str.slice(0, _CELL_LIMIT),
            "Ссылка": _text(work, link_col) if link_col else _text(work),
        },
        index=work.index,
    )
    for title, key, raw in _METRICS:
        frame[title] = _metric(work, key, raw)
    return frame.reset_index(drop=True)


def safe_messages_filename(project_name: str, period_label: str) -> str:
    """Имя файла латиницей — кириллическое браузер сохранял как «download»."""
    return ascii_filename("messages", project_name, period_label, ext="xlsx", fallback="messages")


def messages_to_xlsx(
    messages: pd.DataFrame,
    *,
    project_name: str = "",
    period_label: str = "",
    filters: list[str] | None = None,
) -> bytes:
    """Excel: лист «Сообщения» и лист «Параметры» с описанием отбора."""
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    frame = messages_export_frame(messages)
    params = pd.DataFrame(
        [
            ("Проект", project_name or "—"),
            ("Период", period_label or "—"),
            ("Отбор", "; ".join(filters or []) or "все сообщения"),
            ("Сообщений", len(frame)),
            ("Выгружено", datetime.now().strftime("%d.%m.%Y %H:%M")),
        ],
        columns=["Параметр", "Значение"],
    )
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name=SHEET_MESSAGES, index=False)
        params.to_excel(writer, sheet_name=SHEET_PARAMS, index=False)
        sheet = writer.sheets[SHEET_MESSAGES]
        widths = {
            "Дата": 17, "Площадка": 16, "Сообщество / канал": 24, "Автор": 20,
            "Тип сообщения": 14, "Тональность": 13, "Теги": 24, "Инфоповод": 28,
            "Текст": 70, "Ссылка": 30, "Аудитория": 12, "Охват": 12, "Вовлеченность": 14,
        }
        for index, column in enumerate(frame.columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = widths.get(column, 14)
            sheet.cell(row=1, column=index).font = Font(bold=True)
        sheet.freeze_panes = "A2"
        columns = list(frame.columns)
        if "Дата" in columns:
            date_col = columns.index("Дата") + 1
            for row in range(2, len(frame) + 2):
                sheet.cell(row=row, column=date_col).number_format = "DD.MM.YYYY HH:MM"
        text_col = columns.index("Текст") + 1
        link_col = columns.index("Ссылка") + 1
        for row in range(2, len(frame) + 2):
            sheet.cell(row=row, column=text_col).alignment = Alignment(wrap_text=True, vertical="top")
            link = sheet.cell(row=row, column=link_col)
            if str(link.value or "").startswith("http"):
                link.hyperlink = str(link.value)
                link.font = Font(color="0563C1", underline="single")
        params_sheet = writer.sheets[SHEET_PARAMS]
        params_sheet.column_dimensions["A"].width = 14
        params_sheet.column_dimensions["B"].width = 80
        params_sheet.cell(row=1, column=1).font = Font(bold=True)
        params_sheet.cell(row=1, column=2).font = Font(bold=True)
    return buffer.getvalue()


def describe_filters(
    *,
    event_title: str = "",
    tags: list[str] | None = None,
    match_all: bool = False,
    platforms: list[str] | None = None,
    tones: list[str] | None = None,
    types: list[str] | None = None,
    search: str = "",
    slice_tags: list[str] | None = None,
) -> list[str]:
    """Отбор словами — для листа «Параметры» и подписи у кнопки."""
    parts: list[str] = []
    if slice_tags:
        parts.append("срез по тегам: " + ", ".join(slice_tags))
    if event_title:
        parts.append(f"инфоповод: {event_title}")
    if tags:
        joiner = " и " if match_all and len(tags) > 1 else ", "
        parts.append(("теги: " if len(tags) > 1 else "тег: ") + joiner.join(tags))
    if platforms:
        parts.append(("площадки: " if len(platforms) > 1 else "площадка: ") + ", ".join(platforms))
    if tones:
        parts.append("тональность: " + ", ".join(tone.lower() for tone in tones))
    if types:
        parts.append(("типы: " if len(types) > 1 else "тип: ") + ", ".join(types))
    if search.strip():
        parts.append(f"поиск: «{search.strip()}»")
    return parts
