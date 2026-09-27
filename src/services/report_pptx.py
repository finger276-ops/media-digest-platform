# -*- coding: utf-8 -*-
"""Отчёт в PowerPoint — тот же набор данных, что у Word, PDF и PNG.

Итог работы аналитика для заказчика — почти всегда презентация, и раньше её
собирали руками из PNG и скриншотов. Здесь она строится из того же payload
(report_export.summary_export_payload): те же разделы, выбранные в «Разделах
отчёта», то же брендирование, те же прочерки вместо ложных нулей.

Диаграммы — родные диаграммы PowerPoint, а не картинки: аналитик может
поменять цвета, подписи и порядок, не пересобирая отчёт.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .metrics_compute import (
    METRIC_TITLES,
    VOLUME_METRICS,
    format_int,
    metric_missing,
    metric_text,
    no_sentiment_line,
    percent_text,
)
from .report_export import _classify_summary_line, resolve_report_sections

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.6)
INK = RGBColor.from_string("111827")
MUTED = RGBColor.from_string("6B7280")
CARD_FILL = RGBColor.from_string("F3F4F6")
TONE_COLORS = ("16A34A", "9CA3AF", "DC2626")
# Строк текста саммари на слайд: больше — мелкий шрифт, который с экрана не
# читается; длинные абзацы считаются за несколько строк.
SUMMARY_LINES_PER_SLIDE = 14
SUMMARY_CHARS_PER_LINE = 110


def _rgb(value: Any, fallback: str) -> RGBColor:
    text = str(value or "").strip().lstrip("#")
    try:
        return RGBColor.from_string(text.upper() if len(text) == 6 else fallback)
    except ValueError:
        return RGBColor.from_string(fallback)


class _Deck:
    def __init__(self, payload: dict[str, Any]):
        self.payload = payload
        self.accent = _rgb(payload.get("accent_color"), "2563EB")
        self.background = _rgb(payload.get("background_color"), "FFFFFF")
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_W
        self.prs.slide_height = SLIDE_H
        self.blank = self.prs.slide_layouts[6]
        self.footer = str(
            payload.get("footer_text") or payload.get("client_name") or payload.get("project_name") or ""
        ).strip()

    def text(self, slide, left, top, width, height, value, *, size=14, bold=False,
             color=None, align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(left, top, width, height)
        frame = box.text_frame
        frame.word_wrap = True
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = str(value)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color or INK
        return box

    def slide(self, title: str | None = None):
        slide = self.prs.slides.add_slide(self.blank)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = self.background
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = self.accent
        bar.line.fill.background()
        if title:
            self.text(slide, MARGIN, Inches(0.35), SLIDE_W - 2 * MARGIN, Inches(0.8), title,
                      size=28, bold=True, color=self.accent)
        number = len(self.prs.slides)
        footer = f"{self.footer} · {number}" if self.footer else str(number)
        self.text(slide, MARGIN, SLIDE_H - Inches(0.5), SLIDE_W - 2 * MARGIN, Inches(0.35),
                  footer, size=10, color=MUTED, align=PP_ALIGN.RIGHT)
        return slide

    def save(self) -> bytes:
        out = BytesIO()
        self.prs.save(out)
        return out.getvalue()


def _title_slide(deck: _Deck) -> None:
    payload = deck.payload
    slide = deck.slide()
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.35), SLIDE_H)
    band.fill.solid()
    band.fill.fore_color.rgb = deck.accent
    band.line.fill.background()
    logo = payload.get("logo_bytes") or b""
    if logo:
        try:
            slide.shapes.add_picture(BytesIO(logo), SLIDE_W - MARGIN - Inches(2.4), Inches(0.5),
                                     height=Inches(1.0))
        except Exception:  # noqa: BLE001 — неподдерживаемый формат логотипа: слайд без него
            pass
    deck.text(slide, Inches(1.0), Inches(2.2), Inches(11), Inches(1.2),
              payload.get("report_title") or "Дайджест упоминаний", size=40, bold=True, color=deck.accent)
    deck.text(slide, Inches(1.0), Inches(3.4), Inches(11), Inches(0.8),
              payload.get("client_name") or payload.get("project_name") or "Проект", size=24, bold=True)
    deck.text(slide, Inches(1.0), Inches(4.3), Inches(11), Inches(0.6),
              f"Период: {payload.get('period_label') or 'выбранный период'}", size=16, color=MUTED)
    if payload.get("created_at"):
        deck.text(slide, Inches(1.0), Inches(4.8), Inches(11), Inches(0.5),
                  f"Дата выгрузки: {payload.get('created_at')}", size=12, color=MUTED)


def _card(deck: _Deck, slide, left, top, width, height, label: str, value: str) -> None:
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_FILL
    card.line.fill.background()
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.08), height)
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = deck.accent
    stripe.line.fill.background()
    deck.text(slide, left + Inches(0.3), top + Inches(0.2), width - Inches(0.4), Inches(0.4),
              label, size=14, color=MUTED)
    deck.text(slide, left + Inches(0.3), top + Inches(0.6), width - Inches(0.4), Inches(0.8),
              value, size=32, bold=True)


def _metrics_slide(deck: _Deck) -> None:
    payload = deck.payload
    slide = deck.slide("Основные метрики")
    cards = [("Сообщения", format_int(payload.get("messages", 0)))] + [
        (METRIC_TITLES[key].capitalize(), metric_text(payload, key)) for key in VOLUME_METRICS
    ]
    width = (SLIDE_W - 2 * MARGIN - Inches(0.3) * 3) / 4
    for index, (label, value) in enumerate(cards):
        _card(deck, slide, MARGIN + index * (width + Inches(0.3)), Inches(1.5), Emu(int(width)),
              Inches(1.5), label, value)
    missing = [METRIC_TITLES[key] for key in VOLUME_METRICS if metric_missing(payload, key)]
    if missing:
        deck.text(slide, MARGIN, Inches(3.15), SLIDE_W - 2 * MARGIN, Inches(0.4),
                  f"«—» — нет в выгрузке: {', '.join(missing)}.", size=12, color=MUTED)
    sequence = [item for item in payload.get("comparison_sequence") or [] if isinstance(item, dict)]
    if len(sequence) >= 2:
        data = CategoryChartData()
        data.categories = [str(item.get("label") or item.get("period_id") or "") for item in sequence]
        data.add_series("Сообщений", [int(item.get("messages", 0) or 0) for item in sequence])
        frame = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, MARGIN, Inches(3.6),
                                       SLIDE_W - 2 * MARGIN, Inches(3.2), data)
        chart = frame.chart
        chart.has_legend = False
        chart.has_title = True
        chart.chart_title.text_frame.text = "Динамика сообщений"
        series = chart.plots[0].series[0]
        series.format.line.color.rgb = deck.accent
        series.format.line.width = Pt(2.5)


def _sentiment_slide(deck: _Deck) -> None:
    payload = deck.payload
    slide = deck.slide("Тональность")
    if not payload.get("sentiment_markup", True):
        deck.text(slide, MARGIN, Inches(2.0), SLIDE_W - 2 * MARGIN, Inches(1.5),
                  no_sentiment_line("Тональность"), size=18, color=MUTED)
        return
    values = [int(payload.get(key, 0) or 0) for key in ("positive", "neutral", "negative")]
    total = max(1, int(payload.get("total", 0) or 0))
    data = CategoryChartData()
    data.categories = ["Позитив", "Нейтрал", "Негатив"]
    data.add_series("Сообщений", values)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, MARGIN, Inches(1.3), Inches(6.5),
                                   Inches(5.4), data)
    chart = frame.chart
    chart.has_legend = True
    plot = chart.plots[0]
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.show_percentage = True
    labels.show_value = False
    labels.number_format = "0%"
    labels.number_format_is_linked = False
    for point, color in zip(plot.series[0].points, TONE_COLORS):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RGBColor.from_string(color)
    lines = [
        f"{label}: {percent_text(value, total)} · {format_int(value)} сообщ."
        for label, value in zip(("Позитив", "Нейтрал", "Негатив"), values)
    ]
    deck.text(slide, Inches(7.6), Inches(2.6), Inches(5), Inches(2.5), "\n".join(lines), size=20)


def _bar_slide(deck: _Deck, title: str, items: list[dict[str, Any]]) -> None:
    slide = deck.slide(title)
    shown = [item for item in items[:5] if item.get("name")]
    data = CategoryChartData()
    # Столбцы горизонтальной диаграммы идут снизу вверх — переворачиваем,
    # чтобы первым сверху стоял самый крупный.
    data.categories = [str(item["name"])[:60] for item in reversed(shown)]
    data.add_series("Сообщений", [int(item.get("messages", 0) or 0) for item in reversed(shown)])
    frame = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, MARGIN, Inches(1.3),
                                   SLIDE_W - 2 * MARGIN, Inches(5.4), data)
    chart = frame.chart
    chart.has_legend = False
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.number_format = "# ##0"
    plot.data_labels.number_format_is_linked = False
    series = plot.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = deck.accent


def _events_slide(deck: _Deck, items: list[dict[str, Any]]) -> None:
    slide = deck.slide("Топ инфоповодов")
    shown = [item for item in items[:5] if item.get("name")]
    rows = len(shown) + 1
    table = slide.shapes.add_table(rows, 2, MARGIN, Inches(1.4), SLIDE_W - 2 * MARGIN,
                                   Inches(0.6) * rows).table
    table.columns[0].width = Emu(int(SLIDE_W - 2 * MARGIN - Inches(2.2)))
    table.columns[1].width = Inches(2.2)
    for column, header in enumerate(("Инфоповод", "Сообщений")):
        cell = table.cell(0, column)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = deck.accent
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor.from_string("FFFFFF")
    for index, item in enumerate(shown, start=1):
        for column, value in enumerate((str(item["name"]), format_int(item.get("messages", 0)))):
            cell = table.cell(index, column)
            cell.text = value
            cell.text_frame.paragraphs[0].runs[0].font.size = Pt(16)


def _bullets_slide(deck: _Deck, title: str, lines: list[tuple[str, str]]) -> None:
    slide = deck.slide(title)
    box = slide.shapes.add_textbox(MARGIN, Inches(1.3), SLIDE_W - 2 * MARGIN, SLIDE_H - Inches(2.0))
    frame = box.text_frame
    frame.word_wrap = True
    first = True
    for kind, text in lines:
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        run = paragraph.add_run()
        run.text = f"• {text}" if kind == "bullet" else text
        run.font.size = Pt(18 if kind == "heading" else 16)
        run.font.bold = kind == "heading"
        run.font.color.rgb = deck.accent if kind == "heading" else INK
        paragraph.space_after = Pt(6)


def summary_pages(summary_text: str) -> list[list[tuple[str, str]]]:
    """Разложить текст саммари по слайдам, не теряя ни строки."""
    pages: list[list[tuple[str, str]]] = []
    page: list[tuple[str, str]] = []
    used = 0
    for raw in str(summary_text or "").split("\n"):
        kind, text = _classify_summary_line(raw)
        if not text:
            continue
        cost = max(1, -(-len(text) // SUMMARY_CHARS_PER_LINE))
        # Подзаголовок не оставляем последней строкой слайда.
        if page and (used + cost > SUMMARY_LINES_PER_SLIDE or (kind == "heading" and used + cost + 1 > SUMMARY_LINES_PER_SLIDE)):
            pages.append(page)
            page, used = [], 0
        page.append((kind, text))
        used += cost
    if page:
        pages.append(page)
    return pages


def generate_summary_pptx(payload: dict[str, Any]) -> bytes:
    """Презентация по выбранным разделам отчёта."""
    deck = _Deck(payload)
    sections = set(resolve_report_sections(payload.get("sections")))
    _title_slide(deck)
    if "metrics" in sections:
        _metrics_slide(deck)
    if "sentiment" in sections:
        _sentiment_slide(deck)
    if "top_tags" in sections and payload.get("top_tags"):
        _bar_slide(deck, "Топ тегов", payload["top_tags"])
    if "top_events" in sections and payload.get("top_events"):
        _events_slide(deck, payload["top_events"])
    highlights = [str(x).strip() for x in payload.get("summary_highlights") or [] if str(x).strip()]
    if "highlights" in sections and highlights:
        _bullets_slide(deck, "Главное", [("bullet", text) for text in highlights])
    if "summary_text" in sections:
        pages = summary_pages(payload.get("summary_text") or "")
        for index, page in enumerate(pages):
            title = "Саммари периода" if index == 0 else "Саммари периода (продолжение)"
            _bullets_slide(deck, title, page)
    return deck.save()
