# -*- coding: utf-8 -*-
"""Отчёт в Word: инфографика первой страницей, затем текстовые блоки.
"""

from __future__ import annotations

import logging
from io import BytesIO
from typing import Any


from .metrics_compute import (
    format_int,
    no_sentiment_line,
    percent_text,
)
from .observability import report_failure
from .project_settings import valid_hex_color
from .report_common import (
    _classify_summary_line,
    _docx_metric,
    message_types_line,
    resolve_report_sections,
    audience_lines,
    top_sources_lines,
    _VISUAL_SECTIONS,
)
from .report_png import generate_summary_infographic_png

LOGGER = logging.getLogger("platform.report_export")


def _docx_bottom_border(paragraph, color_hex: str, size: int = 16) -> None:
    """Цветная линия под абзацем - тот же приём, что акцентная полоса в
    шапке PNG-инфографики, только средствами Word (нет прямого API, поэтому
    через oxml - стандартный, задокументированный обходной путь)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color_hex.lstrip("#"))
    borders.append(bottom)
    p_pr.append(borders)


def generate_summary_docx(payload: dict[str, Any]) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt, Inches, Cm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except Exception as exc:
        raise RuntimeError(
            "Для выгрузки Word добавьте python-docx в requirements.txt."
        ) from exc

    accent = valid_hex_color(payload.get("accent_color"), "#2563eb")
    accent_rgb = RGBColor.from_string(accent.lstrip("#"))
    muted_rgb = RGBColor.from_string("6b7280")
    ink_rgb = RGBColor.from_string("111827")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    # Тот же шрифт, что и в PNG-инфографике ("font.family": "DejaVu Sans" в
    # generate_summary_infographic_png) - раньше здесь стоял Arial, и страницы
    # саммари визуально не совпадали с картинкой на первой странице того же
    # документа. Шрифт не встраивается в .docx (в отличие от PDF, где он
    # зашит через TTFont) - если у читателя его нет, Word подставит похожий
    # рубленый шрифт, это мягкая деградация, не поломка.
    docx_font = "DejaVu Sans"
    styles = doc.styles
    styles["Normal"].font.name = docx_font
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].font.color.rgb = ink_rgb
    styles["Heading 1"].font.name = docx_font
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.color.rgb = accent_rgb
    # Тот же акцентный цвет, что заголовки блоков и полоса на карточках
    # метрик в PNG-инфографике - страницы саммари не должны выглядеть
    # отдельным, неоформленным документом рядом с картинкой на первой странице.
    styles["Heading 2"].font.name = docx_font
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 2"].font.color.rgb = accent_rgb
    # Подзаголовки ВНУТРИ текста саммари (build_auto_summary размечает их
    # "## ...") - на ступень мельче "Heading 2", чтобы не спорить визуально
    # с заголовками разделов отчёта.
    styles["Heading 3"].font.name = docx_font
    styles["Heading 3"].font.size = Pt(11.5)
    styles["Heading 3"].font.color.rgb = accent_rgb

    p = doc.add_paragraph()
    r = p.add_run(str(payload.get("report_title") or "Дайджест упоминаний"))
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = accent_rgb
    p.paragraph_format.space_after = Pt(4)

    p2 = doc.add_paragraph()
    r2 = p2.add_run(
        str(payload.get("client_name") or payload.get("project_name") or "Проект")
    )
    r2.bold = True
    r2.font.size = Pt(13)
    r2.font.color.rgb = ink_rgb
    p2.paragraph_format.space_after = Pt(6)

    meta = doc.add_paragraph()
    meta_lines = [
        f"Период: {payload.get('period_label') or 'выбранный период'}",
        f"Дата выгрузки: {payload.get('created_at') or ''}",
    ]
    for i, line in enumerate(meta_lines):
        text = line if i == len(meta_lines) - 1 else line + "\n"
        run = meta.add_run(text)
        run.font.color.rgb = muted_rgb
        run.font.size = Pt(9.5)
    meta.paragraph_format.space_after = Pt(10)
    _docx_bottom_border(meta, accent)

    sections = set(resolve_report_sections(payload.get("sections")))
    has_visual_sections = bool(_VISUAL_SECTIONS & sections)

    if has_visual_sections:
        try:
            infographic_png = generate_summary_infographic_png(payload)
            pic_p = doc.add_paragraph()
            pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = pic_p.add_run()
            run.add_picture(BytesIO(infographic_png), width=Inches(6.4))
            doc.add_page_break()
        except Exception as exc:  # noqa: BLE001 — Word без инфографики лучше, чем без Word
            # Клиент получит документ и не узнает, что страницы не хватает, —
            # поэтому владелец должен узнать вместо него.
            LOGGER.warning("Инфографика для Word не собралась", exc_info=True)
            report_failure("выгрузка Word: инфографика не собралась", exc)

    total = max(1, int(payload.get("total", 0) or 0))
    if "metrics" in sections or "sentiment" in sections:
        doc.add_heading("Основные метрики", level=2)
        if "metrics" in sections:
            doc.add_paragraph(
                f"Сообщений — {format_int(payload.get('messages', 0))}; "
                f"аудитория — {_docx_metric(payload, 'audience')}; "
                f"охват — {_docx_metric(payload, 'reach')}; "
                f"вовлеченность — {_docx_metric(payload, 'engagement')}."
            )
            types_line = message_types_line(payload)
            if types_line:
                doc.add_paragraph(types_line)
        if "sentiment" in sections and not payload.get("sentiment_markup", True):
            doc.add_paragraph(no_sentiment_line("Тональность"))
        elif "sentiment" in sections:
            doc.add_paragraph(
                f"Тональность: позитив — {percent_text(int(payload.get('positive', 0) or 0), total)}; "
                f"нейтрал — {percent_text(int(payload.get('neutral', 0) or 0), total)}; "
                f"негатив — {percent_text(int(payload.get('negative', 0) or 0), total)}."
            )

    if "top_tags" in sections or "top_events" in sections:
        doc.add_heading("Что включить в отчет", level=2)
        if "top_tags" in sections:
            top_tags = payload.get("top_tags") or []
            if top_tags:
                doc.add_paragraph(
                    "Топ тегов: "
                    + "; ".join(
                        str(x.get("name") or "") for x in top_tags[:5] if x.get("name")
                    )
                )
        if "top_events" in sections:
            top_events = payload.get("top_events") or []
            if top_events:
                doc.add_paragraph(
                    "Топ инфоповодов: "
                    + "; ".join(
                        str(x.get("name") or "")
                        for x in top_events[:5]
                        if x.get("name")
                    )
                )

    source_lines = top_sources_lines(payload)
    if "top_sources" in sections and source_lines:
        doc.add_heading("Площадки", level=2)
        doc.add_paragraph("Топ площадок: " + "; ".join(source_lines) + ".")

    audience = audience_lines(payload)
    if "audience" in sections and audience:
        doc.add_heading("Аудитория", level=2)
        for line in audience:
            doc.add_paragraph(line)

    if "summary_text" in sections:
        doc.add_heading("Саммари периода", level=2)
        for raw in str(payload.get("summary_text") or "").split("\n"):
            kind, text = _classify_summary_line(raw)
            if not text:
                continue
            if kind == "heading":
                doc.add_heading(text, level=3)
            elif kind == "bullet":
                para = doc.add_paragraph(text, style="List Bullet")
                para.paragraph_format.space_after = Pt(2)
            else:
                para = doc.add_paragraph(text)
                para.paragraph_format.space_after = Pt(4)

    out = BytesIO()
    doc.save(out)
    return out.getvalue()
