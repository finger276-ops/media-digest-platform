# -*- coding: utf-8 -*-
"""Генерация выгрузок саммари: инфографика PNG, Word, PDF.

Framework-independent (без Streamlit), как services/brand_metrics.py — платформа
только собирает данные в payload, а сама отрисовка/сборка документа не знает
про Streamlit. UI-обёртка (кнопки скачивания) — в report_export_ui.py.
"""

from __future__ import annotations

import logging
import os
import re
import textwrap
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd

from .chart_style import SENTIMENT_COLOR_RANGE
from .dashboard_config import (
    DEFAULT_REPORT_SECTIONS,
    REPORT_SECTION_OPTIONS,
)
from .cached_store import download_storage_file
from .metrics_compute import format_int, percent_text
from .observability import report_failure
from .project_settings import report_branding_from_project_settings, valid_hex_color
from .report_highlights import event_title_column, top_report_events, top_report_tags

LOGGER = logging.getLogger("platform.report_export")

# Разделы, у которых вообще есть что нарисовать на PNG-инфографике. Если
# аналитик оставил только "Полный текст саммари", инфографика с одним
# заголовком и пустым телом — лишняя страница, а не полезный блок.
_VISUAL_SECTIONS = {"metrics", "sentiment", "top_tags", "top_events", "highlights"}


def first_existing_col(df: pd.DataFrame, columns: list[str | None]) -> str | None:
    if df is None or not isinstance(df, pd.DataFrame):
        return None
    for col in columns or []:
        if col and col in df.columns:
            return str(col)
    return None


def clean_summary_for_export(text: str) -> str:
    value = str(text or "").strip()
    value = value.replace("**", "")
    return value


def _classify_summary_line(raw: str) -> tuple[str, str]:
    """Разобрать одну строку summary_text на (тип, видимый текст).

    Строки текста саммари (автотекст из summary_ui.build_auto_summary или
    вручную отредактированный) размечены минимально: "## " - подзаголовок
    раздела (тот же маркер, что Streamlit понимает "из коробки" в live-
    превью через st.markdown), "• "/"- " - пункт списка. И PDF, и DOCX
    рисуют эти три вида по-разному вместо одного и того же стиля абзаца на
    каждую строку."""
    line = raw.strip()
    if line.startswith("## "):
        return "heading", line[3:].strip()
    if line.startswith("• ") or line.startswith("- "):
        return "bullet", line[2:].strip()
    return "text", line


def safe_export_filename(project_name: str, period_label: str, ext: str) -> str:
    raw = f"summary_{project_name}_{period_label}"
    safe = re.sub(r"[^0-9A-Za-zА-Яа-я_.-]+", "_", raw, flags=re.UNICODE).strip("_")
    return f"{safe[:140] or 'summary'}.{ext}"


def export_top_tags(
    messages: pd.DataFrame | None, limit: int = 5
) -> list[dict[str, Any]]:
    """Топ тегов для PNG/DOCX/PDF — та же выборка, что и превью на «Обзоре»
    («Что включить в отчёт», client_insights_ui.top_client_tags) и карточка
    для ИИ (ai_summary._tags_block): все три берут top_report_tags."""
    if messages is None or not isinstance(messages, pd.DataFrame) or messages.empty:
        return []
    try:
        stats = top_report_tags(messages, limit=limit)
    except Exception:  # noqa: BLE001 — выгрузка без топ-тегов лучше, чем никакая
        LOGGER.warning("Топ-теги для выгрузки не посчитались", exc_info=True)
        return []
    result: list[dict[str, Any]] = []
    for _, row in stats.iterrows():
        tag = str(row.get("Тег") or "").strip()
        if not tag:
            continue
        result.append(
            {
                "name": tag,
                "messages": int(row.get("Сообщений", 0) or 0),
                "audience": int(row.get("Аудитория", 0) or 0),
                "reach": int(row.get("Охват", 0) or 0),
                "engagement": int(row.get("Вовлеченность", 0) or 0),
            }
        )
    return result


def export_top_events(
    events: pd.DataFrame | None, limit: int = 5
) -> list[dict[str, Any]]:
    """Топ инфоповодов для PNG/DOCX/PDF — та же выборка, что top_client_events
    и ai_summary._events_block (см. export_top_tags)."""
    work = top_report_events(events, limit=limit)
    if work.empty:
        return []
    title_col = event_title_column(work)
    count_col = first_existing_col(work, ["message_count", "messages", "Сообщений"])
    reach_col = first_existing_col(
        work, ["views", "reach", "Охват", "Просмотры", "Просмотров"]
    )
    engagement_col = first_existing_col(
        work, ["engagement", "Вовлеченность", "Вовлечённость"]
    )
    result: list[dict[str, Any]] = []
    for _, row in work.iterrows():
        name = str(row.get(title_col) or "").strip() if title_col else ""
        if not name:
            continue
        result.append(
            {
                "name": name,
                "messages": int(row.get(count_col, 0) or 0) if count_col else 0,
                "reach": int(row.get(reach_col, 0) or 0) if reach_col else 0,
                "engagement": (
                    int(row.get(engagement_col, 0) or 0) if engagement_col else 0
                ),
            }
        )
    return result


def summary_highlights(summary_text: str, limit: int = 4) -> list[str]:
    lines: list[str] = []
    for raw in str(summary_text or "").replace("\r", "\n").split("\n"):
        kind, line = _classify_summary_line(raw)
        # Подзаголовки ("## ...") - не наблюдение, а название раздела; сами
        # по себе в "Главное" не годятся (там и так рядом отдельный блок
        # метрик/тональности - см. _VISUAL_SECTIONS).
        if kind == "heading":
            continue
        if line and line not in lines:
            lines.append(line)
        if len(lines) >= limit:
            break
    if lines:
        return lines
    return ["Саммари пока не заполнено."]


def _load_report_logo_bytes(branding: dict[str, Any] | None) -> bytes:
    """Load logo bytes from Storage path or, as fallback, from public URL."""
    branding = branding or {}
    storage_path = str(branding.get("logo_storage_path") or "").strip()
    if storage_path:
        try:
            return download_storage_file(storage_path)
        except Exception:  # noqa: BLE001 — ниже есть запасной путь по ссылке
            LOGGER.warning("Логотип из Storage не скачался: %s", storage_path)
    url = str(branding.get("logo_url") or "").strip()
    if url.startswith(("http://", "https://")):
        try:
            from urllib.request import urlopen

            with urlopen(
                url, timeout=6
            ) as response:  # nosec - user-provided report asset URL
                return response.read()
        except Exception:  # noqa: BLE001 — отчёт без логотипа лучше, чем без отчёта
            LOGGER.warning("Логотип по ссылке не скачался: %s", url)
            return b""
    return b""


def _logo_image_from_payload(payload: dict[str, Any]):
    logo_bytes = payload.get("logo_bytes") or b""
    if not logo_bytes:
        return None
    try:
        from PIL import Image as PILImage

        return PILImage.open(BytesIO(logo_bytes)).convert("RGBA")
    except Exception:  # noqa: BLE001 — битая картинка не должна ломать выгрузку
        LOGGER.warning("Логотип не открылся как изображение", exc_info=True)
        return None


def resolve_report_sections(sections: list[str] | None) -> list[str]:
    """Нормализовать выбор блоков конструктора: неизвестные id отбрасываются,
    пустой/некорректный выбор откатывается на полный набор по умолчанию —
    отчёт без единого блока никому не нужен и обычно означает баг вызова,
    а не осознанный выбор аналитика."""
    if not sections:
        return list(DEFAULT_REPORT_SECTIONS)
    result = [s for s in sections if s in REPORT_SECTION_OPTIONS]
    return result or list(DEFAULT_REPORT_SECTIONS)


def summary_export_payload(
    project_name: str,
    period_label: str,
    summary_text: str,
    metrics: dict[str, Any],
    messages: pd.DataFrame | None = None,
    events_agg: pd.DataFrame | None = None,
    *,
    branding: dict[str, Any] | None = None,
    sections: list[str] | None = None,
) -> dict[str, Any]:
    sent = metrics.get("sentiment", {}) if isinstance(metrics, dict) else {}
    sections = resolve_report_sections(sections)
    branding = report_branding_from_project_settings(
        {"report_branding": branding or {}}, project_name=project_name
    )
    return {
        "project_name": project_name,
        "client_name": branding.get("client_name") or project_name,
        "report_title": branding.get("report_title") or "Дайджест упоминаний",
        "accent_color": branding.get("accent_color") or "#2563eb",
        "background_color": branding.get("background_color") or "#ffffff",
        "footer_text": branding.get("footer_text") or "",
        "logo_url": branding.get("logo_url") or "",
        "logo_storage_path": branding.get("logo_storage_path") or "",
        "logo_filename": branding.get("logo_filename") or "",
        "logo_mime_type": branding.get("logo_mime_type") or "",
        "logo_bytes": _load_report_logo_bytes(branding),
        "sections": sections,
        "period_label": period_label,
        "summary_text": clean_summary_for_export(summary_text),
        "summary_highlights": summary_highlights(summary_text),
        "messages": int(metrics.get("messages", 0) or 0),
        "audience": int(metrics.get("audience", 0) or 0),
        "reach": int(metrics.get("reach", 0) or 0),
        "engagement": int(metrics.get("engagement", 0) or 0),
        "positive": int(sent.get("positive", 0) or 0),
        "neutral": int(sent.get("neutral", 0) or 0),
        "negative": int(sent.get("negative", 0) or 0),
        "total": int(sent.get("total", 0) or 0),
        "comparison_sequence": metrics.get("comparison_sequence") or [],
        # limit=5: столько же всегда и рисуют PNG/DOCX/PDF (items[:5]) -
        # раньше "полный" шаблон запрашивал 8, но лишние 3 нигде не
        # показывались, просто отбрасывались слоем отрисовки.
        "top_tags": export_top_tags(messages, limit=5),
        "top_events": export_top_events(events_agg, limit=5),
        "created_at": datetime.now().strftime("%d.%m.%Y %H:%M"),
    }


def _short_label(value: Any, max_len: int = 34) -> str:
    text = str(value or "").strip()
    return text if len(text) <= max_len else text[: max_len - 1].rstrip() + "…"


def _metric_delta_for_export(current: Any, previous: Any) -> str:
    try:
        cur = float(current or 0)
        prev = float(previous or 0)
    except (TypeError, ValueError):
        return ""
    diff = cur - prev
    if prev:
        pct = diff / abs(prev) * 100
        return f"{diff:+,.0f} / {pct:+.1f}%".replace(",", " ")
    if diff:
        return f"{diff:+,.0f}".replace(",", " ")
    return "0"


def _adaptive_font_size(value: str, *, base: int = 15, min_size: int = 10) -> int:
    """Keep large metric values readable inside report cards."""
    length = len(str(value or ""))
    if length > 15:
        return max(min_size, base - 5)
    if length > 12:
        return max(min_size, base - 4)
    if length > 9:
        return max(min_size, base - 2)
    return base


def _draw_export_card(
    ax,
    x: float,
    y: float,
    w: float,
    h: float,
    title: str,
    value: str,
    subtitle: str = "",
    *,
    accent_color: str = "#2563eb",
) -> None:
    from matplotlib.patches import FancyBboxPatch

    patch = FancyBboxPatch(
        (x, y),
        w,
        h,
        boxstyle="round,pad=0.014,rounding_size=0.018",
        linewidth=0.9,
        edgecolor="#d9e0ea",
        facecolor="#f8fafc",
    )
    ax.add_patch(patch)
    ax.plot(
        [x + 0.016, x + 0.016],
        [y + 0.022, y + h - 0.022],
        color=accent_color,
        linewidth=2.2,
    )
    ax.text(
        x + 0.040,
        y + h - 0.026,
        title,
        fontsize=8.8,
        color="#4b5563",
        va="top",
        ha="left",
    )
    ax.text(
        x + 0.040,
        y + h * 0.50,
        value,
        fontsize=_adaptive_font_size(value, base=15, min_size=10),
        color="#111827",
        va="center",
        ha="left",
        fontweight="bold",
    )
    if subtitle:
        ax.text(
            x + 0.040,
            y + 0.018,
            subtitle,
            fontsize=7.4,
            color="#6b7280",
            va="bottom",
            ha="left",
        )


# Отступ до подписи в футере (см. её fontsize/позицию ниже) - последний
# блок ("Главное") не должен налезать на неё, даже если саммари длинное и
# даёт 4 длинных пункта, каждый на 2 строки.
_FOOTER_CLEARANCE = 0.075


def _draw_metrics_section(ax, payload, comparison, accent, top: float) -> float:
    if len(comparison) >= 2:
        previous, current = comparison[-2], comparison[-1]
        metric_cards = [
            (
                "Сообщения",
                current.get("messages", 0),
                _metric_delta_for_export(
                    current.get("messages", 0), previous.get("messages", 0)
                ),
            ),
            (
                "Аудитория",
                current.get("audience", 0),
                _metric_delta_for_export(
                    current.get("audience", 0), previous.get("audience", 0)
                ),
            ),
            (
                "Охват",
                current.get("reach", 0),
                _metric_delta_for_export(
                    current.get("reach", 0), previous.get("reach", 0)
                ),
            ),
            (
                "Вовлеченность",
                current.get("engagement", 0),
                _metric_delta_for_export(
                    current.get("engagement", 0), previous.get("engagement", 0)
                ),
            ),
        ]
        ax.text(
            0.060,
            top,
            f"Последний период: {_short_label(current.get('label'), 48)}",
            fontsize=8.2,
            color="#6b7280",
            va="top",
            ha="left",
        )
    else:
        metric_cards = [
            ("Сообщения", payload.get("messages", 0), ""),
            ("Аудитория", payload.get("audience", 0), ""),
            ("Охват", payload.get("reach", 0), ""),
            ("Вовлеченность", payload.get("engagement", 0), ""),
        ]

    xs = [0.060, 0.525]
    ys = [top - 0.107, top - 0.227]
    for idx, (title, value, subtitle) in enumerate(metric_cards):
        _draw_export_card(
            ax,
            xs[idx % 2],
            ys[idx // 2],
            0.405,
            0.095,
            title,
            format_int(value),
            f"к пред. периоду: {subtitle}" if subtitle else "",
            accent_color=accent,
        )
    # 0.227 - высота самого блока (до низа второй строки карточек), + 0.050 -
    # зазор до заголовка следующего блока в исходной раскладке (0.862 -> 0.585).
    return top - 0.277


def _draw_sentiment_section(ax, fig, payload, comparison, top: float) -> float:
    from matplotlib.patches import Rectangle

    total = max(1, int(payload.get("total", 0) or 0))
    pos = int(payload.get("positive", 0) or 0)
    neu = int(payload.get("neutral", 0) or 0)
    neg = int(payload.get("negative", 0) or 0)
    if len(comparison) >= 2:
        sent = comparison[-1].get("sentiment", {}) or {}
        total = max(1, int(sent.get("total", 0) or 0))
        pos = int(sent.get("positive", 0) or 0)
        neu = int(sent.get("neutral", 0) or 0)
        neg = int(sent.get("negative", 0) or 0)

    ax.text(
        0.060,
        top,
        "Тональность",
        fontsize=12,
        fontweight="bold",
        color="#111827",
        va="top",
        ha="left",
    )
    pie_bottom = top - 0.158
    pie_ax = fig.add_axes([0.070, pie_bottom, 0.220, 0.145])
    pie_ax.axis("equal")
    values = [max(pos, 0), max(neu, 0), max(neg, 0)]
    # Те же цвета, что и на живом дашборде (services/chart_style.py) - иначе
    # тональность выглядела бы разными оттенками зелёного/красного на экране
    # и в выгруженном PNG/PDF/DOCX одного и того же периода.
    colors = list(SENTIMENT_COLOR_RANGE)
    labels = ["Позитив", "Нейтрал", "Негатив"]
    if sum(values) <= 0:
        values = [1]
        pie_colors = ["#d1d5db"]
    else:
        pie_colors = colors
    pie_ax.pie(
        values,
        colors=pie_colors,
        startangle=90,
        counterclock=False,
        wedgeprops={"width": 0.42, "edgecolor": "white"},
    )
    pie_ax.text(
        0,
        0.05,
        format_int(total),
        ha="center",
        va="center",
        fontsize=12.5,
        fontweight="bold",
        color="#111827",
    )
    pie_ax.text(
        0, -0.13, "сообщений", ha="center", va="center", fontsize=7.5, color="#6b7280"
    )
    pie_ax.set_xticks([])
    pie_ax.set_yticks([])

    y0 = top - 0.040
    for i, (lab, val, col) in enumerate(zip(labels, [pos, neu, neg], colors)):
        yy = y0 - i * 0.041
        ax.add_patch(
            Rectangle(
                (0.330, yy - 0.010), 0.014, 0.014, facecolor=col, edgecolor="none"
            )
        )
        ax.text(0.352, yy, lab, fontsize=9.2, color="#111827", va="center", ha="left")
        ax.text(
            0.490,
            yy,
            format_int(val),
            fontsize=9.2,
            color="#111827",
            va="center",
            ha="right",
            fontweight="bold",
        )
        ax.text(
            0.510,
            yy,
            percent_text(val, total),
            fontsize=8.4,
            color="#6b7280",
            va="center",
            ha="left",
        )
    # pie_bottom (top - 0.158) - низ самого донат-графика; + 0.045 - зазор до
    # заголовка следующего блока в исходной раскладке (0.585 -> 0.382... то
    # есть до низа доната остаётся 0.427, а следующий блок стартовал на 0.382).
    return pie_bottom - 0.045


def _draw_top_lists_section(
    ax, payload, top: float, *, show_tags: bool, show_events: bool
) -> float:
    if show_tags:
        top_tags = payload.get("top_tags") or []
        ax.text(
            0.060,
            top,
            "Топ тегов",
            fontsize=11.5,
            fontweight="bold",
            color="#111827",
            va="top",
            ha="left",
        )
        y = top - 0.028
        if top_tags:
            for item in top_tags[:5]:
                ax.text(
                    0.070,
                    y,
                    f"• {_short_label(item.get('name'), 28)}",
                    fontsize=8.4,
                    color="#111827",
                    va="top",
                    ha="left",
                )
                ax.text(
                    0.430,
                    y,
                    f"{format_int(item.get('messages', 0))} сообщ.",
                    fontsize=7.8,
                    color="#6b7280",
                    va="top",
                    ha="right",
                )
                y -= 0.028
        else:
            ax.text(
                0.070,
                y,
                "Нет тегов для отображения",
                fontsize=8.4,
                color="#6b7280",
                va="top",
                ha="left",
            )

    if show_events:
        top_events = payload.get("top_events") or []
        ax.text(
            0.525,
            top,
            "Топ инфоповодов",
            fontsize=11.5,
            fontweight="bold",
            color="#111827",
            va="top",
            ha="left",
        )
        y = top - 0.028
        if top_events:
            for item in top_events[:5]:
                ax.text(
                    0.535,
                    y,
                    f"• {_short_label(item.get('name'), 29)}",
                    fontsize=8.4,
                    color="#111827",
                    va="top",
                    ha="left",
                )
                ax.text(
                    0.935,
                    y,
                    f"{format_int(item.get('messages', 0))} сообщ.",
                    fontsize=7.8,
                    color="#6b7280",
                    va="top",
                    ha="right",
                )
                y -= 0.028
        else:
            ax.text(
                0.535,
                y,
                "Нет инфоповодов для отображения",
                fontsize=8.4,
                color="#6b7280",
                va="top",
                ha="left",
            )
    # Фиксированная высота блока (под 5 позиций) вне зависимости от того,
    # сколько реально показано, - так же, как было в исходной раскладке.
    return top - 0.177


def _draw_highlights_section(ax, payload, top: float) -> float:
    ax.text(
        0.060,
        top,
        "Главное",
        fontsize=11.5,
        fontweight="bold",
        color="#111827",
        va="top",
        ha="left",
    )
    summary_y = top - 0.027
    line_count = 0
    # Раньше единственным ограничителем было "не больше 8 строк" - на
    # длинном автосаммари (4 пункта по 2 строки) это давало текст, который
    # реально наезжал на подпись в футере (она стоит на фиксированной
    # 0.045 независимо от того, сколько текста выше). Останавливаемся,
    # как только следующая строка попала бы в зону футера, а не только по
    # счётчику строк.
    for block in (payload.get("summary_highlights") or [])[:4]:
        if summary_y < _FOOTER_CLEARANCE:
            break
        wrapped = textwrap.wrap(str(block), width=86) or [str(block)]
        bullet = True
        for seg in wrapped[:2]:
            if summary_y < _FOOTER_CLEARANCE:
                break
            prefix = "• " if bullet else "  "
            ax.text(
                0.070,
                summary_y,
                prefix + seg,
                fontsize=8.4,
                color="#111827",
                va="top",
                ha="left",
            )
            summary_y -= 0.021
            line_count += 1
            bullet = False
            if line_count >= 8:
                break
        summary_y -= 0.004
        if line_count >= 8:
            break
    return summary_y


def generate_summary_infographic_png(payload: dict[str, Any]) -> bytes:
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib.patches import Rectangle
    except Exception as exc:
        raise RuntimeError(
            "Для инфографики нужен matplotlib>=3.8 в requirements.txt."
        ) from exc

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 9,
            "axes.unicode_minus": False,
        }
    )
    fig = plt.figure(figsize=(8.27, 11.69), dpi=170)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")

    background = valid_hex_color(payload.get("background_color"), "#ffffff")
    accent = valid_hex_color(payload.get("accent_color"), "#2563eb")
    fig.patch.set_facecolor(background)
    ax.add_patch(Rectangle((0, 0), 1, 1, facecolor=background, edgecolor="none"))

    project = _short_label(
        payload.get("client_name") or payload.get("project_name") or "Проект", 42
    )
    report_title = _short_label(
        payload.get("report_title") or "Дайджест упоминаний", 44
    )
    period = _short_label(payload.get("period_label") or "выбранный период", 56)
    created = str(payload.get("created_at") or "")
    comparison = payload.get("comparison_sequence") or []
    sections = set(resolve_report_sections(payload.get("sections")))

    # Header
    ax.add_patch(
        Rectangle((0, 0.885), 1, 0.115, facecolor=accent, edgecolor="none", alpha=0.96)
    )
    ax.text(
        0.060,
        0.965,
        report_title,
        fontsize=11.5,
        fontweight="bold",
        color="white",
        va="top",
        ha="left",
    )
    ax.text(
        0.060,
        0.935,
        project,
        fontsize=19,
        fontweight="bold",
        color="white",
        va="top",
        ha="left",
    )
    ax.text(
        0.060,
        0.904,
        period,
        fontsize=8.6,
        color="#e5e7eb",
        va="top",
        ha="left",
    )
    ax.text(0.935, 0.965, created, fontsize=8.0, color="#e5e7eb", va="top", ha="right")

    logo_img = _logo_image_from_payload(payload)
    if logo_img is not None:
        from matplotlib.patches import FancyBboxPatch

        box = FancyBboxPatch(
            (0.785, 0.902),
            0.150,
            0.055,
            boxstyle="round,pad=0.006,rounding_size=0.010",
            linewidth=0,
            facecolor="white",
            alpha=0.94,
        )
        ax.add_patch(box)
        logo_ax = fig.add_axes([0.795, 0.910, 0.130, 0.040])
        logo_ax.imshow(logo_img)
        logo_ax.axis("off")

    # Тело инфографики - курсор сверху вниз: каждый включённый блок рисуется
    # от текущего cursor и возвращает позицию, где должен начаться следующий
    # (зазор до следующего блока уже включён в возврат - см. комментарии в
    # каждой _draw_*_section). Выключенный блок просто не сдвигает курсор -
    # следующий встаёт на его место, без дыр.
    cursor = 0.862
    show_tags = "top_tags" in sections
    show_events = "top_events" in sections

    if "metrics" in sections:
        cursor = _draw_metrics_section(ax, payload, comparison, accent, cursor)

    if "sentiment" in sections:
        cursor = _draw_sentiment_section(ax, fig, payload, comparison, cursor)

    if show_tags or show_events:
        cursor = _draw_top_lists_section(
            ax, payload, cursor, show_tags=show_tags, show_events=show_events
        )

    if "highlights" in sections:
        cursor = _draw_highlights_section(ax, payload, cursor)

    if not (sections & _VISUAL_SECTIONS):
        ax.text(
            0.060,
            cursor,
            "Все аналитические блоки отключены в настройках выгрузки.",
            fontsize=9,
            color="#6b7280",
            va="top",
            ha="left",
        )

    footer_text = str(
        payload.get("footer_text")
        or "Инфографика сформирована автоматически на основе выбранного периода и текущего саммари."
    )
    ax.text(
        0.060,
        0.045,
        _short_label(footer_text, 110),
        fontsize=7.8,
        color="#6b7280",
        va="bottom",
        ha="left",
    )

    out = BytesIO()
    # Do not use bbox_inches="tight": it changes image geometry and can cause PDF scaling/cropping.
    fig.savefig(out, format="png", facecolor=background)
    plt.close(fig)
    return out.getvalue()


def _docx_bottom_border(paragraph, color_hex: str, size: int = 16) -> None:
    """Цветная линия под абзацем - тот же приём, что акцентная полоса в
    шапке PNG-инфографики, только средствами Word (нет прямого API, поэтому
    через oxml - стандартный, задокументированный обходной путь)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color_hex.lstrip("#"))
    borders.append(bottom)
    p_pr.append(borders)


def generate_summary_docx(payload: dict[str, Any]) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt, Inches, Cm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except Exception as exc:
        raise RuntimeError(
            "Для выгрузки Word добавьте python-docx в requirements.txt."
        ) from exc

    accent = valid_hex_color(payload.get("accent_color"), "#2563eb")
    accent_rgb = RGBColor.from_string(accent.lstrip("#"))
    muted_rgb = RGBColor.from_string("6b7280")
    ink_rgb = RGBColor.from_string("111827")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    # Тот же шрифт, что и в PNG-инфографике ("font.family": "DejaVu Sans" в
    # generate_summary_infographic_png) - раньше здесь стоял Arial, и страницы
    # саммари визуально не совпадали с картинкой на первой странице того же
    # документа. Шрифт не встраивается в .docx (в отличие от PDF, где он
    # зашит через TTFont) - если у читателя его нет, Word подставит похожий
    # рубленый шрифт, это мягкая деградация, не поломка.
    docx_font = "DejaVu Sans"
    styles = doc.styles
    styles["Normal"].font.name = docx_font
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].font.color.rgb = ink_rgb
    styles["Heading 1"].font.name = docx_font
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.color.rgb = accent_rgb
    # Тот же акцентный цвет, что заголовки блоков и полоса на карточках
    # метрик в PNG-инфографике - страницы саммари не должны выглядеть
    # отдельным, неоформленным документом рядом с картинкой на первой странице.
    styles["Heading 2"].font.name = docx_font
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 2"].font.color.rgb = accent_rgb
    # Подзаголовки ВНУТРИ текста саммари (build_auto_summary размечает их
    # "## ...") - на ступень мельче "Heading 2", чтобы не спорить визуально
    # с заголовками разделов отчёта.
    styles["Heading 3"].font.name = docx_font
    styles["Heading 3"].font.size = Pt(11.5)
    styles["Heading 3"].font.color.rgb = accent_rgb

    p = doc.add_paragraph()
    r = p.add_run(str(payload.get("report_title") or "Дайджест упоминаний"))
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = accent_rgb
    p.paragraph_format.space_after = Pt(4)

    p2 = doc.add_paragraph()
    r2 = p2.add_run(
        str(payload.get("client_name") or payload.get("project_name") or "Проект")
    )
    r2.bold = True
    r2.font.size = Pt(13)
    r2.font.color.rgb = ink_rgb
    p2.paragraph_format.space_after = Pt(6)

    meta = doc.add_paragraph()
    meta_lines = [
        f"Период: {payload.get('period_label') or 'выбранный период'}",
        f"Дата выгрузки: {payload.get('created_at') or ''}",
    ]
    for i, line in enumerate(meta_lines):
        text = line if i == len(meta_lines) - 1 else line + "\n"
        run = meta.add_run(text)
        run.font.color.rgb = muted_rgb
        run.font.size = Pt(9.5)
    meta.paragraph_format.space_after = Pt(10)
    _docx_bottom_border(meta, accent)

    sections = set(resolve_report_sections(payload.get("sections")))
    has_visual_sections = bool(_VISUAL_SECTIONS & sections)

    if has_visual_sections:
        try:
            infographic_png = generate_summary_infographic_png(payload)
            pic_p = doc.add_paragraph()
            pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = pic_p.add_run()
            run.add_picture(BytesIO(infographic_png), width=Inches(6.4))
            doc.add_page_break()
        except Exception as exc:  # noqa: BLE001 — Word без инфографики лучше, чем без Word
            # Клиент получит документ и не узнает, что страницы не хватает, —
            # поэтому владелец должен узнать вместо него.
            LOGGER.warning("Инфографика для Word не собралась", exc_info=True)
            report_failure("выгрузка Word: инфографика не собралась", exc)

    total = max(1, int(payload.get("total", 0) or 0))
    if "metrics" in sections or "sentiment" in sections:
        doc.add_heading("Основные метрики", level=2)
        if "metrics" in sections:
            doc.add_paragraph(
                f"Сообщений — {format_int(payload.get('messages', 0))}; "
                f"аудитория — {format_int(payload.get('audience', 0))}; "
                f"охват — {format_int(payload.get('reach', 0))}; "
                f"вовлеченность — {format_int(payload.get('engagement', 0))}."
            )
        if "sentiment" in sections:
            doc.add_paragraph(
                f"Тональность: позитив — {percent_text(int(payload.get('positive', 0) or 0), total)}; "
                f"нейтрал — {percent_text(int(payload.get('neutral', 0) or 0), total)}; "
                f"негатив — {percent_text(int(payload.get('negative', 0) or 0), total)}."
            )

    if "top_tags" in sections or "top_events" in sections:
        doc.add_heading("Что включить в отчет", level=2)
        if "top_tags" in sections:
            top_tags = payload.get("top_tags") or []
            if top_tags:
                doc.add_paragraph(
                    "Топ тегов: "
                    + "; ".join(
                        str(x.get("name") or "") for x in top_tags[:5] if x.get("name")
                    )
                )
        if "top_events" in sections:
            top_events = payload.get("top_events") or []
            if top_events:
                doc.add_paragraph(
                    "Топ инфоповодов: "
                    + "; ".join(
                        str(x.get("name") or "")
                        for x in top_events[:5]
                        if x.get("name")
                    )
                )

    if "summary_text" in sections:
        doc.add_heading("Саммари периода", level=2)
        for raw in str(payload.get("summary_text") or "").split("\n"):
            kind, text = _classify_summary_line(raw)
            if not text:
                continue
            if kind == "heading":
                doc.add_heading(text, level=3)
            elif kind == "bullet":
                para = doc.add_paragraph(text, style="List Bullet")
                para.paragraph_format.space_after = Pt(2)
            else:
                para = doc.add_paragraph(text)
                para.paragraph_format.space_after = Pt(4)

    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _pdf_font_candidates() -> list[tuple[str, str | None]]:
    """Return regular/bold TTF candidates that support Cyrillic.

    Streamlit Cloud images do not always include system DejaVu/Noto fonts. If we
    fall back to ReportLab core Helvetica, Cyrillic is rendered as black squares.
    To make PDF export portable, we also look for the DejaVu Sans bundled with
    matplotlib when the package is installed.
    """
    candidates: list[tuple[str, str | None]] = []

    env_path = os.getenv("PLATFORM_PDF_FONT_PATH", "").strip()
    if env_path:
        candidates.append((env_path, None))

    # Optional project-local fonts. Do not commit font files unless licensing is clear;
    # this path is only for private deployments that provide their own font.
    here = Path(__file__).resolve().parent.parent
    candidates.extend(
        [
            (
                str(here / "assets" / "fonts" / "DejaVuSans.ttf"),
                str(here / "assets" / "fonts" / "DejaVuSans-Bold.ttf"),
            ),
            (
                str(here.parent / "assets" / "fonts" / "DejaVuSans.ttf"),
                str(here.parent / "assets" / "fonts" / "DejaVuSans-Bold.ttf"),
            ),
        ]
    )

    # Common Linux/Streamlit Cloud system fonts.
    candidates.extend(
        [
            (
                "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
            ),
            (
                "/usr/share/fonts/truetype/dejavu/DejaVuSansCondensed.ttf",
                "/usr/share/fonts/truetype/dejavu/DejaVuSansCondensed-Bold.ttf",
            ),
            (
                "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",
                "/usr/share/fonts/truetype/noto/NotoSans-Bold.ttf",
            ),
            (
                "/usr/share/fonts/truetype/freefont/FreeSans.ttf",
                "/usr/share/fonts/truetype/freefont/FreeSansBold.ttf",
            ),
            (
                "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
                "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf",
            ),
            (
                "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
                "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
            ),
        ]
    )

    try:
        import matplotlib  # type: ignore

        mpl_fonts = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
        candidates.append(
            (str(mpl_fonts / "DejaVuSans.ttf"), str(mpl_fonts / "DejaVuSans-Bold.ttf"))
        )
    except Exception:  # noqa: BLE001 — matplotlib необязателен, это запасной шрифт
        pass

    return candidates


def _pdf_font_name() -> str:
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except Exception as exc:
        raise RuntimeError(
            "Для выгрузки PDF добавьте reportlab в requirements.txt."
        ) from exc

    for regular_path, bold_path in _pdf_font_candidates():
        regular = Path(regular_path)
        if not regular.exists():
            continue
        try:
            pdfmetrics.registerFont(TTFont("PlatformSans", str(regular)))
            bold = Path(bold_path) if bold_path else None
            if bold is not None and bold.exists():
                pdfmetrics.registerFont(TTFont("PlatformSans-Bold", str(bold)))
                try:
                    pdfmetrics.registerFontFamily(
                        "PlatformSans",
                        normal="PlatformSans",
                        bold="PlatformSans-Bold",
                        italic="PlatformSans",
                        boldItalic="PlatformSans-Bold",
                    )
                except Exception:  # noqa: BLE001 — без семейства жирный заменится обычным
                    pass
            return "PlatformSans"
        except Exception:  # noqa: BLE001 — кандидат не подошёл, пробуем следующий
            continue

    raise RuntimeError(
        "Не найден TTF-шрифт с поддержкой кириллицы для PDF. "
        "Проверьте, что установлен matplotlib>=3.8 или задайте PLATFORM_PDF_FONT_PATH."
    )


def _pdf_bold_font_name(font_name: str) -> str:
    """Имя жирного варианта, если он реально зарегистрирован - иначе тот же
    обычный шрифт (жирный текст останется обычным, не подменится Helvetica)."""
    from reportlab.pdfbase import pdfmetrics

    bold_name = f"{font_name}-Bold"
    try:
        pdfmetrics.getFont(bold_name)
        return bold_name
    except Exception:  # noqa: BLE001 - нет жирного варианта, работаем без него
        return font_name


class _PdfMetricsBlock:
    """Строка из 4 карточек метрик - те же данные и цвета, что на месте PNG
    раньше, но нарисованы напрямую на canvas PDF, тем же шрифтом, что и
    остальной текст страницы."""

    def __init__(self, cards, subtitle, accent_color, ink_color, muted_color, font_name, bold_font_name):
        self.cards = cards  # [(title, value, delta_subtitle), ...] до 4 штук
        self.subtitle = subtitle
        self.accent_color = accent_color
        self.ink_color = ink_color
        self.muted_color = muted_color
        self.font_name = font_name
        self.bold_font_name = bold_font_name

    def height(self, width):
        from reportlab.lib.units import cm

        return (0.5 * cm if self.subtitle else 0) + 2 * (2.05 * cm) + 0.3 * cm

    def draw(self, canv, x, y, width):
        """Рисует блок так, что (x, y) - левый ВЕРХНИЙ угол; возвращает y низа блока."""
        from reportlab.lib import colors as rl_colors
        from reportlab.lib.units import cm

        gap = 0.3 * cm
        card_w = (width - gap) / 2
        card_h = 2.05 * cm
        top = y
        if self.subtitle:
            canv.setFillColor(self.muted_color)
            canv.setFont(self.font_name, 8.2)
            canv.drawString(x, top - 10, self.subtitle)
            top -= 0.5 * cm
        positions = [
            (x, top - card_h),
            (x + card_w + gap, top - card_h),
            (x, top - card_h - gap - card_h),
            (x + card_w + gap, top - card_h - gap - card_h),
        ]
        for (cx, cy), (card_title, value, subtitle) in zip(positions, self.cards):
            canv.setFillColor(rl_colors.HexColor("#f8fafc"))
            canv.setStrokeColor(rl_colors.HexColor("#d9e0ea"))
            canv.setLineWidth(0.7)
            canv.roundRect(cx, cy, card_w, card_h, 6, stroke=1, fill=1)
            canv.setStrokeColor(self.accent_color)
            canv.setLineWidth(2.2)
            canv.line(cx + 10, cy + 10, cx + 10, cy + card_h - 10)
            canv.setFillColor(rl_colors.HexColor("#4b5563"))
            canv.setFont(self.font_name, 9)
            canv.drawString(cx + 20, cy + card_h - 20, card_title)
            canv.setFillColor(self.ink_color)
            size = _adaptive_font_size(value, base=17, min_size=11)
            canv.setFont(self.bold_font_name, size)
            canv.drawString(cx + 20, cy + card_h / 2 - 6, value)
            if subtitle:
                canv.setFillColor(self.muted_color)
                canv.setFont(self.font_name, 7.4)
                canv.drawString(cx + 20, cy + 10, f"к пред. периоду: {subtitle}")
        return top - card_h - gap - card_h


class _PdfSentimentBlock:
    """Донат-диаграмма тональности + легенда - те же цвета, что и на живом
    дашборде (SENTIMENT_COLOR_RANGE), нарисованные через canvas.wedge вместо
    растровой картинки."""

    HEIGHT_CM = 3.6

    def __init__(self, values, colors_hex, labels, total, background_color, font_name, bold_font_name):
        self.values = values
        self.colors_hex = colors_hex
        self.labels = labels
        self.total = max(1, total)
        self.background_color = background_color
        self.font_name = font_name
        self.bold_font_name = bold_font_name

    def height(self, width):
        from reportlab.lib.units import cm

        return self.HEIGHT_CM * cm

    def draw(self, canv, x, y, width):
        from reportlab.lib import colors as rl_colors
        from reportlab.lib.units import cm

        block_h = self.height(width)
        r = block_h / 2 - 0.15 * cm
        cx, cy = x + r + 0.2 * cm, y - block_h / 2
        total_v = sum(self.values) or 1
        start = 90.0
        any_positive = any(v > 0 for v in self.values)
        colors_list = self.colors_hex if any_positive else ["#d1d5db"]
        values = self.values if any_positive else [1]
        for val, col in zip(values, colors_list):
            if val <= 0:
                continue
            extent = -360.0 * (val / total_v)
            canv.setFillColor(rl_colors.HexColor(col))
            canv.setStrokeColor(self.background_color)
            canv.setLineWidth(1.4)
            canv.wedge(cx - r, cy - r, cx + r, cy + r, start, extent, stroke=1, fill=1)
            start += extent
        canv.setFillColor(self.background_color)
        hole_r = r * 0.56
        canv.circle(cx, cy, hole_r, stroke=0, fill=1)
        canv.setFillColor(rl_colors.HexColor("#111827"))
        canv.setFont(self.bold_font_name, 13)
        canv.drawCentredString(cx, cy + 2, format_int(self.total))
        canv.setFillColor(rl_colors.HexColor("#6b7280"))
        canv.setFont(self.font_name, 7)
        canv.drawCentredString(cx, cy - 11, "сообщений")

        legend_x = cx + r + 1.1 * cm
        row_h = block_h / max(1, len(self.labels))
        ly = y - row_h / 2 + 3
        for lab, val, col in zip(self.labels, self.values, self.colors_hex):
            canv.setFillColor(rl_colors.HexColor(col))
            canv.rect(legend_x, ly - 5, 9, 9, stroke=0, fill=1)
            canv.setFillColor(rl_colors.HexColor("#111827"))
            canv.setFont(self.font_name, 9.2)
            canv.drawString(legend_x + 15, ly - 4, lab)
            canv.setFont(self.bold_font_name, 9.2)
            canv.drawRightString(legend_x + 3.9 * cm, ly - 4, format_int(val))
            canv.setFillColor(rl_colors.HexColor("#6b7280"))
            canv.setFont(self.font_name, 8.4)
            canv.drawString(legend_x + 4.05 * cm, ly - 4, percent_text(val, self.total))
            ly -= row_h
        return y - block_h


class _PdfTopListsBlock:
    """«Топ тегов» / «Топ инфоповодов» - две колонки, рисуются напрямую, без
    ручного textwrap: короткие строки, перенос не нужен (см. _short_label)."""

    HEIGHT_CM = 4.1

    def __init__(self, tags, events, show_tags, show_events, ink_color, muted_color, font_name, bold_font_name):
        self.tags = tags
        self.events = events
        self.show_tags = show_tags
        self.show_events = show_events
        self.ink_color = ink_color
        self.muted_color = muted_color
        self.font_name = font_name
        self.bold_font_name = bold_font_name

    def height(self, width):
        from reportlab.lib.units import cm

        return self.HEIGHT_CM * cm

    def _draw_column(self, canv, x, top, col_w, heading_text, items, name_key, count_label_fmt):
        from reportlab.lib import colors as rl_colors

        canv.setFillColor(self.ink_color)
        canv.setFont(self.bold_font_name, 11)
        canv.drawString(x, top - 12, heading_text)
        y = top - 32
        if items:
            for item in items[:5]:
                canv.setFillColor(self.ink_color)
                canv.setFont(self.font_name, 8.6)
                canv.drawString(x, y, f"• {_short_label(item.get('name'), 30)}")
                canv.setFillColor(self.muted_color)
                canv.setFont(self.font_name, 7.8)
                canv.drawRightString(
                    x + col_w, y, count_label_fmt(item)
                )
                y -= 17
        else:
            canv.setFillColor(self.muted_color)
            canv.setFont(self.font_name, 8.6)
            canv.drawString(x, y, "Нет данных для отображения")
        return y

    def draw(self, canv, x, y, width):
        from reportlab.lib.units import cm

        gap = 0.6 * cm
        col_w = (width - gap) / 2
        if self.show_tags:
            self._draw_column(
                canv,
                x,
                y,
                col_w,
                "Топ тегов",
                self.tags,
                "name",
                lambda item: f"{format_int(item.get('messages', 0))} сообщ.",
            )
        if self.show_events:
            self._draw_column(
                canv,
                x + col_w + gap,
                y,
                col_w,
                "Топ инфоповодов",
                self.events,
                "name",
                lambda item: f"{format_int(item.get('messages', 0))} сообщ.",
            )
        return y - self.height(width)


def generate_summary_pdf(payload: dict[str, Any]) -> bytes:
    try:
        from reportlab.lib import colors as rl_colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
        from reportlab.platypus.flowables import Flowable
        from xml.sax.saxutils import escape as xml_escape
    except Exception as exc:
        raise RuntimeError(
            "Для выгрузки PDF добавьте reportlab в requirements.txt."
        ) from exc

    font_name = _pdf_font_name()
    bold_font_name = _pdf_bold_font_name(font_name)
    accent = valid_hex_color(payload.get("accent_color"), "#2563eb")
    background = valid_hex_color(payload.get("background_color"), "#ffffff")
    accent_color = rl_colors.HexColor(accent)
    background_color = rl_colors.HexColor(background)
    ink_color = rl_colors.HexColor("#111827")
    muted_color = rl_colors.HexColor("#6b7280")

    sections = set(resolve_report_sections(payload.get("sections")))
    total = max(1, int(payload.get("total", 0) or 0))
    comparison = payload.get("comparison_sequence") or []

    # --- шапка/футер: рисуются на КАЖДОЙ странице через onPage-колбэк, а не
    # один раз как первая страница-картинка - если саммари длинное и уходит
    # на страницу 2+, брендирование не теряется. ---
    report_title_text = _short_label(
        payload.get("report_title") or "Дайджест упоминаний", 60
    )
    project_text = _short_label(
        payload.get("client_name") or payload.get("project_name") or "Проект", 46
    )
    meta_text = str(payload.get("period_label") or "выбранный период")
    created_text = str(payload.get("created_at") or "")
    footer_text = _short_label(
        str(
            payload.get("footer_text")
            or "Сформировано автоматически на основе выбранного периода и текущего саммари."
        ),
        140,
    )
    logo_img = _logo_image_from_payload(payload)
    HEADER_H = 2.7 * cm
    FOOTER_H = 0.9 * cm
    MARGIN = 1.7 * cm

    def _draw_page_frame(canv, doc):
        canv.saveState()
        page_w, page_h = doc.pagesize
        canv.setFillColor(background_color)
        canv.rect(0, 0, page_w, page_h, stroke=0, fill=1)
        canv.setFillColor(accent_color)
        canv.rect(0, page_h - HEADER_H, page_w, HEADER_H, stroke=0, fill=1)
        canv.setFillColor(rl_colors.white)
        canv.setFont(font_name, 10)
        canv.drawString(MARGIN, page_h - 0.85 * cm, report_title_text)
        canv.setFont(bold_font_name, 15)
        canv.drawString(MARGIN, page_h - 1.55 * cm, project_text)
        canv.setFont(font_name, 8.3)
        canv.setFillColor(rl_colors.HexColor("#e5e7eb"))
        canv.drawString(MARGIN, page_h - 2.05 * cm, meta_text)
        canv.setFont(font_name, 7.8)
        canv.drawRightString(page_w - MARGIN, page_h - 0.85 * cm, created_text)
        if logo_img is not None:
            try:
                from reportlab.lib.utils import ImageReader

                logo_w, logo_h = 2.6 * cm, 1.0 * cm
                canv.drawImage(
                    ImageReader(logo_img),
                    page_w - MARGIN - logo_w,
                    page_h - 2.0 * cm,
                    width=logo_w,
                    height=logo_h,
                    mask="auto",
                    preserveAspectRatio=True,
                    anchor="c",
                )
            except Exception:  # noqa: BLE001 - без логотипа PDF всё равно нужен
                LOGGER.warning("Логотип не встроился в PDF", exc_info=True)
        canv.setFillColor(muted_color)
        canv.setFont(font_name, 7.6)
        canv.drawString(MARGIN, FOOTER_H / 2, footer_text)
        canv.restoreState()

    out = BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=A4,
        leftMargin=MARGIN,
        rightMargin=MARGIN,
        topMargin=HEADER_H + 0.5 * cm,
        bottomMargin=FOOTER_H + 0.4 * cm,
    )
    base = getSampleStyleSheet()
    normal = ParagraphStyle(
        "PlatformNormal",
        parent=base["Normal"],
        fontName=font_name,
        fontSize=10,
        leading=14,
        textColor=ink_color,
    )
    heading = ParagraphStyle(
        "PlatformHeading",
        parent=normal,
        fontName=bold_font_name,
        fontSize=12,
        leading=16,
        spaceBefore=4,
        spaceAfter=6,
        textColor=accent_color,
    )
    # Подзаголовки ВНУТРИ текста саммари ("## ..." от build_auto_summary) -
    # на ступень мельче heading, чтобы не спорить визуально с заголовками
    # разделов отчёта ("Саммари периода" и т.д.).
    subheading = ParagraphStyle(
        "PlatformSubheading",
        parent=normal,
        fontName=bold_font_name,
        fontSize=10.5,
        leading=14,
        spaceBefore=8,
        spaceAfter=3,
        textColor=accent_color,
    )
    bullet_style = ParagraphStyle(
        "PlatformBullet",
        parent=normal,
        leftIndent=12,
        bulletIndent=0,
        spaceAfter=2,
    )

    # Обёртка, которая просто зовёт draw(canv, x, top, width) у одного из
    # блоков выше - так каждый блок (карточки/донат/топ-списки) остаётся
    # ОБЫЧНЫМ платипус-флоублом и участвует в общей вёрстке страницы:
    # не помещается - переносится на следующую, как любой другой абзац.
    class _BlockFlowable(Flowable):
        def __init__(self, block):
            Flowable.__init__(self)
            self.block = block
            self.width = 0
            self._h = 0

        def wrap(self, avail_width, avail_height):
            self.width = avail_width
            self._h = self.block.height(avail_width)
            return self.width, self._h

        def draw(self):
            self.block.draw(self.canv, 0, self._h, self.width)

    story: list[Any] = []

    if "metrics" in sections:
        if len(comparison) >= 2:
            previous, current = comparison[-2], comparison[-1]
            cards = [
                (
                    "Сообщения",
                    format_int(current.get("messages", 0)),
                    _metric_delta_for_export(current.get("messages", 0), previous.get("messages", 0)),
                ),
                (
                    "Аудитория",
                    format_int(current.get("audience", 0)),
                    _metric_delta_for_export(current.get("audience", 0), previous.get("audience", 0)),
                ),
                (
                    "Охват",
                    format_int(current.get("reach", 0)),
                    _metric_delta_for_export(current.get("reach", 0), previous.get("reach", 0)),
                ),
                (
                    "Вовлеченность",
                    format_int(current.get("engagement", 0)),
                    _metric_delta_for_export(current.get("engagement", 0), previous.get("engagement", 0)),
                ),
            ]
            subtitle = f"Последний период: {_short_label(current.get('label'), 48)}"
        else:
            cards = [
                ("Сообщения", format_int(payload.get("messages", 0)), ""),
                ("Аудитория", format_int(payload.get("audience", 0)), ""),
                ("Охват", format_int(payload.get("reach", 0)), ""),
                ("Вовлеченность", format_int(payload.get("engagement", 0)), ""),
            ]
            subtitle = ""
        story.append(
            _BlockFlowable(
                _PdfMetricsBlock(cards, subtitle, accent_color, ink_color, muted_color, font_name, bold_font_name)
            )
        )
        story.append(Spacer(1, 10))

    if "sentiment" in sections:
        pos = int(payload.get("positive", 0) or 0)
        neu = int(payload.get("neutral", 0) or 0)
        neg = int(payload.get("negative", 0) or 0)
        sent_total = total
        if len(comparison) >= 2:
            sent = comparison[-1].get("sentiment", {}) or {}
            sent_total = max(1, int(sent.get("total", 0) or 0))
            pos = int(sent.get("positive", 0) or 0)
            neu = int(sent.get("neutral", 0) or 0)
            neg = int(sent.get("negative", 0) or 0)
        story.append(Paragraph("Тональность", heading))
        story.append(
            _BlockFlowable(
                _PdfSentimentBlock(
                    [pos, neu, neg],
                    list(SENTIMENT_COLOR_RANGE),
                    ["Позитив", "Нейтрал", "Негатив"],
                    sent_total,
                    background_color,
                    font_name,
                    bold_font_name,
                )
            )
        )
        story.append(Spacer(1, 10))

    show_tags = "top_tags" in sections
    show_events = "top_events" in sections
    if show_tags or show_events:
        story.append(
            _BlockFlowable(
                _PdfTopListsBlock(
                    payload.get("top_tags") or [],
                    payload.get("top_events") or [],
                    show_tags,
                    show_events,
                    ink_color,
                    muted_color,
                    font_name,
                    bold_font_name,
                )
            )
        )
        story.append(Spacer(1, 10))

    if "highlights" in sections:
        story.append(Paragraph("Главное", heading))
        highlights = (payload.get("summary_highlights") or [])[:4]
        for block in highlights:
            story.append(Paragraph("• " + xml_escape(str(block)), normal))
        story.append(Spacer(1, 6))

    if "summary_text" in sections:
        if story:
            story.append(PageBreak())
        story.append(Paragraph("Саммари периода", heading))
        for raw in str(payload.get("summary_text") or "").split("\n"):
            kind, text = _classify_summary_line(raw)
            if not text:
                continue
            if kind == "heading":
                story.append(Paragraph(xml_escape(text), subheading))
            elif kind == "bullet":
                story.append(Paragraph(xml_escape(text), bullet_style, bulletText="•"))
            else:
                story.append(Paragraph(xml_escape(text), normal))

    if not story:
        # Аналитик снял вообще все разделы - пустой PDF выглядел бы как баг,
        # а не как осознанный (пустой) выбор.
        story.append(
            Paragraph("Все разделы отчёта отключены в настройках выгрузки.", normal)
        )

    doc.build(story, onFirstPage=_draw_page_frame, onLaterPages=_draw_page_frame)
    return out.getvalue()
