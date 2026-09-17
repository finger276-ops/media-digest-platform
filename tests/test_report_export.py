"""Проверка выгрузок саммари: PNG-инфографика, Word, PDF.

Раньше этот код жил внутри app.py и не имел собственных тестов. Проверяем,
что после переноса в services/report_export.py генерация по-прежнему
собирает валидные файлы (по сигнатуре формата), а не падает и не отдаёт
пустышку.
"""

import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
for _p in (REPO / "src", REPO / "scripts", REPO / "tests"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import pandas as pd  # noqa: E402
from docx import Document  # noqa: E402
from io import BytesIO  # noqa: E402

from services.dashboard_config import (  # noqa: E402
    DEFAULT_REPORT_SECTIONS,
    REPORT_SECTION_OPTIONS,
)
from services.metrics_compute import overview_metrics  # noqa: E402
from services.report_export import (  # noqa: E402
    export_top_events,
    export_top_tags,
    first_existing_col,
    generate_summary_docx,
    generate_summary_infographic_png,
    generate_summary_pdf,
    resolve_report_sections,
    safe_export_filename,
    summary_export_payload,
)

failures = []


def check(label, condition, detail=""):
    print(("  ✓ " if condition else "  ✗ ") + label + (f" — {detail}" if detail and not condition else ""))
    if not condition:
        failures.append(label)


MESSAGES = pd.DataFrame(
    [
        {
            "audience": 5000,
            "views": 12000,
            "engagement": 300,
            "sentiment": "позитив",
            "tags": "ТЕХНОНИКОЛЬ|Маркетплейсы",
            "event_title": "Запуск нового завода",
        },
        {
            "audience": 3000,
            "views": 8000,
            "engagement": 150,
            "sentiment": "негатив",
            "tags": "ТЕХНОНИКОЛЬ|Жалобы",
            "event_title": "Жалобы на монтаж",
        },
        {
            "audience": 2000,
            "views": 4000,
            "engagement": 50,
            "sentiment": "нейтрал",
            "tags": "Маркетплейсы",
            "event_title": "Отраслевая статистика",
        },
    ]
)

EVENTS_AGG = pd.DataFrame(
    [
        {"title": "Запуск нового завода", "message_count": 1, "views": 12000, "engagement": 300},
        {"title": "Жалобы на монтаж", "message_count": 1, "views": 8000, "engagement": 150},
        {"title": "Без сюжета", "message_count": 5, "views": 500, "engagement": 5},
    ]
)

BRANDING = {
    "client_name": "ТЕХНОНИКОЛЬ",
    "report_title": "Дайджест упоминаний",
    "accent_color": "#2563eb",
    "background_color": "#ffffff",
    "footer_text": "Подготовлено аналитическим отделом",
    # logo_* намеренно пустые — не должен идти запрос в Supabase Storage.
}

print("1. Вспомогательные функции экспорта")
check(
    "safe_export_filename убирает недопустимые символы",
    safe_export_filename("ТЕХНОНИКОЛЬ/тест", "24.04–30.04", "docx").endswith(".docx"),
)
check(
    "first_existing_col находит первую существующую колонку",
    first_existing_col(MESSAGES, ["missing", "tags"]) == "tags",
)
top_tags = export_top_tags(MESSAGES, limit=5)
check("export_top_tags вернул непустой список", bool(top_tags), str(top_tags))
check(
    "тег с наибольшим охватом первый",
    top_tags[0]["name"] == "ТЕХНОНИКОЛЬ",
    str(top_tags),
)
top_events = export_top_events(EVENTS_AGG, limit=5)
check(
    "export_top_events отфильтровал «Без сюжета»",
    all(e["name"] != "Без сюжета" for e in top_events),
    str(top_events),
)

print("2. Сборка payload из данных периода")
metrics = overview_metrics(MESSAGES)
payload = summary_export_payload(
    "ТЕХНОНИКОЛЬ",
    "24.04.2026–30.04.2026",
    "Ключевое событие — запуск нового завода.\nНегатив по монтажу требует внимания.",
    metrics,
    messages=MESSAGES,
    events_agg=EVENTS_AGG,
    report_template="full",
    branding=BRANDING,
)
check("payload содержит числа из metrics, а не выдуманные", payload["messages"] == 3, str(payload["messages"]))
check("payload несёт бренд-настройки без обращения к Storage", payload["client_name"] == "ТЕХНОНИКОЛЬ")
check("шаблон отчета сохранён", payload["report_template"] == "full")
check("топ-тегов и топ-инфоповодов собраны в payload", bool(payload["top_tags"]) and bool(payload["top_events"]))

print("3. PNG-инфографика")
try:
    png_bytes = generate_summary_infographic_png(payload)
    check("PNG не пустой", len(png_bytes) > 1000, f"{len(png_bytes)} bytes")
    check("PNG-сигнатура на месте", png_bytes[:8] == b"\x89PNG\r\n\x1a\n")
except Exception as exc:
    check("PNG сгенерирован без исключений", False, str(exc))

print("4. Word (.docx)")
try:
    docx_bytes = generate_summary_docx(payload)
    check("DOCX не пустой", len(docx_bytes) > 1000, f"{len(docx_bytes)} bytes")
    check("DOCX — это ZIP-контейнер (сигнатура PK)", docx_bytes[:2] == b"PK")
except Exception as exc:
    check("DOCX сгенерирован без исключений", False, str(exc))

print("5. PDF")
try:
    pdf_bytes = generate_summary_pdf(payload)
    check("PDF не пустой", len(pdf_bytes) > 1000, f"{len(pdf_bytes)} bytes")
    check("PDF-сигнатура на месте", pdf_bytes[:5] == b"%PDF-")
except Exception as exc:
    check("PDF сгенерирован без исключений", False, str(exc))

print("6. Пустые данные не должны падать")
try:
    empty_metrics = overview_metrics(pd.DataFrame())
    empty_payload = summary_export_payload(
        "Пустой проект", "период", "", empty_metrics, messages=None, events_agg=None
    )
    generate_summary_infographic_png(empty_payload)
    generate_summary_docx(empty_payload)
    generate_summary_pdf(empty_payload)
    check("генерация на пустых данных не падает", True)
except Exception as exc:
    check("генерация на пустых данных не падает", False, str(exc))

print("7. Конструктор отчёта: resolve_report_sections нормализует выбор")
check(
    "пустой список -> полный набор по умолчанию",
    resolve_report_sections([]) == DEFAULT_REPORT_SECTIONS,
    str(resolve_report_sections([])),
)
check("None -> полный набор по умолчанию", resolve_report_sections(None) == DEFAULT_REPORT_SECTIONS)
check(
    "неизвестные id отбрасываются, известные остаются в своём порядке",
    resolve_report_sections(["metrics", "unknown_id", "sentiment"]) == ["metrics", "sentiment"],
    str(resolve_report_sections(["metrics", "unknown_id", "sentiment"])),
)
check(
    "только неизвестные id -> откат на полный набор (не пустой отчёт по ошибке вызова)",
    resolve_report_sections(["bogus"]) == DEFAULT_REPORT_SECTIONS,
)

print("8. Конструктор отчёта: DOCX реально включает/выключает разделы")


def _docx_headings(docx_bytes):
    doc = Document(BytesIO(docx_bytes))
    return [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]


full_docx = generate_summary_docx(
    summary_export_payload(
        "ТЕХНОНИКОЛЬ",
        "24.04.2026–30.04.2026",
        "Текст саммари.",
        metrics,
        messages=MESSAGES,
        events_agg=EVENTS_AGG,
        report_template="full",
    )
)
full_headings = _docx_headings(full_docx)
check(
    "с полным набором разделов все заголовки на месте",
    {"Основные метрики", "Что включить в отчет", "Саммари периода"} <= set(full_headings),
    str(full_headings),
)

summary_only_docx = generate_summary_docx(
    summary_export_payload(
        "ТЕХНОНИКОЛЬ",
        "24.04.2026–30.04.2026",
        "Текст саммари.",
        metrics,
        messages=MESSAGES,
        events_agg=EVENTS_AGG,
        report_template="full",
        sections=["summary_text"],
    )
)
summary_only_headings = _docx_headings(summary_only_docx)
check(
    "только «summary_text» — нет заголовков метрик/тегов, саммари есть",
    "Основные метрики" not in summary_only_headings
    and "Что включить в отчет" not in summary_only_headings
    and "Саммари периода" in summary_only_headings,
    str(summary_only_headings),
)

metrics_only_docx = generate_summary_docx(
    summary_export_payload(
        "ТЕХНОНИКОЛЬ",
        "24.04.2026–30.04.2026",
        "Текст саммари.",
        metrics,
        messages=MESSAGES,
        events_agg=EVENTS_AGG,
        report_template="full",
        sections=["metrics"],
    )
)
metrics_only_headings = _docx_headings(metrics_only_docx)
check(
    "только «metrics» — заголовок метрик есть, саммари нет",
    "Основные метрики" in metrics_only_headings
    and "Саммари периода" not in metrics_only_headings,
    str(metrics_only_headings),
)

print("9. Конструктор отчёта: PNG/PDF не падают ни на одном наборе разделов, включая пустой")
try:
    combos = (
        [],
        ["metrics"],
        ["top_tags"],
        ["top_events"],
        ["highlights"],
        list(REPORT_SECTION_OPTIONS.keys()),
    )
    for combo in combos:
        combo_payload = summary_export_payload(
            "ТЕХНОНИКОЛЬ",
            "24.04.2026–30.04.2026",
            "Текст.",
            metrics,
            messages=MESSAGES,
            events_agg=EVENTS_AGG,
            sections=combo,
        )
        png = generate_summary_infographic_png(combo_payload)
        check(
            f"PNG для набора {combo or '[] (откат на полный)'} — валидная сигнатура",
            png[:8] == b"\x89PNG\r\n\x1a\n",
        )
        pdf = generate_summary_pdf(combo_payload)
        check(
            f"PDF для набора {combo or '[] (откат на полный)'} — валидная сигнатура",
            pdf[:5] == b"%PDF-",
        )
except Exception as exc:
    check("генерация не падает ни на одном наборе разделов", False, str(exc))

print("10. Конструктор отчёта: PNG с меньшим набором разделов реально компактнее")
full_png = generate_summary_infographic_png(
    summary_export_payload(
        "ТЕХНОНИКОЛЬ",
        "24.04.2026–30.04.2026",
        "Текст саммари подлиннее, чтобы блок «Главное» тоже дал контент для сравнения размеров.",
        metrics,
        messages=MESSAGES,
        events_agg=EVENTS_AGG,
        sections=list(REPORT_SECTION_OPTIONS.keys()),
    )
)
metrics_only_png = generate_summary_infographic_png(
    summary_export_payload(
        "ТЕХНОНИКОЛЬ",
        "24.04.2026–30.04.2026",
        "Текст саммари подлиннее, чтобы блок «Главное» тоже дал контент для сравнения размеров.",
        metrics,
        messages=MESSAGES,
        events_agg=EVENTS_AGG,
        sections=["metrics"],
    )
)
check(
    "PNG с одним разделом реально другой (по размеру), а не тот же файл с игнорируемым параметром",
    len(metrics_only_png) != len(full_png),
    f"{len(metrics_only_png)} vs {len(full_png)}",
)

print("11. Текстовые страницы Word/PDF окрашены под брендирование, а не голый чёрный текст")
styled_payload = summary_export_payload(
    "ТЕХНОНИКОЛЬ",
    "24.04.2026–30.04.2026",
    "Текст саммари.",
    metrics,
    messages=MESSAGES,
    events_agg=EVENTS_AGG,
    branding={**BRANDING, "accent_color": "#7c3aed"},
)
styled_docx = Document(BytesIO(generate_summary_docx(styled_payload)))
check(
    "заголовок отчёта покрашен в акцентный цвет проекта",
    str(styled_docx.styles["Heading 2"].font.color.rgb) == "7C3AED",
    str(styled_docx.styles["Heading 2"].font.color.rgb),
)
check(
    "обычный текст — не чистый чёрный, а «чернильный» тон инфографики",
    str(styled_docx.styles["Normal"].font.color.rgb) == "111827",
    str(styled_docx.styles["Normal"].font.color.rgb),
)
title_run = next(p for p in styled_docx.paragraphs if p.text == "Дайджест упоминаний").runs[0]
check(
    "заголовок отчёта (титул) тоже акцентный, а не дефолтный чёрный",
    str(title_run.font.color.rgb) == "7C3AED",
    str(title_run.font.color.rgb),
)

# PDF: пиксели не проверяем (как и раньше в этом файле - только сигнатура),
# но разный accent_color обязан давать разные байты - иначе цвет из payload
# в PDF просто не доехал.
pdf_purple = generate_summary_pdf(
    summary_export_payload(
        "ТЕХНОНИКОЛЬ",
        "24.04.2026–30.04.2026",
        "Текст саммари.",
        metrics,
        messages=MESSAGES,
        events_agg=EVENTS_AGG,
        branding={**BRANDING, "accent_color": "#7c3aed"},
        sections=["summary_text"],
    )
)
pdf_green = generate_summary_pdf(
    summary_export_payload(
        "ТЕХНОНИКОЛЬ",
        "24.04.2026–30.04.2026",
        "Текст саммари.",
        metrics,
        messages=MESSAGES,
        events_agg=EVENTS_AGG,
        branding={**BRANDING, "accent_color": "#16a34a"},
        sections=["summary_text"],
    )
)
check(
    "разный accent_color -> разные байты PDF (цвет реально доезжает до документа)",
    pdf_purple != pdf_green,
)

print()
if failures:
    print(f"ПРОВАЛЕНО: {len(failures)} → {failures}")
    raise SystemExit(1)
print("Проверки выгрузок саммари пройдены.")
