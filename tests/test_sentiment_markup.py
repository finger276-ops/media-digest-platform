# -*- coding: utf-8 -*-
"""Тональность без разметки: прочерк вместо ложного нуля — во всём продукте.

Платформа сама тональность не определяет, а берёт её из выгрузки. Если
колонка «Тональность» пуста у всех сообщений, прежний подсчёт относил их к
нейтральным, и экран, файл отчёта, саммари и карточка для ИИ говорили
«Нейтрал 100 %, негатив 0 %, риск негатива низкий» — ложное «всё спокойно».

Проверяется: общий признак разметки (services.metrics_compute), законный ноль
(всё размечено «нейтральная»), частичная разметка (считается как раньше),
выгрузка только с флагом негатива, признак на уровне выгрузки-периода,
сравнение периодов, тексты саммари и карточки ИИ, файлы Word/PDF/PNG и
сквозной прогон приложения на проекте без разметки.

Мутационные проверки (что ломает какой тест):
- has_sentiment_markup всегда True -> краснеют сценарии A/F и все «—» на экране;
- убрать флаг негатива из _row_markup -> краснеет «только флаг → размечено»;
- убрать 'nan' из EMPTY_SENTIMENT_VALUES -> краснеет сценарий F;
- в prepare_dashboard_messages считать признак по строке, а не по периоду ->
  краснеет «день из пустых строк внутри размеченной выгрузки»;
- в _tone_cards убрать ветку «—» -> краснеет сквозной «Обзор»;
- вернуть в summary_ui прежнюю строку негатива -> краснеет «саммари»;
- вернуть в report_export donut без проверки -> краснеет «PNG/PDF».
"""

import json
import os
import sys
from io import BytesIO
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
for _p in (REPO / "src", REPO / "tests"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

os.environ.setdefault("SUPABASE_URL", "https://test.supabase.co")
os.environ.setdefault("SUPABASE_SERVICE_ROLE_KEY", "test-key")
os.environ.setdefault("PLATFORM_ADMIN_PASSWORD", "test-admin")

import pandas as pd  # noqa: E402

from services import brand_metrics  # noqa: E402
from services import metrics_compute as mc  # noqa: E402

failures = []


def check(label, condition, detail=""):
    print(("  ✓ " if condition else "  ✗ ") + label + (f" — {detail}" if detail and not condition else ""))
    if not condition:
        failures.append(label)


def frame(sentiments, period="p1", flags=None, days=None, **extra):
    rows = []
    for i, sentiment in enumerate(sentiments):
        row = {
            "message_id": f"{period}_{i}",
            "period_id": period,
            "sentiment": sentiment,
            "views": 10,
            "audience": 100,
            "engagement": 1,
            "chat_profile": f"c{i}",
            "author": f"a{i}",
            "chat_title": f"Канал {i % 3}",
            "text_clean": f"Сообщение {i}",
            "datetime": days[i] if days else "2026-04-24T10:00:00",
        }
        if flags is not None:
            row["is_negative"] = flags[i]
        for key, value in extra.items():
            row[key] = value[i] if isinstance(value, list) else value
        rows.append(row)
    return pd.DataFrame(rows)


SCENARIOS = {
    "A пусто": (frame([""] * 10), False),
    "F nan/None/NULL/пробел": (frame([None, float("nan"), "nan", "None", "NULL", " "]), False),
    "B всё «нейтральная»": (frame(["нейтральная"] * 10), True),
    "C только флаг негатива": (frame([""] * 6, flags=[True, True, False, False, False, False]), True),
    "C' флаг везде False": (frame([""] * 6, flags=[False] * 6), False),
    "E частичная": (frame(["негативная", "негативная", "позитивная"] + [""] * 7), True),
    "EN английская и «Отрицательная»": (frame(["negative", "Отрицательная", "neutral"]), True),
}

print("1. Признак разметки на сырых и подготовленных кадрах")
for name, (data, expected) in SCENARIOS.items():
    for variant, df in (("сырой", data), ("подготовленный", mc.prepare_dashboard_messages(data))):
        got = mc.has_sentiment_markup(df)
        check(f"{name} / {variant}: разметка {'есть' if expected else 'нет'}", got == expected, str(got))
        check(f"{name} / {variant}: has_markup в sentiment_counts", mc.sentiment_counts(df)["has_markup"] == expected)
# Прежние числа не меняются: форма словаря только расширена.
counts = {name: mc.sentiment_counts(mc.prepare_dashboard_messages(df)) for name, (df, _) in SCENARIOS.items()}
check("A: прежние числа (всё в нейтрале)",
      [counts["A пусто"][k] for k in ("positive", "neutral", "negative", "total")] == [0, 10, 0, 10],
      str(counts["A пусто"]))
check("E: частичная считается как раньше",
      [counts["E частичная"][k] for k in ("positive", "neutral", "negative", "total")] == [1, 7, 2, 10],
      str(counts["E частичная"]))
check("C: флаг негатива учтён",
      counts["C только флаг негатива"]["negative"] == 2, str(counts["C только флаг негатива"]))
check("индексы бренда пользуются тем же признаком",
      brand_metrics.has_sentiment_markup is mc.has_sentiment_markup
      and brand_metrics.NO_SENTIMENT_REASON == mc.NO_SENTIMENT_REASON)

print("2. Пустой период — «нет данных», а не «нет разметки»")
empty = mc.sentiment_counts(pd.DataFrame())
check("пустой кадр: нули и has_markup False", empty == {"positive": 0, "neutral": 0, "negative": 0, "total": 0, "has_markup": False}, str(empty))
check("пустой кадр: прочерка «нет разметки» нет", mc.sentiment_unmarked(empty) is False)

print("3. Признак определяется по выгрузке-периоду, срез его наследует")
days = ["2026-04-24T10:00:00"] * 3 + ["2026-04-25T10:00:00"] * 3
p1 = frame(["негативная", "позитивная", "", "", "", ""], "p1", days=days)
p2 = frame([""] * 6, "p2", days=[d.replace("04-2", "05-0") for d in days])
both = mc.prepare_dashboard_messages(pd.concat([p1, p2], ignore_index=True))
day25 = both[(both["period_id"] == "p1") & both["datetime"].str.startswith("2026-04-25")]
check("день из одних пустых строк внутри размеченной выгрузки — законный ноль, не прочерк",
      mc.has_sentiment_markup(day25))
check("неразмеченный период целиком — прочерк", not mc.has_sentiment_markup(both[both["period_id"] == "p2"]))
check("размеченный + неразмеченный вместе — как раньше", mc.has_sentiment_markup(both))
stale = pd.concat(
    [mc.prepare_dashboard_messages(frame([""] * 3, "p1")), frame([""] * 2, "p9")], ignore_index=True
)
check("подготовленный кадр + сырые пустые строки (NaN в признаке) — разметки нет", not mc.has_sentiment_markup(stale))
stale_marked = pd.concat(
    [mc.prepare_dashboard_messages(frame([""] * 3, "p1")), frame(["негативная"], "p9")], ignore_index=True
)
check("… а с сырой размеченной строкой — есть", mc.has_sentiment_markup(stale_marked))

from services.period_comparison import (  # noqa: E402
    comparison_row,
    comparison_visual_rows,
    daily_metrics_for_comparison,
)

daily = daily_metrics_for_comparison(both)
check("дни размеченной выгрузки размечены, дни неразмеченной — нет",
      [d["sentiment"]["has_markup"] for d in daily] == [True, True, False, False],
      str([(d["label"], d["sentiment"]["has_markup"]) for d in daily]))

print("4. Старые словари без признака")
check("словарь без ключа и без сообщений — считается размеченным", mc.sentiment_unmarked({"total": 5}) is False)
check("словарь без ключа + сообщения без разметки — прочерк",
      mc.sentiment_unmarked({"total": 10}, SCENARIOS["A пусто"][0]) is True)
check("has_markup False — прочерк", mc.sentiment_unmarked({"total": 3, "has_markup": False}) is True)

print("5. Сравнение периодов")
unmarked_period = mc.overview_metrics(mc.prepare_dashboard_messages(frame([""] * 4, "p1")))
marked_period = mc.overview_metrics(mc.prepare_dashboard_messages(frame(["негативная", "", "", ""], "p2")))
neutral_period = mc.overview_metrics(mc.prepare_dashboard_messages(frame(["нейтральная"] * 4, "p3")))
for item in (unmarked_period, marked_period, neutral_period):
    total = item["sentiment"]["total"]
    for key in ("positive", "neutral", "negative"):
        item[f"{key}_share"] = item["sentiment"][key] / total
    item["label"] = "период"
row_unmarked = comparison_row(unmarked_period, None)
check("неразмеченный период: доли — прочерк",
      [row_unmarked[k] for k in ("Позитив", "Нейтрал", "Негатив")] == ["—", "—", "—"], str(row_unmarked))
row_after = comparison_row(marked_period, unmarked_period)
check("размеченный после неразмеченного: доли есть, изменения — прочерк",
      row_after["Негатив"] == "25%" and row_after["Δ негатива"] == "—", str(row_after))
row_neutral = comparison_row(neutral_period, neutral_period)
check("всё нейтральное: законный 0 % и изменение 0,0 п.п.",
      row_neutral["Негатив"] == "0%" and row_neutral["Δ негатива"] == "0,0 п.п.", str(row_neutral))
visual = comparison_visual_rows([unmarked_period, marked_period])
check("данные графиков помечают неразмеченные точки",
      visual["Тональность размечена"].tolist() == [False, True], str(visual.to_dict("records")))

print("6. Тексты: саммари, клиентский обзор, карточка для ИИ")
from client_insights_ui import build_tag_change_table  # noqa: E402
from services.ai_summary import KIND_RISKS, TASK_PROMPTS, build_data_card, comparison_block  # noqa: E402
from summary_ui import build_auto_summary  # noqa: E402

no_markup = mc.prepare_dashboard_messages(frame([""] * 10, tags=["Кровля"] * 10))
neutral = mc.prepare_dashboard_messages(frame(["нейтральная"] * 10, tags=["Кровля"] * 10))
periods = pd.DataFrame([{"period_id": "p1", "period_name": "Апрель"}, {"period_id": "p2", "period_name": "Май"}])

summary_none = build_auto_summary(no_markup, pd.DataFrame(), periods, ["p1"])
check("саммари без разметки не пишет «Негативных сообщений: 0»", "Негативных сообщений: 0" not in summary_none, summary_none[:300])
check("саммари без разметки называет причину", mc.NO_SENTIMENT_REASON in summary_none)
check("риск негатива «не оценён», а не «низкий»",
      "Риск негатива: не оценён" in summary_none and "Риск негатива: низкий" not in summary_none, summary_none)
check("блок метрик без «нейтрал 10 (100,0%)»", "нейтрал 10" not in summary_none)
summary_neutral = build_auto_summary(neutral, pd.DataFrame(), periods, ["p1"])
check("всё нейтральное: законный ноль остаётся",
      "Негативных сообщений: 0 (0.0%)." in summary_neutral and "Риск негатива: низкий" in summary_neutral,
      summary_neutral[:300])

card_args = dict(project_name="Проект", periods=periods, period_ids=["p1"], events_agg=pd.DataFrame(),
                 metrics=None, brand_cards=None, include_excerpts=True)
card_none = build_data_card(messages=no_markup, negative_excerpts=True, **card_args)
check("карточка ИИ: «Тональность: нет данных» вместо трёх чисел",
      "Тональность: нет данных" in card_none and "нейтрал 10" not in card_none, card_none[:400])
check("карточка ИИ: у тегов нет «негатив 0»", ", негатив 0" not in card_none)
check("карточка ИИ для рисков: не пишет «Негативных сообщений в периоде нет»",
      "Негативных сообщений в периоде нет" not in card_none and "выделить нельзя" in card_none)
card_flag = build_data_card(messages=mc.prepare_dashboard_messages(SCENARIOS["C только флаг негатива"][0]),
                            negative_excerpts=True, **card_args)
check("карточка ИИ: при одном флаге негатива выдержки есть", "Выдержки негативных сообщений (" in card_flag, card_flag[-400:])
check("задание для рисков запрещает вывод «негатива нет» без разметки", "не размечена" in TASK_PROMPTS[KIND_RISKS])

mixed = {"comparison": {"previous": dict(unmarked_period, label="Апрель"),
                        "current": dict(marked_period, label="Май")}}
dynamics = comparison_block(mixed)
check("динамика: нет «было 0, стало 1» между неразмеченным и размеченным периодом",
      "негативные сообщения: было" not in dynamics and "нет данных для сравнения" in dynamics, dynamics)

tag_frame = mc.prepare_dashboard_messages(pd.concat(
    [frame([""] * 4, "p1", tags="Кровля"), frame(["негативная", "", "", ""], "p2", tags="Кровля")],
    ignore_index=True,
))
tag_changes = build_tag_change_table(tag_frame, periods, ["p1", "p2"])
check("изменения тегов: без разметки в прошлом периоде Δ негатива — прочерк",
      not tag_changes.empty and set(tag_changes["Δ негатива"]) == {"—"}, str(tag_changes.to_dict("records")))
# Пустой прошлый период — «нет данных», а не «нет разметки»: сравнение как раньше.
only_current = mc.prepare_dashboard_messages(frame(["негативная", "негативная", "", ""], "p2", tags="Кровля"))
empty_prev_changes = build_tag_change_table(only_current, periods, ["p1", "p2"])
check("изменения тегов: пустой прошлый период не прячет Δ негатива",
      not empty_prev_changes.empty and list(empty_prev_changes["Δ негатива"]) == [2],
      str(empty_prev_changes.to_dict("records")))

from overview_ui import _tone_cards  # noqa: E402

marked_sent = mc.sentiment_counts(mc.prepare_dashboard_messages(frame(["негативная", "", "", ""], "p2")))
empty_sent = mc.sentiment_counts(pd.DataFrame())
check("«Сравнение периодов»: изменение к пустой точке показывается, как раньше",
      _tone_cards(marked_sent, empty_sent, require_prev_messages=False)[2]["delta"] == "+25,0 п.п.",
      str(_tone_cards(marked_sent, empty_sent, require_prev_messages=False)[2]))
check("шапка «Обзора»: к пустому прошлому периоду изменения нет, как раньше",
      _tone_cards(marked_sent, empty_sent)[2]["delta"] is None)
check("изменения нет, если прошлый период не размечен",
      _tone_cards(marked_sent, mc.sentiment_counts(no_markup))[2]["delta"] is None)

print("7. Файлы отчёта: Word, PDF, PNG")
import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402

from services import report_export  # noqa: E402

payload_none = report_export.summary_export_payload("Проект", "Апрель", "Текст.", mc.overview_metrics(no_markup),
                                                    messages=no_markup, events_agg=pd.DataFrame())
payload_neutral = report_export.summary_export_payload("Проект", "Апрель", "Текст.", mc.overview_metrics(neutral),
                                                       messages=neutral, events_agg=pd.DataFrame())
check("payload помечает отсутствие разметки",
      payload_none["sentiment_markup"] is False and payload_neutral["sentiment_markup"] is True)
docx_text = "\n".join(p.text for p in Document(BytesIO(report_export.generate_summary_docx(payload_none))).paragraphs)
check("Word без разметки: нет «нейтрал — 100%», есть причина",
      "нейтрал — 100%" not in docx_text and mc.NO_SENTIMENT_REASON in docx_text, docx_text[:400])
docx_neutral = "\n".join(p.text for p in Document(BytesIO(report_export.generate_summary_docx(payload_neutral))).paragraphs)
check("Word при законном нуле: доли на месте", "нейтрал — 100%" in docx_neutral)


def sentiment_axes(payload):
    fig = plt.figure(figsize=(8.27, 11.69))
    ax = fig.add_axes([0, 0, 1, 1])
    before = len(fig.axes)
    report_export._draw_sentiment_section(ax, fig, payload, [], 0.6)
    added = len(fig.axes) - before
    plt.close(fig)
    return added


check("PNG без разметки: кольцевая диаграмма не рисуется", sentiment_axes(payload_none) == 0)
check("PNG при законном нуле: диаграмма есть", sentiment_axes(payload_neutral) == 1)

created = []
original_block = report_export._PdfSentimentBlock


class _SpyBlock(original_block):
    def __init__(self, *args, **kwargs):
        created.append(args[0])
        super().__init__(*args, **kwargs)


report_export._PdfSentimentBlock = _SpyBlock
try:
    report_export.generate_summary_pdf(payload_none)
    check("PDF без разметки: блок-диаграмма не создан", created == [], str(created))
    report_export.generate_summary_pdf(payload_neutral)
    check("PDF при законном нуле: диаграмма есть", len(created) == 1, str(created))
finally:
    report_export._PdfSentimentBlock = original_block

print("8. Негатив по флагу и английской разметке в тегах и инфоповодах")
from services.manual_moderation import recompute_event_counts  # noqa: E402
from services.tag_compute import build_tag_statistics_compute  # noqa: E402

flag_tags = build_tag_statistics_compute(frame([""] * 6, flags=[True, True, False, False, False, False], tags="Кровля"))
check("тег: негатив по флагу is_negative", int(flag_tags.loc[0, "Негатив"]) == 2, str(flag_tags.to_dict("records")))
en_tags = build_tag_statistics_compute(frame(["negative", "Отрицательная", "neutral"], tags="Кровля"))
check("тег: негатив по «negative»/«Отрицательная»", int(en_tags.loc[0, "Негатив"]) == 2, str(en_tags.to_dict("records")))
events = pd.DataFrame([{"event_id": "e1", "event_title": "Тема", "message_count": 0, "negative_count": 0}])
event_messages = frame([""] * 3, flags=[True, True, True], event_id="e1")
recounted = recompute_event_counts(events, event_messages)
check("инфоповод: счётчик негатива по флагу", int(recounted.loc[0, "negative_count"]) == 3, str(recounted.to_dict("records")))

print("9. Сквозной прогон приложения на проекте без разметки тональности")
from datetime import datetime, timezone  # noqa: E402

from fake_supabase import FakeClient  # noqa: E402

import platform_store as store  # noqa: E402

CLIENT = FakeClient()
store.get_supabase_client = lambda: CLIENT
now = datetime.now(timezone.utc).isoformat()
PROJECT = "nomark"
CLIENT.db["platform_projects"] = [
    {"project_id": PROJECT, "project_name": "Без разметки", "status": "active",
     "settings": {}, "created_at": now, "updated_at": now}
]
CLIENT.db["platform_periods"] = [
    {"project_id": PROJECT, "period_id": pid, "period_name": name, "date_from": start, "date_to": end,
     "source_filename": "f.xlsx", "status": "active", "manifest": {}, "uploaded_at": now}
    for pid, name, start, end in (("q1", "Апрель", "2026-04-24", "2026-04-30"),
                                  ("q2", "Май", "2026-05-01", "2026-05-07"))
]
THEMES = ["Запуск завода", "Жалобы на монтаж", "Отраслевая статистика"]


def message_row(pid, index, theme, day, **extra):
    payload = {
        "message_id": f"{pid}_m{index}", "period_id": pid, "date": day, "datetime": f"{day}T10:00:00",
        "sentiment": "", "views": 1000 * (index + 1), "audience": 5000, "engagement": 10, "likes": 10,
        "comments": 0, "reposts": 0, "text_clean": f"Сообщение {index} про {theme}",
        "message_link": f"https://example.com/{pid}/{index}", "platform": "vk.com",
        "author": f"user{index}", "tags": theme, "event_title": theme,
    }
    payload.update(extra)
    return {"project_id": PROJECT, "period_id": pid, "table_name": "messages",
            "row_id": payload["message_id"], "payload": payload}


def event_row(pid, index, theme, start, end):
    return {"project_id": PROJECT, "period_id": pid, "table_name": "events", "row_id": f"{pid}_e{index}",
            "payload": {"event_id": f"{pid}_e{index}", "period_id": pid, "event_title": theme,
                        "event_summary": f"Инфоповод про {theme}", "message_count": 3, "negative_count": 0,
                        "chat_count": 2, "importance_score": 10 - index, "start_date": start,
                        "end_date": end, "main_tags": theme}}


rows = []
for pid, day, end in (("q1", "2026-04-24", "2026-04-30"), ("q2", "2026-05-01", "2026-05-07")):
    rows += [message_row(pid, i, THEMES[i % 3], day) for i in range(9)]
    rows += [event_row(pid, i, theme, day, end) for i, theme in enumerate(THEMES)]
# Отзывы без оценки и без тональности: отбирать претензии не по чему.
rows += [
    message_row("q1", 100 + i, "Кровля", "2026-04-25", platform="wildberries.ru", platform_type="Отзывы",
                message_type="Комментарий", title="ТН / Гибкая черепица", event_title="",
                text_clean=f"Отзыв {i}: Недостатки: пришло в рваном пакете")
    for i in range(3)
]
CLIENT.db["platform_table_rows"] = rows

from streamlit.testing.v1 import AppTest  # noqa: E402

at = AppTest.from_file(str(REPO / "streamlit_app.py"), default_timeout=120)
at.session_state["platform_is_admin"] = True
at.run()
check("приложение стартовало", not at.exception, str(at.exception))
period_select = [m for m in at.sidebar.multiselect if str(m.label) == "Периоды"]
if period_select:
    period_select[0].set_value(["q1", "q2"]).run()


def texts():
    out = []
    for kind in ("markdown", "caption", "info", "success", "warning"):
        out += [str(el.value) for el in getattr(at, kind)]
    return out


def metric_values(label):
    return [str(m.value) for m in at.metric if str(m.label) == label]


def open_section(name):
    {str(b.label): b for b in at.sidebar.button}[name].click().run()


check("«Обзор»: карточки тональности — прочерк",
      metric_values("Негатив")[:1] == ["—"] and metric_values("Нейтрал")[:1] == ["—"],
      str([(str(m.label), str(m.value)) for m in at.metric]))
check("«Обзор»: нет «Нейтрал 100%»", "100%" not in metric_values("Нейтрал"), str(metric_values("Нейтрал")))
page = " ".join(texts())
check("клиентский обзор: нет «Критичный негатив не выявлен»", "Критичный негатив не выявлен" not in page)
check("клиентский обзор: сигнал «Тональность не размечена»", "Тональность не размечена" in page)
check("клиентский обзор: «Риск негатива» — прочерк", metric_values("Риск негатива")[:1] == ["—"],
      str(metric_values("Риск негатива")))
change_tables = [el.value for el in at.dataframe if "Δ негатива" in getattr(el.value, "columns", [])]
check("таблица изменений тегов на экране: прочерк не превращается обратно в «0»",
      bool(change_tables)
      and set(change_tables[0]["Δ негатива"]) == {"—"}
      and set(change_tables[0]["Негатив сейчас"]) == {"—"},
      str(change_tables[0].head() if change_tables else "нет таблицы"))

open_section("Теги")
check("«Теги» открылись", not at.exception, str(at.exception))
page = " ".join(texts())
check("строка метрик: «тональность не размечена», не «негатив 0%»",
      mc.NO_SENTIMENT_LABEL in page and "негатив 0%" not in page, page[:300])
tag_tables = [el.value for el in at.dataframe if "Доля негатива" in getattr(el.value, "columns", [])]
check("таблица тегов: негатив — прочерк",
      bool(tag_tables) and set(tag_tables[0]["Доля негатива"]) == {"—"},
      str(tag_tables[0].head() if tag_tables else "нет таблицы"))

open_section("Инфоповоды")
check("«Инфоповоды» открылись", not at.exception, str(at.exception))
check("инфоповоды: подпись о причине", any("Доля негатива по инфоповодам не показана" in t for t in texts()))
specs = [json.loads(el.proto.spec) for el in at.get("vega_lite_chart") if getattr(el, "proto", None) is not None]
check("график топ-инфоповодов без шкалы «Доля негатива»",
      not any(s.get("encoding", {}).get("color", {}).get("field") == "Доля негатива" for s in specs))

open_section("Отзывы")
check("«Отзывы» открылись", not at.exception, str(at.exception))
page = " ".join(texts())
check("отзывы: нет «ни одной претензии»", "ни одной претензии" not in page, page[:300])
check("отзывы: сказано, что отбирать не по чему", "Претензии не отобрать" in page)
check("отзывы: «Претензий» — прочерк", metric_values("Претензий")[:1] == ["—"], str(metric_values("Претензий")))

open_section("Отчёт")
check("«Отчёт» открылся", not at.exception, str(at.exception))
page = " ".join(texts() + [str(el.value) for el in at.text_area])
check("отчёт: предупреждение о блоке «Тональность»", "будет пометка вместо диаграммы" in page)
check("отчёт: автосаммари не утверждает «Риск негатива: низкий»", "Риск негатива: низкий" not in page)

print()
if failures:
    print(f"ПРОВАЛЕНО: {len(failures)} → {failures}")
    raise SystemExit(1)
print("Проверки тональности без разметки пройдены.")
