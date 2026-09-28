# -*- coding: utf-8 -*-
"""Отобранные сообщения — в Excel из раздела «Сообщения».

Мутационные проверки (что ломает какой тест):
- выгружается страница, а не весь отбор -> «вся лента: в файл — все
  найденные, не страница» краснеет;
- в ключевых сообщениях выгружаются 15 из топа -> «ключевые: в файл — весь
  отбор, не топ-15» краснеет;
- метрика без проверки выгрузки периода -> «нет охвата в выгрузке — прочерк,
  а не 0» краснеет;
- «Без сюжета» не отсекается -> «служебный «Без сюжета» — пустой инфоповод»
  краснеет;
- ссылка без гиперссылки -> «ссылка кликабельна» краснеет;
- описание отбора без поиска -> «отбор записан на втором листе» краснеет;
- свежие не сверху -> «свежие сообщения сверху» краснеет;
- имя файла с кириллицей -> «имя файла — латиницей, без запрещённых
  символов» краснеет.
"""

import sys
from io import BytesIO
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
for _p in (REPO / "src", REPO / "tests"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from services.message_export import (  # noqa: E402
    describe_filters,
    messages_export_frame,
    messages_to_xlsx,
    safe_messages_filename,
)

failures = []


def check(label, condition, detail=""):
    print(("  ✓ " if condition else "  ✗ ") + label + (f" — {detail}" if detail and not condition else ""))
    if not condition:
        failures.append(label)


print("1. Таблица для Excel")
SAMPLE = pd.DataFrame({
    "datetime": ["2026-04-02T10:00:00", "2026-04-03T11:00:00", "2026-04-01T09:00:00"],
    "platform": ["vk.com", "t.me", ""],
    "message_link": ["https://vk.com/wall-1_1", "", "https://otzovik.com/r/1"],
    "chat_title": ["Кровельщики", "Стройка", ""],
    "author": ["a", "b", "c"],
    "message_type": ["Пост", "", "Комментарий"],
    "sentiment": ["негатив", "нейтрал", "позитив"],
    "tags": ["Технониколь|Кровля", "", "Кнауф"],
    "event_title": ["Протечки кровли", "Без сюжета", ""],
    "text_clean": ["первый", "второй", "х" * 40000],
    "audience": [100, 200, 300], "views": [10, 0, 5], "engagement": [1, 2, 3],
    "_period_has_reach": [True, False, True],
})
frame = messages_export_frame(SAMPLE)
check("колонки по-русски", list(frame.columns) == ["Дата", "Площадка", "Сообщество / канал", "Автор", "Тип сообщения",
                                                  "Тональность", "Теги", "Инфоповод", "Текст", "Ссылка", "Аудитория",
                                                  "Охват", "Вовлеченность"], str(list(frame.columns)))
check("свежие сообщения сверху", list(frame["Автор"]) == ["b", "a", "c"], str(list(frame["Автор"])))
by_author = {row["Автор"]: row for row in frame.to_dict("records")}
check("площадка — домен, в том числе из ссылки", by_author["b"]["Площадка"] == "telegram.org"
      and by_author["c"]["Площадка"] == "otzovik.com", str(frame["Площадка"].tolist()))
check("нет охвата в выгрузке — прочерк, а не 0", by_author["b"]["Охват"] == "—" and by_author["a"]["Охват"] == 10,
      str(frame["Охват"].tolist()))
check("служебный «Без сюжета» — пустой инфоповод", by_author["b"]["Инфоповод"] == ""
      and by_author["a"]["Инфоповод"] == "Протечки кровли")
check("теги через запятую", by_author["a"]["Теги"] == "Технониколь, Кровля")
check("длинный текст обрезан под предел ячейки Excel", len(by_author["c"]["Текст"]) == 32000)
# Кириллическое имя браузер сохранял как «download» без расширения.
check("имя файла — латиницей, без запрещённых символов", safe_messages_filename("Кнауф: кровля", "01.04–07.04") ==
      "messages_Knauf_krovlya_01.04_07.04.xlsx", safe_messages_filename("Кнауф: кровля", "01.04–07.04"))

print("2. Файл Excel")
filters = describe_filters(tags=["Технониколь", "Кровля"], match_all=True, types=["Пост"], search="протека",
                           slice_tags=["Технониколь"], event_title="Протечки кровли")
book = load_workbook(BytesIO(messages_to_xlsx(SAMPLE, project_name="Кнауф", period_label="Апрель", filters=filters)))
check("два листа", book.sheetnames == ["Сообщения", "Параметры"], str(book.sheetnames))
sheet = book["Сообщения"]
header = [cell.value for cell in sheet[1]]
link_cell = sheet.cell(row=3, column=header.index("Ссылка") + 1)
check("ссылка кликабельна", link_cell.hyperlink is not None and link_cell.hyperlink.target == "https://vk.com/wall-1_1",
      str(link_cell.value))
check("дата — дата Excel, а не текст", sheet.cell(row=2, column=1).is_date)
check("шапка закреплена", sheet.freeze_panes == "A2")
params = {row[0].value: row[1].value for row in book["Параметры"].iter_rows(min_row=2)}
check("отбор записан на втором листе", params.get("Отбор") == "срез по тегам: Технониколь; инфоповод: Протечки кровли; "
      "теги: Технониколь и Кровля; тип: Пост; поиск: «протека»", str(params.get("Отбор")))
check("проект, период, число сообщений", params.get("Проект") == "Кнауф" and params.get("Период") == "Апрель"
      and params.get("Сообщений") == 3, str(params))
check("пустой отбор — «все сообщения»", describe_filters() == [])

print("3. Кнопка в разделе «Сообщения»")
from streamlit.testing.v1 import AppTest  # noqa: E402


def export_app():
    import pandas as pd
    import streamlit as st

    import messages_ui

    if not getattr(messages_ui, "_spy_installed", False):
        original = messages_ui._render_excel_button

        def spy(export_set, **kwargs):
            st.session_state["exported"] = (len(export_set), list(kwargs["filters"]))
            original(export_set, **kwargs)

        messages_ui._render_excel_button = spy
        messages_ui._spy_installed = True
    rows = []
    for i in range(60):
        tags = "Технониколь" if i < 20 else "Кнауф"
        rows.append({"message_id": f"m{i}", "tags": tags, "message_type": "Пост" if i % 2 else "Комментарий",
                     "text_clean": ("протекает кровля " if i < 30 else "трещина ") + str(i),
                     "datetime": f"2026-04-{1 + i % 28:02d}T10:00:00", "engagement": i, "views": 10, "audience": 100,
                     "sentiment": "нейтральная", "message_link": f"https://vk.com/wall-1_{i}"})
    messages_ui.render_messages_block(pd.DataFrame(rows), project_id="p", project_name="Кнауф", period_label="Апрель",
                                      slice_tags=st.session_state.get("slice"))


at = AppTest.from_function(export_app, default_timeout=60)
at.run()
check("раздел открылся", not at.exception, str(at.exception))
buttons = [el for el in at.get("download_button")]
check("кнопка «Скачать в Excel» на месте", any("Скачать в Excel · 60 сообщ." in str(el.proto.label) for el in buttons),
      str([el.proto.label for el in buttons]))
check("ключевые: в файл — весь отбор, не топ-15", at.session_state["exported"][0] == 60, str(at.session_state["exported"]))
next(m for m in at.multiselect if str(m.label) == "Теги").select("Технониколь").run()
check("выбран тег — в файл только его сообщения", at.session_state["exported"] == (20, ["тег: Технониколь"]),
      str(at.session_state["exported"]))
next(m for m in at.multiselect if str(m.label) == "Теги").set_value([]).run()
next(r for r in at.radio if str(r.label) == "Режим просмотра сообщений").set_value("Вся лента").run()
next(s for s in at.selectbox if str(s.label) == "Сообщений на странице").set_value(25).run()
next(t for t in at.text_input if str(t.label) == "Поиск по всей ленте").input("протекает").run()
# Найдено 30, на странице 25: в файл идут все 30.
check("вся лента: в файл — все найденные, не страница", at.session_state["exported"][0] == 30,
      str(at.session_state["exported"]))
next(m for m in at.multiselect if str(m.label) == "Теги").select("Технониколь").run()
check("поиск записан в отбор", at.session_state["exported"] == (20, ["тег: Технониколь", "поиск: «протекает»"]),
      str(at.session_state["exported"]))
at.session_state["slice"] = ["Технониколь"]
at.run()
check("срез по тегам записан в отбор", at.session_state["exported"][1][0] == "срез по тегам: Технониколь",
      str(at.session_state["exported"]))

print()
if failures:
    print(f"ПРОВАЛЕНО: {len(failures)}")
    for f in failures:
        print(f"  - {f}")
    raise SystemExit(1)
print("Выгрузка отобранных сообщений в Excel работает.")
