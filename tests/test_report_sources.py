# -*- coding: utf-8 -*-
"""Раздел отчёта «Топ площадок»: Word, PDF, PNG, PowerPoint.

Мутационные проверки (что ломает какой тест):
- площадки по названию сообщества, а не по домену -> «площадки — домены,
  самые частые первыми» краснеет;
- Word без проверки раздела -> «раздел выключен — в Word площадок нет»
  краснеет;
- PDF без площадок -> «площадки в PDF» краснеет;
- PNG: строка не под метриками -> «PNG: площадки строкой под метриками»
  краснеет;
- PNG: без отдельной строки, когда метрики выключены -> «PNG: метрики
  выключены — площадки отдельной строкой» краснеет;
- PNG: площадки отдельным блоком с заголовком -> «PNG: с площадками
  «Главное» теряет не больше строки» краснеет;
- PowerPoint без слайда -> «слайд «Топ площадок» с диаграммой» краснеет.
"""

import sys
from io import BytesIO
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
for _p in (REPO / "src", REPO / "tests"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import docx  # noqa: E402
import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402

import services.report_export as report_export  # noqa: E402
from services.dashboard_config import DEFAULT_REPORT_SECTIONS, REPORT_SECTION_OPTIONS  # noqa: E402
from services.report_pptx import generate_summary_pptx  # noqa: E402

failures = []


def check(label, condition, detail=""):
    print(("  ✓ " if condition else "  ✗ ") + label + (f" — {detail}" if detail and not condition else ""))
    if not condition:
        failures.append(label)


MESSAGES = pd.DataFrame({
    "platform": ["vk.com"] * 6 + ["t.me"] * 3 + [""],
    "chat_title": ["Барахолка Луганск"] * 6 + ["Такси чат"] * 3 + [""],
    "message_link": [""] * 9 + ["https://otzovik.com/review/1"],
    "text_clean": ["x"] * 10,
    "tags": ["Тег"] * 10,
    "sentiment": ["негатив"] * 10,
    # С типами под карточками метрик две строки: типы и площадки.
    "message_type": ["Пост"] * 7 + ["Комментарий"] * 3,
})
METRICS = {"messages": 10, "sentiment": {"total": 10, "negative": 10}}
SUMMARY = "\n".join("- Очень длинный тезис саммари, который занимает две строки в инфографике и проверяет, "
                    f"что текст не наезжает на подпись внизу листа отчёта номер {i}" for i in range(4))


def payload(sections=None, messages=MESSAGES, metrics=METRICS):
    return report_export.summary_export_payload("Кнауф", "Апрель", SUMMARY, dict(metrics), messages=messages,
                                                events_agg=pd.DataFrame({"event_title": [f"Инфоповод {i}" for i in range(6)],
                                                                         "messages": [6 - i for i in range(6)]}),
                                                sections=sections)


print("1. Данные раздела")
check("раздел в конструкторе отчёта", REPORT_SECTION_OPTIONS.get("top_sources") == "Топ площадок"
      and "top_sources" in DEFAULT_REPORT_SECTIONS)
top = payload()["top_sources"]
check("площадки — домены, самые частые первыми", [item["name"] for item in top] == ["vk.com", "telegram.org", "otzovik.com"],
      str(top))
check("доля от всех сообщений", abs(top[0]["share"] - 0.6) < 1e-9 and top[0]["messages"] == 6)
check("нет сообщений — нет площадок", payload(messages=pd.DataFrame())["top_sources"] == [])

print("2. Word и PDF")


def docx_text(data):
    return "\n".join(p.text for p in docx.Document(BytesIO(data)).paragraphs)


text = docx_text(report_export.generate_summary_docx(payload()))
check("площадки в Word", "Топ площадок: vk.com — 6 сообщ. (60%); telegram.org — 3 сообщ. (30%); "
      "otzovik.com — 1 сообщ. (10%)." in text, text[-400:])
check("в Word нет названий сообществ", "Барахолка" not in text and "Такси чат" not in text)
check("раздел выключен — в Word площадок нет",
      "Топ площадок" not in docx_text(report_export.generate_summary_docx(payload(sections=["metrics", "highlights"]))))

import reportlab.platypus as platypus  # noqa: E402

pdf_texts = []
_Paragraph = platypus.Paragraph


class _Spy(_Paragraph):
    def __init__(self, text, *args, **kwargs):
        pdf_texts.append(str(text))
        super().__init__(text, *args, **kwargs)


platypus.Paragraph = _Spy
try:
    report_export.generate_summary_pdf(payload())
finally:
    platypus.Paragraph = _Paragraph
check("площадки в PDF", "Площадки" in pdf_texts and "• vk.com — 6 сообщ. (60%)" in pdf_texts,
      str([t for t in pdf_texts if "сообщ." in t]))

print("3. PNG")
highlight_lines = {}
_highlights = report_export._draw_highlights_section


def _spy_highlights(ax, data, top):
    highlight_lines["top"] = top
    return _highlights(ax, data, top)


report_export._draw_highlights_section = _spy_highlights


def png_texts(sections=None):
    captured = []
    original_savefig = plt.Figure.savefig

    def savefig(fig, *args, **kwargs):
        captured.extend(t.get_text() for ax in fig.axes for t in ax.texts)
        return original_savefig(fig, *args, **kwargs)

    plt.Figure.savefig = savefig
    try:
        report_export.generate_summary_infographic_png(payload(sections=sections))
    finally:
        plt.Figure.savefig = original_savefig
    return captured


texts = png_texts()
top_with_sources = highlight_lines["top"]
check("PNG: площадки строкой под метриками", "Площадки: vk.com 60% · telegram.org 30% · otzovik.com 10%" in texts
      and not any(t == "Площадки" for t in texts), str([t for t in texts if "Площадки" in t]))
png_texts(sections=[s for s in DEFAULT_REPORT_SECTIONS if s != "top_sources"])
# Строка «Главного» — 0.021 высоты листа: сдвиг до 0.011 стоит не больше одной строки.
shift = highlight_lines["top"] - top_with_sources
check("PNG: с площадками «Главное» теряет не больше строки", 0 < shift <= 0.011, f"сдвиг {shift:.3f}")
texts = png_texts(sections=["top_sources", "highlights"])
check("PNG: метрики выключены — площадки отдельной строкой", any(t.startswith("Площадки: vk.com") for t in texts),
      str(texts[:12]))
check("PNG: раздел выключен — площадок нет", not any("Площадки" in t for t in png_texts(sections=["metrics"])))
report_export._draw_highlights_section = _highlights

print("4. PowerPoint")


def slide(prs, title):
    for item in prs.slides:
        if any(sh.has_text_frame and sh.text_frame.text == title for sh in item.shapes):
            return item
    return None


prs = Presentation(BytesIO(generate_summary_pptx(payload())))
sources_slide = slide(prs, "Топ площадок")
chart = next((sh.chart for sh in sources_slide.shapes if sh.has_chart), None) if sources_slide else None
# Столбцы идут снизу вверх: самая крупная площадка — последней категорией, то есть сверху.
check("слайд «Топ площадок» с диаграммой", chart is not None and chart.chart_type == XL_CHART_TYPE.BAR_CLUSTERED
      and list(chart.plots[0].categories) == ["otzovik.com", "telegram.org", "vk.com"],
      str(list(chart.plots[0].categories)) if chart is not None else "нет слайда")
check("раздел выключен — нет слайда",
      slide(Presentation(BytesIO(generate_summary_pptx(payload(sections=["metrics"])))), "Топ площадок") is None)
check("нет площадок — нет слайда",
      slide(Presentation(BytesIO(generate_summary_pptx(payload(messages=pd.DataFrame())))), "Топ площадок") is None)

print()
if failures:
    print(f"ПРОВАЛЕНО: {len(failures)}")
    for f in failures:
        print(f"  - {f}")
    raise SystemExit(1)
print("«Топ площадок» есть во всех форматах отчёта.")
