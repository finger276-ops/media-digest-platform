# -*- coding: utf-8 -*-
"""Типы сообщений в «Сравнении периодов» и в отчёте (Word, PDF, PNG, PowerPoint).

Мутационные проверки (что ломает какой тест):
- type_comparison_row считает изменение без проверки типов прошлой точки ->
  «точка после точки без типа — изменения нет» краснеет;
- type_comparison_row с обратным знаком -> «изменение доли к предыдущей
  точке» краснеет;
- comparison_type_order без «Тип не указан» в конце -> «Тип не указан — в
  конце таблицы» краснеет;
- type_share_rows берёт точки без типа -> «график — только точки с типом»
  краснеет;
- карточки сравнения без типов прошлой точки -> «карточки типов в
  сравнении — к предыдущей точке» краснеет;
- вид «Типы сообщений» без проверки наличия типов -> «нет типа — нет вида
  таблицы» краснеет;
- payload без message_types -> «строка о типах в Word» краснеет;
- строка в Word без проверки раздела -> «раздел «Основные метрики» выключен —
  строки нет» краснеет;
- PDF без строки о типах -> «строка о типах в PDF» краснеет;
- PNG без строки о типах -> «строка о типах в PNG» краснеет;
- слайд без проверки типов -> «нет типа — нет слайда» краснеет.
"""

import os
import sys
from io import BytesIO
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
for _p in (REPO / "src", REPO / "tests"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

os.environ["SUPABASE_URL"] = "https://test.supabase.co"
os.environ["SUPABASE_SERVICE_ROLE_KEY"] = "test-key"
os.environ["PLATFORM_ADMIN_PASSWORD"] = "test-admin"

import pandas as pd  # noqa: E402

from services.period_comparison import (  # noqa: E402
    build_comparison_table,
    comparison_type_order,
    type_share_rows,
)

failures = []


def check(label, condition, detail=""):
    print(("  ✓ " if condition else "  ✗ ") + label + (f" — {detail}" if detail and not condition else ""))
    if not condition:
        failures.append(label)


print("1. Сравнение периодов: таблица и данные графика")
# Май — ни одного сообщения; Июль — сообщения есть, а колонки «Тип сообщения»
# в его выгрузке нет. «Тип не указан» по числу (5) обогнал бы «Репост» (4).
SEQ = [
    {"label": "Март", "messages": 10, "message_types": [("Пост", 6), ("Комментарий", 4)]},
    {"label": "Апрель", "messages": 10, "message_types": [("Пост", 3), ("Комментарий", 2), ("Тип не указан", 5)]},
    {"label": "Май", "messages": 0, "message_types": []},
    {"label": "Июнь", "messages": 8, "message_types": [("Пост", 4), ("Репост", 4)]},
    {"label": "Июль", "messages": 5, "message_types": []},
    {"label": "Август", "messages": 4, "message_types": [("Пост", 4)]},
]
check("Тип не указан — в конце таблицы",
      comparison_type_order(SEQ) == ["Пост", "Комментарий", "Репост", "Тип не указан"], str(comparison_type_order(SEQ)))
table = build_comparison_table(SEQ, "Типы сообщений")
check("колонки: тип и изменение его доли", list(table.columns)[:5] == ["Период", "Пост", "Δ Пост", "Комментарий", "Δ Комментарий"],
      str(list(table.columns)))
rows = {row["Период"]: row for row in table.to_dict("records")}
check("число и доля в ячейке", rows["Март"]["Пост"] == "6 · 60%", str(rows["Март"]))
check("у первой точки изменения нет", rows["Март"]["Δ Пост"] == "—")
check("изменение доли к предыдущей точке", rows["Апрель"]["Δ Пост"] == "-30,0 п.п."
      and rows["Апрель"]["Δ Комментарий"] == "-20,0 п.п.", str(rows["Апрель"]))
check("пустая точка — прочерки, а не нули", rows["Май"]["Пост"] == "—" and rows["Май"]["Δ Пост"] == "—")
check("точка после пустой — изменения нет", rows["Июнь"]["Δ Пост"] == "—" and rows["Июнь"]["Пост"] == "4 · 50%",
      str(rows["Июнь"]))
check("точка после точки без типа — изменения нет", rows["Июль"]["Пост"] == "—" and rows["Август"]["Δ Пост"] == "—"
      and rows["Август"]["Пост"] == "4 · 100%", str(rows["Август"]))
shares = type_share_rows(SEQ)
check("график — только точки с типом", set(shares["Полный период"]) == {"Март", "Апрель", "Июнь", "Август"},
      str(sorted(set(shares["Полный период"]))))
check("доли точки в сумме — 100%", all(abs(v - 1) < 1e-9 for v in shares.groupby("Полный период")["Доля"].sum()))

print("2. Сравнение периодов на экране")
from streamlit.testing.v1 import AppTest  # noqa: E402


def comparison_app():
    import pandas as pd
    import streamlit as st

    from overview_ui import render_period_comparison_metrics

    variant = st.session_state.get("variant", "")
    kinds = {"p1": ["Пост"] * 6 + ["Комментарий"] * 4, "p2": ["Пост"] * 4 + ["Комментарий"] * 4 + ["Репост"] * 2}
    rows = []
    for pid, day in (("p1", "2026-04-02"), ("p2", "2026-05-02")):
        for i, kind in enumerate(kinds[pid]):
            row = {"period_id": pid, "message_id": f"{pid}_{i}", "datetime": f"{day}T10:00:00", "date": day,
                   "text_clean": "x", "views": 10, "audience": 100, "engagement": 1, "sentiment": "нейтральная"}
            if variant != "no_types":
                row["message_type"] = kind
            rows.append(row)
    periods = pd.DataFrame([
        {"period_id": "p1", "period_name": "Апрель", "date_from": "2026-04-01", "date_to": "2026-04-07"},
        {"period_id": "p2", "period_name": "Май", "date_from": "2026-05-01", "date_to": "2026-05-07"},
    ])
    render_period_comparison_metrics(pd.DataFrame(rows), periods, ["p1", "p2"], granularity="period")


def comparison(variant=""):
    at = AppTest.from_function(comparison_app, default_timeout=60)
    at.session_state["variant"] = variant
    at.run()
    return at


at = comparison()
check("сравнение открылось", not at.exception, str(at.exception))
cards = {str(m.label): m for m in at.metric}
check("карточки типов в сравнении — к предыдущей точке",
      "Пост" in cards and str(cards["Пост"].value) == "4 · 40%" and str(cards["Пост"].delta) == "-20,0 п.п."
      and str(cards["Репост"].delta) == "+20,0 п.п.",
      str({k: (str(v.value), str(v.delta)) for k, v in cards.items()}))
check("график типов на экране", any("Динамика типов сообщений" in str(m.value) for m in at.markdown))
view = next((r for r in at.radio if str(r.label) == "Показатель"), None)
check("вид таблицы «Типы сообщений»", view is not None and "Типы сообщений" in list(view.options),
      str(list(view.options)) if view is not None else "")
view.set_value("Типы сообщений").run()
frames = [el.value for el in at.dataframe]
types_table = next((t for t in frames if "Δ Пост" in t.columns), None)
check("таблица типов по периодам", types_table is not None and list(types_table["Пост"]) == ["6 · 60%", "4 · 40%"],
      str([list(t.columns) for t in frames]))

at = comparison("no_types")
check("без типа — сравнение открылось", not at.exception, str(at.exception))
check("нет типа — нет карточек типов", not {"Пост", "Комментарий"} & {str(m.label) for m in at.metric})
view = next((r for r in at.radio if str(r.label) == "Показатель"), None)
check("нет типа — нет вида таблицы", view is not None and "Типы сообщений" not in list(view.options))
check("нет типа — вместо графика подпись", any("Тип сообщения в выгрузке не указан" in str(c.value) for c in at.caption))

print("3. Отчёт")
import docx  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402

from services import report_export  # noqa: E402
from services.report_pptx import generate_summary_pptx  # noqa: E402

LINE = "Типы сообщений: Пост — 412 (67%), Комментарий — 150 (24%), Репост — 51 (8%)."
METRICS = {"messages": 613, "sentiment": {"total": 613, "neutral": 613},
           "message_types": [("Пост", 412), ("Комментарий", 150), ("Репост", 51)]}
NO_TYPES = {"messages": 613, "sentiment": {"total": 613, "neutral": 613}, "message_types": []}


def payload(metrics, sections=None, messages=None):
    return report_export.summary_export_payload("Кнауф", "3 апреля", "Итог периода.", metrics,
                                                messages=messages, sections=sections)


def docx_text(data):
    return "\n".join(p.text for p in docx.Document(BytesIO(data)).paragraphs)


check("типы в данных отчёта", payload(METRICS)["message_types"] == METRICS["message_types"])
from_messages = payload({"messages": 3, "sentiment": {}},
                        messages=pd.DataFrame({"message_type": ["Пост", "Пост", "Репост"], "text_clean": ["x"] * 3}))
check("метрики без типов — типы из сообщений", from_messages["message_types"] == [("Пост", 2), ("Репост", 1)],
      str(from_messages["message_types"]))
check("строка о типах в Word", LINE in docx_text(report_export.generate_summary_docx(payload(METRICS))))
check("раздел «Основные метрики» выключен — строки нет",
      "Типы сообщений" not in docx_text(report_export.generate_summary_docx(payload(METRICS, sections=["sentiment"]))))
check("нет типа — в Word строки нет",
      "Типы сообщений" not in docx_text(report_export.generate_summary_docx(payload(NO_TYPES))))

import reportlab.platypus as platypus  # noqa: E402

pdf_texts = []
_Paragraph = platypus.Paragraph


class _SpyParagraph(_Paragraph):
    def __init__(self, text, *args, **kwargs):
        pdf_texts.append(str(text))
        super().__init__(text, *args, **kwargs)


platypus.Paragraph = _SpyParagraph
try:
    pdf = report_export.generate_summary_pdf(payload(METRICS))
finally:
    platypus.Paragraph = _Paragraph
check("строка о типах в PDF", pdf[:4] == b"%PDF" and LINE in pdf_texts, str([t for t in pdf_texts if "Тип" in t]))

import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402


def png_metrics(metrics):
    fig = plt.figure()
    ax = fig.add_axes([0, 0, 1, 1])
    end = report_export._draw_metrics_section(ax, payload(metrics), [], "#2563eb", 0.862)
    texts = [t.get_text() for t in ax.texts]
    plt.close(fig)
    return end, texts


end_types, texts = png_metrics(METRICS)
end_plain, plain_texts = png_metrics(NO_TYPES)
check("строка о типах в PNG", LINE in texts, str(texts))
check("строка в PNG почти не сдвигает блоки ниже", 0 < end_plain - end_types <= 0.011, f"{end_plain} → {end_types}")
check("нет типа — в PNG строки нет", not any("Типы сообщений" in t for t in plain_texts))
check("PNG собирается целиком", report_export.generate_summary_infographic_png(payload(METRICS))[:4] == b"\x89PNG")


def slide(prs, title):
    for item in prs.slides:
        if any(sh.has_text_frame and sh.text_frame.text == title for sh in item.shapes):
            return item
    return None


prs = Presentation(BytesIO(generate_summary_pptx(payload(METRICS))))
types_slide = slide(prs, "Типы сообщений")
chart = next((sh.chart for sh in types_slide.shapes if sh.has_chart), None) if types_slide else None
check("слайд «Типы сообщений» с кольцевой диаграммой", chart is not None and chart.chart_type == XL_CHART_TYPE.DOUGHNUT
      and list(chart.plots[0].categories) == ["Пост", "Комментарий", "Репост"])
titles = [next((sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text), "") for s in prs.slides]
check("слайд типов — сразу после основных метрик",
      titles.index("Типы сообщений") == titles.index("Основные метрики") + 1 if "Типы сообщений" in titles else False,
      str(titles))
try:
    no_types_slide = slide(Presentation(BytesIO(generate_summary_pptx(payload(NO_TYPES)))), "Типы сообщений")
    check("нет типа — нет слайда", no_types_slide is None)
except Exception as exc:  # noqa: BLE001 — пустая диаграмма роняет всю выгрузку
    check("нет типа — нет слайда", False, repr(exc))
check("раздел выключен — нет слайда",
      slide(Presentation(BytesIO(generate_summary_pptx(payload(METRICS, sections=["sentiment"])))), "Типы сообщений") is None)

print()
if failures:
    print(f"ПРОВАЛЕНО: {len(failures)}")
    for f in failures:
        print(f"  - {f}")
    raise SystemExit(1)
print("Типы сообщений есть в сравнении периодов и во всех форматах отчёта.")
