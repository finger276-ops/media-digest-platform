# -*- coding: utf-8 -*-
"""Отчёт в PowerPoint: те же разделы и правила, что у Word, PDF и PNG.

Мутационные проверки (что ломает какой тест):
- _metrics_slide печатает format_int вместо metric_text -> «охват без
  колонки — прочерк» краснеет;
- _sentiment_slide без проверки разметки -> «без разметки — пометка, а не
  диаграмма» краснеет;
- generate_summary_pptx без проверки sections -> «выбран один раздел —
  только он» краснеет;
- summary_pages теряет строки на границе слайда -> «весь текст саммари на
  слайдах» краснеет;
- _bar_slide без разворота порядка -> «самый крупный тег — сверху» краснеет;
- кнопки нет на странице -> «кнопка «Скачать PowerPoint» на месте»
  краснеет.
"""

import os
import sys
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
for _p in (REPO / "src", REPO / "tests"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

os.environ["SUPABASE_URL"] = "https://test.supabase.co"
os.environ["SUPABASE_SERVICE_ROLE_KEY"] = "test-key"
os.environ["PLATFORM_ADMIN_PASSWORD"] = "test-admin"

import pandas as pd  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402

from services import metrics_compute as mc  # noqa: E402
from services import report_export  # noqa: E402
from services.report_pptx import generate_summary_pptx, summary_pages  # noqa: E402

failures = []


def check(label, condition, detail=""):
    print(("  ✓ " if condition else "  ✗ ") + label + (f" — {detail}" if detail and not condition else ""))
    if not condition:
        failures.append(label)


def messages(sentiments, *, views=0, declared="audience|engagement"):
    n = len(sentiments)
    return mc.prepare_dashboard_messages(
        pd.DataFrame(
            {
                "period_id": ["a"] * n,
                "views": [views] * n,
                "audience": [100 + i for i in range(n)],
                "engagement": [i for i in range(n)],
                "sentiment": sentiments,
                "tags": ["Тарифы|Водители", "Тарифы", "Тарифы", "Приложение"][:n] + ["Тарифы"] * max(0, n - 4),
                "text_clean": ["Текст"] * n,
                mc.SOURCE_METRICS_COLUMN: [declared] * n,
            }
        )
    )


SUMMARY = "## Метрики периода\nСообщений: 4\n• Первый тезис\n" + "\n".join(
    f"Абзац {i}: " + "обсуждение тарифов и водителей " * 8 for i in range(25)
) + "\nПоследняя строка саммари"


def build(msgs, sections=None, **branding):
    metrics = mc.overview_metrics(msgs)
    metrics["comparison_sequence"] = [dict(metrics, label="01.04"), dict(metrics, label="02.04", messages=9)]
    payload = report_export.summary_export_payload(
        "Такси", "Апрель", SUMMARY, metrics, messages=msgs, events_agg=pd.DataFrame(
            [{"title": "Новый тариф", "event_title": "Новый тариф", "message_count": 12, "importance_score": 5}]
        ), branding=branding or None, sections=sections,
    )
    return Presentation(BytesIO(generate_summary_pptx(payload)))


def slide_texts(slide):
    return " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


def charts(slide):
    return [sh.chart for sh in slide.shapes if getattr(sh, "has_chart", False) and sh.has_chart]


def find(prs, title):
    for slide in prs.slides:
        if title in slide_texts(slide):
            return slide
    return None


print("1. Все разделы")
full = build(messages(["негативная", "позитивная", "нейтральная", "нейтральная"]))
titles = [slide_texts(s) for s in full.slides]
for title in ("Дайджест упоминаний", "Основные метрики", "Тональность", "Топ тегов", "Топ инфоповодов", "Главное", "Саммари периода"):
    check(f"слайд «{title}» на месте", any(title in t for t in titles), str([t[:30] for t in titles]))
check("формат 16:9", full.slide_width > full.slide_height * 1.7)
metrics_slide = find(full, "Основные метрики")
card_values = [sh.text_frame.text for sh in metrics_slide.shapes if sh.has_text_frame] if metrics_slide else []
check("охват без колонки — прочерк", card_values.count("—") == 1 and "нет в выгрузке: охват" in slide_texts(metrics_slide),
      str(card_values))
check("вовлечённость в выгрузке есть — число", "6" in card_values, str(card_values))
check("динамика — родная линейная диаграмма", metrics_slide is not None and any(
    c.chart_type == XL_CHART_TYPE.LINE_MARKERS for c in charts(metrics_slide)))
tone = find(full, "Тональность")
tone_charts = charts(tone) if tone else []
check("тональность — кольцевая диаграмма с тремя долями", bool(tone_charts) and tone_charts[0].chart_type == XL_CHART_TYPE.DOUGHNUT
      and list(tone_charts[0].plots[0].categories) == ["Позитив", "Нейтрал", "Негатив"])
tags = find(full, "Топ тегов")
tag_chart = charts(tags)[0] if tags and charts(tags) else None
check("самый крупный тег — сверху", tag_chart is not None and list(tag_chart.plots[0].categories)[-1] == "Тарифы",
      str(list(tag_chart.plots[0].categories)) if tag_chart else "")
events = find(full, "Топ инфоповодов")
tables = [sh.table for sh in events.shapes if sh.has_table] if events else []
check("инфоповоды — таблица", bool(tables) and tables[0].cell(1, 0).text == "Новый тариф", "")
all_text = " ".join(titles)
check("весь текст саммари на слайдах", all(f"Абзац {i}:" in all_text for i in range(25)) and "Последняя строка саммари" in all_text)
check("длинное саммари разбито на несколько слайдов", sum("Саммари периода" in t for t in titles) >= 3)

print("2. Прочерки и выбор разделов")
unmarked = build(messages(["", "", "", ""]))
tone = find(unmarked, "Тональность")
check("без разметки — пометка, а не диаграмма", tone is not None and not charts(tone) and "нет данных" in slide_texts(tone).lower(),
      slide_texts(tone)[:200] if tone else "")
only_metrics = build(messages(["негативная"] * 4), sections=["metrics"])
check("выбран один раздел — только он", len(only_metrics.slides) == 2, str([slide_texts(s)[:25] for s in only_metrics.slides]))
branded = build(messages(["негативная"] * 4), sections=["metrics"], accent_color="#ff0000", client_name="Клиент")
bar = [sh for sh in branded.slides[1].shapes if sh.shape_type == 1]
check("акцентный цвет проекта", any(str(sh.fill.fore_color.rgb) == "FF0000" for sh in bar if sh.fill.type == 1))
check("имя клиента на титуле", "Клиент" in slide_texts(branded.slides[0]))

print("3. Разбиение саммари")
pages = summary_pages("## Заголовок\n" + "\n".join(f"строка {i}" for i in range(40)))
check("ни одна строка не потеряна", sum(len(p) for p in pages) == 41, str([len(p) for p in pages]))
check("подзаголовок не последний на слайде", all(p[-1][0] != "heading" for p in pages))
check("пустой текст — ни одного слайда", summary_pages("") == [])

print("4. Кнопка на странице «Отчёт»")
from fake_supabase import FakeClient  # noqa: E402

import platform_store as store  # noqa: E402

CLIENT = FakeClient()
store.get_supabase_client = lambda: CLIENT
NOW = datetime.now(timezone.utc).isoformat()
CLIENT.db["platform_projects"] = [
    {"project_id": "pp", "project_name": "Такси", "status": "active", "settings": {}, "created_at": NOW, "updated_at": NOW}
]
CLIENT.db["platform_periods"] = [
    {"project_id": "pp", "period_id": "a", "period_name": "Апрель", "date_from": "2026-04-01", "date_to": "2026-04-07",
     "source_filename": "f.xlsx", "status": "active", "manifest": {}, "uploaded_at": NOW}
]
CLIENT.db["platform_table_rows"] = [
    {"project_id": "pp", "period_id": "a", "table_name": "messages", "row_id": f"m{i}",
     "payload": {"message_id": f"m{i}", "period_id": "a", "date": "2026-04-02", "datetime": "2026-04-02T10:00:00",
                 "sentiment": "негативная", "views": 10, "audience": 100, "engagement": 1, "tags": "Тарифы",
                 "text_clean": f"Сообщение {i}", "chat_title": "Чат", "author": f"a{i}"}}
    for i in range(5)
]
import streamlit as st  # noqa: E402
from streamlit.testing.v1 import AppTest  # noqa: E402

st.cache_data.clear()
at = AppTest.from_file(str(REPO / "streamlit_app.py"), default_timeout=180)
at.session_state["platform_project_id"] = "pp"
at.session_state["platform_project_role"] = "viewer"
at.session_state["platform_nav_page"] = "Отчёт"
at.session_state["period_select_pp"] = ["a"]
at.run()
check("раздел открылся", not at.exception, str(at.exception))
labels = [str(b.label) for b in at.get("download_button")]
check("кнопка «Скачать PowerPoint» на месте", "Скачать PowerPoint" in labels, str(labels))
errors = [str(e.value) for e in at.error]
check("выгрузка PowerPoint собралась без ошибки", not any("PowerPoint" in e for e in errors), str(errors))

print()
if failures:
    print(f"ПРОВАЛЕНО: {len(failures)}")
    for f in failures:
        print(f"  - {f}")
    raise SystemExit(1)
print("Отчёт в PowerPoint собирается по тем же правилам, что Word и PDF.")
