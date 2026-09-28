# -*- coding: utf-8 -*-
"""Раздел «Аудитория»: пол, возраст и география авторов — расчёт, экран,
отчёт (Word, PDF, PowerPoint) и карточка данных для ИИ.

Мутационные проверки (что ломает какой тест):
- пол по сообщениям, а не по авторам -> «пол — доли авторов, автор учтён
  один раз» краснеет;
- возраст 0 не отбрасывается как заглушка -> «возраст 0 — заглушка, не
  0–24» краснеет;
- сообщения без автора складываются в одного «автора» -> «сообщения без
  автора — разные авторы» краснеет;
- «не определено» считается регионом -> «„не определено“ — не регион»
  краснеет;
- равные регионы не по алфавиту -> «равные регионы — по алфавиту» краснеет;
- склонение без 11–14 -> «11 регионов, 12 упоминаний» краснеет;
- экран без объяснения при отсутствии данных -> «без данных — объяснение,
  откуда они берутся» краснеет;
- Word без проверки раздела -> «раздел выключен — в Word аудитории нет»
  краснеет;
- PowerPoint: регионы не перевёрнуты -> «крупнейший регион — сверху»
  краснеет;
- карточка ИИ без покрытия -> «ИИ видит, у скольких авторов пол известен»
  краснеет.
"""

import json
import os
import sys
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
for _p in (REPO / "src", REPO / "tests"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

os.environ["SUPABASE_URL"] = "https://test.supabase.co"
os.environ["SUPABASE_SERVICE_ROLE_KEY"] = "test-key"
os.environ["PLATFORM_ADMIN_PASSWORD"] = "test-admin"

import docx  # noqa: E402
import matplotlib  # noqa: E402

matplotlib.use("Agg")
import pandas as pd  # noqa: E402
from pptx import Presentation  # noqa: E402

from services.audience import (  # noqa: E402
    age_group,
    audience_summary,
    gender_label,
    mentions_text,
    places_count_text,
)
from services.formatting import plural  # noqa: E402

failures = []


def check(label, condition, detail=""):
    print(("  ✓ " if condition else "  ✗ ") + label + (f" — {detail}" if detail and not condition else ""))
    if not condition:
        failures.append(label)


def msg(author, gender="", age="", region="", city="", tag="Технониколь"):
    return {"author": author, "author_gender": gender, "author_age": age, "region": region, "city": city,
            "tags": tag, "text_clean": f"Сообщение {author}", "platform": "vk.com", "sentiment": "нейтрал"}


# 7 авторов, 10 сообщений:
# ivan пишет трижды (дважды «мужской», раз «женский»), два анонимных сообщения,
# у petr возраст 0 и регион «не определено».
ROWS = [
    msg("ivan", "Мужской", "35", "Москва", "Москва"),
    msg("ivan", "Мужской", "35", "Москва", "Москва"),
    msg("ivan", "Женский", "35", "Москва", "Москва"),
    msg("olga", "Женский", "47", "Рязанская область", "Рязань"),
    msg("olga", "Женский", "47", "Рязанская область", "Рязань"),
    msg("petr", "Мужской", "0", "не определено", "", tag="Кнауф"),
    msg("anna", "", "", "Москва", "Москва", tag="Кнауф"),
    msg("", "Мужской", "", "Калужская область", "Калуга", tag="Кнауф"),
    msg("", "Мужской", "", "Калужская область", "Калуга", tag="Кнауф"),
    msg("sergey"),
]
for index, row in enumerate(ROWS):
    row["message_id"] = f"m{index}"
MESSAGES = pd.DataFrame(ROWS)
PLAIN = MESSAGES.drop(columns=["author_gender", "author_age", "region", "city"])

print("1. Расчёт")
summary = audience_summary(MESSAGES)
gender = {g["name"]: g for g in summary["gender"]["groups"]}
check("авторов семь: анонимные сообщения — отдельные авторы", summary["authors"] == 7, str(summary["authors"]))
check("сообщения без автора — разные авторы", summary["gender"]["known"] == 5, str(summary["gender"]))
check("пол — доли авторов, автор учтён один раз",
      gender["Мужчины"]["authors"] == 4 and abs(gender["Мужчины"]["share"] - 0.8) < 1e-9
      and gender["Женщины"]["authors"] == 1, str(summary["gender"]))
check("покрытие пола — 5 из 7 авторов", abs(summary["gender"]["share_known"] - 5 / 7) < 1e-9)
ages = {g["name"]: g["authors"] for g in summary["age"]["groups"]}
check("возраст 0 — заглушка, не 0–24", summary["age"]["known"] == 2 and ages.get("0–24") == 0, str(ages))
check("возраст по группам", ages.get("35–44") == 1 and ages.get("45–54") == 1, str(ages))
geo = summary["geo"]
check("„не определено“ — не регион", geo["regions_count"] == 3 and geo["known"] == 5, str(geo))
check("география — по упоминаниям", [(p["name"], p["messages"]) for p in geo["regions"]][:1] == [("Москва", 4)]
      and geo["messages_known"] == 8, str(geo["regions"]))
check("равные регионы — по алфавиту", [p["name"] for p in geo["regions"]]
      == ["Москва", "Калужская область", "Рязанская область"], str(geo["regions"]))
check("города", [p["name"] for p in geo["cities"]] == ["Москва", "Калуга", "Рязань"], str(geo["cities"]))
check("без колонок аудитории — данных нет", audience_summary(PLAIN)["has_data"] is False)
check("пустая выборка — данных нет", audience_summary(pd.DataFrame())["has_data"] is False)
check("возраст числом и диапазоном",
      [age_group(v) for v in ("35", "25-34", "до 18", "55+", 70, "150", "abc")]
      == ["35–44", "25–34", "0–24", "55 и старше", "55 и старше", "", ""])
check("пол в разных написаниях",
      [gender_label(v) for v in ("Мужской", "Ж", "female", "woman", "не определено")]
      == ["Мужчины", "Женщины", "Женщины", "Женщины", ""])
check("склонение: 1 регион, 2 региона, 5 регионов, 21 регион",
      [plural(n, "регион", "региона", "регионов") for n in (1, 2, 5, 21)] == ["регион", "региона", "регионов", "регион"])
check("11 регионов, 12 упоминаний", plural(11, "регион", "региона", "регионов") == "регионов"
      and mentions_text(12) == "12 упоминаний" and mentions_text(1) == "1 упоминание")
check("«из 1 региона и 21 города»", places_count_text({"regions_count": 1, "cities_count": 21})
      == "1 региона и 21 города")

print("2. Отчёт: Word, PDF, PowerPoint")
from services.report_export import (  # noqa: E402
    generate_summary_docx,
    generate_summary_infographic_png,
    generate_summary_pdf,
    summary_export_payload,
)
from services.dashboard_config import DEFAULT_REPORT_SECTIONS, REPORT_SECTION_OPTIONS  # noqa: E402
from services.project_settings import report_sections_from_project_settings  # noqa: E402
from services.report_pptx import generate_summary_pptx  # noqa: E402


def payload(sections=None, messages=MESSAGES):
    return summary_export_payload("Проект", "Июль", "- тезис", {"messages": len(messages), "sentiment": {}},
                                  messages=messages, sections=sections)


check("раздел в конструкторе отчёта и включён по умолчанию",
      REPORT_SECTION_OPTIONS.get("audience", "").startswith("Аудитория") and "audience" in DEFAULT_REPORT_SECTIONS)


def docx_paragraphs(data):
    return [p.text for p in docx.Document(BytesIO(data)).paragraphs]


paragraphs = docx_paragraphs(generate_summary_docx(payload()))
check("Word: заголовок «Аудитория»", "Аудитория" in paragraphs)
check("Word: пол с покрытием", "Пол: мужчины — 80 %, женщины — 20 % (пол известен у 71 % авторов)." in paragraphs,
      str([p for p in paragraphs if p.startswith("Пол")]))
check("Word: возраст с покрытием", "Возраст: 35–44 — 50 %, 45–54 — 50 % (возраст известен у 29 % авторов)."
      in paragraphs, str([p for p in paragraphs if p.startswith("Возраст")]))
check("Word: география", "География: авторы из 3 регионов и 3 городов (место известно у 71 % авторов)." in paragraphs
      and "Регионы: Москва — 4 упоминания; Калужская область — 2 упоминания; Рязанская область — 2 упоминания."
      in paragraphs, str([p for p in paragraphs if p.startswith(("География", "Регионы"))]))
check("раздел выключен — в Word аудитории нет",
      "Аудитория" not in docx_paragraphs(generate_summary_docx(payload(sections=["metrics", "highlights"]))))
check("нет данных об авторах — в Word раздела нет",
      "Аудитория" not in docx_paragraphs(generate_summary_docx(payload(messages=PLAIN))))

import reportlab.platypus as platypus  # noqa: E402

pdf_texts = []
_Paragraph = platypus.Paragraph


class _Spy(_Paragraph):
    def __init__(self, text, *args, **kwargs):
        pdf_texts.append(str(text))
        super().__init__(text, *args, **kwargs)


platypus.Paragraph = _Spy
try:
    generate_summary_pdf(payload())
finally:
    platypus.Paragraph = _Paragraph
check("PDF: блок «Аудитория» с полом", "Аудитория" in pdf_texts
      and any(t.startswith("• Пол: мужчины — 80 %") for t in pdf_texts), str([t for t in pdf_texts if "Пол" in t]))


def slides(data):
    prs = Presentation(BytesIO(data))
    result = {}
    for slide in prs.slides:
        texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        charts = [sh.chart for sh in slide.shapes if sh.has_chart]
        if texts:
            result[texts[0]] = (texts, charts)
    return result


deck = slides(generate_summary_pptx(payload()))
people = deck.get("Аудитория: пол и возраст")
check("PowerPoint: слайд «пол и возраст» с двумя диаграммами", people is not None and len(people[1]) == 2,
      str(list(deck)))
if people:
    sex = people[1][0].plots[0]
    check("PowerPoint: доли пола", list(sex.categories) == ["Мужчины", "Женщины"]
          and [round(v, 3) for v in sex.series[0].values] == [0.8, 0.2], str(list(sex.series[0].values)))
    check("PowerPoint: покрытие под диаграммами", any("Пол известен у 71 % авторов" in t for t in people[0]))
places = deck.get("Аудитория: география")
check("PowerPoint: слайд «география»", places is not None and len(places[1]) == 2, str(list(deck)))
if places:
    # Столбцы горизонтальной диаграммы идут снизу вверх: крупнейший — последним.
    check("крупнейший регион — сверху", list(places[1][0].plots[0].categories)
          == ["Рязанская область", "Калужская область", "Москва"], str(list(places[1][0].plots[0].categories)))
check("раздел выключен — слайдов аудитории нет",
      not any(t.startswith("Аудитория") for t in slides(generate_summary_pptx(payload(sections=["metrics"])))))
check("нет данных — слайдов аудитории нет",
      not any(t.startswith("Аудитория") for t in slides(generate_summary_pptx(payload(messages=PLAIN)))))
png = generate_summary_infographic_png(payload(sections=["audience"]))
check("PNG собирается и с одной «Аудиторией»", png[:4] == b"\x89PNG")
check("набор, сохранённый до «Аудитории», получает её сам",
      "audience" in report_sections_from_project_settings(
          {"report_sections": ["metrics"], "report_sections_known": [s for s in REPORT_SECTION_OPTIONS if s != "audience"]}))
check("выключенная осознанно «Аудитория» не возвращается",
      "audience" not in report_sections_from_project_settings(
          {"report_sections": ["metrics"], "report_sections_known": list(REPORT_SECTION_OPTIONS)}))

print("3. Карточка данных для ИИ")
from services.ai_summary import build_data_card  # noqa: E402


def card(messages):
    return build_data_card(project_name="П", periods=pd.DataFrame(), period_ids=["p1"], messages=messages,
                           events_agg=pd.DataFrame(), include_excerpts=False)


text = card(MESSAGES)
check("ИИ видит, у скольких авторов пол известен",
      "- Пол (известен у 71 % авторов): мужчины — 80 %, женщины — 20 %" in text,
      text[text.find("Аудитория"):][:300])
check("ИИ видит регионы с покрытием", "- Регионы (место известно у 71 % авторов; авторы из 3 регионов и 3 городов)"
      in text)
check("без данных — блока аудитории нет", "Аудитория (" not in card(PLAIN))

print("4. Раздел в приложении")
from fake_supabase import FakeClient  # noqa: E402

import platform_store as store  # noqa: E402

CLIENT = FakeClient()
store.get_supabase_client = lambda: CLIENT
NOW = datetime.now(timezone.utc).isoformat()
CLIENT.db["platform_projects"] = [
    {"project_id": pid, "project_name": name, "status": "active", "settings": {}, "created_at": NOW, "updated_at": NOW}
    for pid, name in (("aud", "Аудитория"), ("plain", "Без аудитории"))
]
CLIENT.db["platform_periods"] = [
    {"project_id": pid, "period_id": "p1", "period_name": "Июль", "date_from": "2026-07-01", "date_to": "2026-07-07",
     "source_filename": "f.xlsx", "status": "active", "manifest": {}, "uploaded_at": NOW}
    for pid in ("aud", "plain")
]
CLIENT.db["platform_manual_rows"] = []
table_rows = []
for pid, frame in (("aud", MESSAGES), ("plain", PLAIN)):
    for row in frame.to_dict("records"):
        payload_row = dict(row, period_id="p1", date="2026-07-02", datetime="2026-07-02T10:00:00", event_title="",
                           message_link=f"https://vk.com/wall-1_{row['message_id']}", views=1, audience=1, engagement=1)
        table_rows.append({"project_id": pid, "period_id": "p1", "table_name": "messages",
                           "row_id": row["message_id"], "payload": payload_row})
CLIENT.db["platform_table_rows"] = table_rows

import streamlit as st  # noqa: E402
from streamlit.testing.v1 import AppTest  # noqa: E402

from audience_ui import NO_DATA_TEXT  # noqa: E402


def open_audience(project, state=None):
    st.cache_data.clear()
    at = AppTest.from_file(str(REPO / "streamlit_app.py"), default_timeout=180)
    at.session_state["platform_project_id"] = project
    at.session_state["platform_project_role"] = "viewer"
    at.session_state["platform_nav_page"] = "Аудитория"
    at.session_state[f"period_select_{project}"] = ["p1"]
    for key, value in (state or {}).items():
        at.session_state[key] = value
    at.run()
    return at


at = open_audience("aud")
check("раздел открылся без исключений", not at.exception, str(at.exception))
check("раздел есть в меню пользователя", "Аудитория" in [str(b.label) for b in at.sidebar.button])
cards = {str(m.label): str(m.value) for m in at.metric}
check("карточки: авторы и покрытие", cards.get("Авторов") == "7" and cards.get("Пол известен") == "71 %"
      and cards.get("Возраст известен") == "29 %" and cards.get("Место известно") == "71 %", str(cards))
# В спецификации кириллица экранирована (\u0413…) — сравниваем разобранный JSON.
charts = [json.dumps(json.loads(el.proto.spec), ensure_ascii=False) for el in at.get("vega_lite_chart")]
check("четыре диаграммы: пол, возраст, регионы, города", len(charts) == 4
      and sum('"Группа"' in spec for spec in charts) == 2 and sum('"Место"' in spec for spec in charts) == 2,
      str(len(charts)))
check("подпись географии", any("из 3 регионов и 3 городов" in str(c.value) for c in at.caption),
      str([str(c.value)[:80] for c in at.caption]))

at = open_audience("aud", {"tag_slice::aud": ["Кнауф"]})
cards = {str(m.label): str(m.value) for m in at.metric}
check("срез по тегу действует и на «Аудиторию»", cards.get("Авторов") == "4", str(cards))

at = open_audience("plain")
check("без данных — объяснение, откуда они берутся", not at.exception
      and any(str(i.value) == NO_DATA_TEXT for i in at.info), str([str(i.value)[:80] for i in at.info]))

print()
if failures:
    print(f"ПРОВАЛЕНО: {len(failures)}")
    for f in failures:
        print(f"  - {f}")
    raise SystemExit(1)
print("«Аудитория» считает пол, возраст и географию по авторам и доходит до отчёта.")
