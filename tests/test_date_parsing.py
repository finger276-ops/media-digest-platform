# -*- coding: utf-8 -*-
"""Даты выгрузки: ячейки Excel и ISO читаются верно, «дд.мм.гггг» — как раньше.

Раньше колонка дат разбиралась одним pd.to_datetime(..., dayfirst=True).
Выгрузки читаются как текст, и ячейка-дата Excel приезжает строкой
«2026-04-05 09:00:00» — pandas читал её как 4 мая, а если в первой строке
день не больше 12, дни с 13-го по 31-е не разбирались вовсе. Апрельский файл
становился периодом «04.01.2026–04.12.2026», половина сообщений оставалась
без даты, period_id считался от неверного названия. Раздельные «Дата» и
«Время» склеивались в «2026-04-01 00:00:00 09:30:00» — это читалось как
4 января. Пояса Z и +03:00 в одной колонке обрывали загрузку.

Разделы:
1. Разбор колонки (services/date_parsing.py): ISO при первом дне ≤ 12 — все
   30 дней апреля; «дд.мм.гггг» — до байта как прежний разбор; смесь поясов —
   московское время без ошибки, время без пояса не сдвигается; номера дней
   Excel; пустое и мусор — NaT.
2. Склейка «Дата» + «Время»: ячейки Excel (полночь, 1900-01-01, доля суток),
   текст склеивается как раньше.
3. Сквозная загрузка через поддельный Supabase: xlsx с настоящими
   ячейками-датами (openpyxl), xlsx в раскладке Brand Analytics с раздельными
   «Дата» и «Время», CSV с ISO-датами и с поясами; название периода
   «01.04.2026–30.04.2026» и period_id от него.
4. Отчёт о загрузке: сообщения без даты — в итоге, в манифесте, строкой;
   на странице «Загрузка файла» — предупреждение при доле больше 5 % (файл
   загружается), справка при меньшей, ничего при нуле; очередь автозагрузки
   и сводка пересборки.
5. Проверка уже загруженных периодов (services/period_date_check.py,
   scripts/check_period_dates.py): признаки неверных дат, база не меняется.

Мутационные проверки (что ломает какой тест):
- в _parse_year_first разбирать время без пояса с dayfirst=True вместо
  format="ISO8601" -> краснеют 1 «ISO: все 30 дней апреля» и «время без
  пояса не сдвигается», 3 у xlsx с ячейками-датами, раскладки Brand
  Analytics и CSV с ISO-датами — название периода «04.01.2026–04.12.2026»,
  period_id, месяцы сообщений (21 проверка);
- разбирать время без пояса с utc=True (сдвиг на московское) -> краснеют
  1 «ISO: время сохранено», «время без пояса не сдвигается», 3 «часы и минуты
  сохранены» (7 проверок);
- не убирать полночь у ячейки-даты при склейке (убрать замену _ISO_MIDNIGHT
  в join_date_and_time) -> краснеют 2 «полночь ячейки-даты убрана»,
  «время-ячейка 1900-01-01», «время долей суток», 3 «раскладка Brand
  Analytics: дата и время склеены без полуночи» (9 проверок);
- не считать сообщения без даты (undated_messages всегда 0) -> краснеют 3
  «без даты — в итоге и манифесте», «пустая дата», 4 «страница: 5 из 60
  (8 %) — предупреждение» и «1 из 120 — справка»;
- остаток строк с года разбирать с dayfirst=True -> краснеет 1 «остаток
  строк с года — без перестановки»;
- разбирать остаток одним вызовом, не отделяя строки с поясом -> краснеют
  1 «остаток строк с года», «дд.мм.гггг со смесью поясов»;
- вернуть прежний признак пояса (без времени перед ним) -> краснеют 1
  «дд-мм-гггг — как прежний разбор» и «полночь, не 03:00»;
- не исключать время без даты из остатка -> краснеют 1 «время без даты —
  не дата», 2 «без даты склейка как раньше», 3 «пустая дата»;
- решать про колонку даты по имени, а не по содержимому -> краснеют 3 и 4
  «файл без колонки даты»;
- обрезать весь комментарий задачи до 160 знаков -> краснеет 4 «строка о
  датах видна целиком»;
- «перепутаны» уже при соседних месяцах или двух месяцах у сообщений ->
  краснеют 5 «05.03–05.04» и «сутки на стыке месяцев»;
- долю суток понимать только меньше единицы, не снимать «1 day,», понимать
  только целые номера дня в склейке -> краснеют 2 «время больше суток» или
  «дробный номер дня», 3 «время больше суток и дробный номер дня».
"""

import copy
import io
import os
import sys
import tempfile
import warnings
from datetime import date, datetime
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
for _p in (REPO / "src", REPO / "scripts", REPO / "tests"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

os.environ["SUPABASE_URL"] = "https://test.supabase.co"
os.environ["SUPABASE_SERVICE_ROLE_KEY"] = "test-key"
os.environ["PLATFORM_ADMIN_PASSWORD"] = "test-admin"

# Прежний разбор в сверках ниже предупреждает, что разбирает даты поштучно.
warnings.filterwarnings("ignore", message="Could not infer format")
warnings.filterwarnings("ignore", message="Parsing dates in")

import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
from pandas.testing import assert_series_equal  # noqa: E402

from fake_supabase import FakeClient, Table  # noqa: E402

import platform_store as store  # noqa: E402

CLIENT = FakeClient()
store.get_supabase_client = lambda: CLIENT
store._retry_sleep = lambda _seconds: None

from services.date_parsing import join_date_and_time, parse_source_datetimes  # noqa: E402
from services.import_report import undated_line, undated_needs_warning  # noqa: E402
from services.ingest import process_canonical, read_canonical_bytes  # noqa: E402
from services.message_normalize import UNDATED_MESSAGES  # noqa: E402
from services.period_date_check import (  # noqa: E402
    messages_date_problems,
    period_date_problems,
)
from services.text_cleaning import parse_datetime  # noqa: E402

failures = []


def check(label, condition, detail=""):
    print(("  ✓ " if condition else "  ✗ ") + label + (f" — {detail}" if detail and not condition else ""))
    if not condition:
        failures.append(label)


def parse_datetime_before(series):
    """Прежний разбор (до версии 4.12.8) — дословно."""
    prepared = (
        series.fillna("").astype(str).str.replace("\xa0", " ", regex=False).str.strip()
    )
    return pd.to_datetime(prepared, errors="coerce", dayfirst=True)


def glue_before(dates, times):
    """Прежняя склейка «Дата» + «Время» — дословно."""
    return [
        (f"{str(d).strip()} {str(t).strip()}".strip() if str(t).strip() else str(d).strip())
        for d, t in zip(dates, times)
    ]


def same_series(got, expected):
    try:
        assert_series_equal(got, expected, check_exact=True, check_names=False)
        return ""
    except AssertionError as exc:
        return str(exc)[:300]


APRIL = [date(2026, 4, day) for day in range(1, 31)]

print("1. Разбор колонки дат")
# Первая строка — 1 апреля: именно такой файл прежний разбор ломал сильнее
# всего (формат «год-день-месяц» по первой строке).
iso = pd.Series([f"{d:%Y-%m-%d} {9 + d.day % 8:02d}:15:00" for d in APRIL])
parsed = parse_datetime(iso)
check("ISO: все 30 дней апреля — апрель, день тот же",
      parsed.notna().all() and (parsed.dt.month == 4).all() and parsed.dt.day.tolist() == list(range(1, 31)),
      str(parsed.tolist()[:5]))
check("ISO: время сохранено", parsed.dt.hour.tolist() == [9 + d.day % 8 for d in APRIL])
before = parse_datetime_before(iso)
check("прежний разбор на этих данных ошибался (тест стережёт настоящую ошибку)",
      before.isna().sum() >= 18 and (before.dropna().dt.month != 4).any(),
      f"NaT: {before.isna().sum()}, месяцы: {sorted(set(before.dropna().dt.month))}")
check("ISO без времени и с T, одной колонкой",
      parse_datetime(pd.Series(["2026-04-05", "2026-04-13T10:00:00", "2026-04-07 09:00", "2026/04/08 10:00"])).tolist()
      == [pd.Timestamp("2026-04-05"), pd.Timestamp("2026-04-13 10:00"), pd.Timestamp("2026-04-07 09:00"),
          pd.Timestamp("2026-04-08 10:00")])

# «дд.мм.гггг» прежний разбор читал верно — результат должен совпасть до
# байта, вместе с типом колонки и пустыми значениями.
RU_COLUMNS = {
    "дд.мм.гггг чч:мм": [f"{d:%d.%m.%Y} {8 + d.day % 12:02d}:{d.day:02d}" for d in APRIL],
    "дд.мм.гггг чч:мм:сс": [f"{d:%d.%m.%Y} 10:{d.day:02d}:{59 - d.day:02d}" for d in APRIL],
    "дд.мм.гггг": [f"{d:%d.%m.%Y}" for d in APRIL],
    "д.м.гггг ч:мм": [f"{d.day}.{d.month}.{d.year} {d.day % 10}:05" for d in APRIL],
    "дд.мм.гг": [f"{d:%d.%m.%y}" for d in APRIL],
    "дд.мм.гг чч:мм": [f"{d:%d.%m.%y} 12:30" for d in APRIL],
    "с пропусками и пробелами": ["", None, " 05.04.2026 09:00", "13.04.2026\xa010:00", float("nan"), "30.04.2026 23:59 "],
    "мусор в начале": ["не дата", "05.04.2026 09:00", "13.04.2026 10:00"],
    "одна дата": ["13.04.2026 10:00"],
    # Хвост «-2026» похож на пояс «-20:26»: такие даты не должны уезжать
    # на 03:00 как время с поясом.
    "дд-мм-гггг": ["24-04-2026", "05-04-2026", "30-04-2026"],
    "дд-мм-гггг, первая ≤ 12": ["05-04-2026", "24-04-2026"],
    "дд-мм-гггг чч:мм": ["05-04-2026 09:00", "24-04-2026 10:15"],
}
for label, values in RU_COLUMNS.items():
    column = pd.Series(values)
    problem = same_series(parse_datetime(column), parse_datetime_before(column))
    check(f"«{label}» — как прежний разбор, до байта", not problem, problem)
# Смесь русских форматов в одной колонке прежний разбор понимал только по
# формату первой строки. Новый разбирает всё, а там, где прежний разобрал,
# совпадает с ним.
mixed = pd.Series(["05.04.2026 09:00", "13.04.2026", "14.04.2026 10:00:00", "2026-04-15 11:00:00", "31.02.2026"])
new_mixed, old_mixed = parse_datetime(mixed), parse_datetime_before(mixed)
known = old_mixed.notna()
check("смесь форматов: где прежний разобрал — то же самое",
      (new_mixed[known] == old_mixed[known]).all(), f"{new_mixed.tolist()} / {old_mixed.tolist()}")
check("смесь форматов: разобрано всё, кроме несуществующего 31.02",
      new_mixed.isna().tolist() == [False, False, False, False, True], str(new_mixed.tolist()))

zones = pd.Series([
    "2026-04-05T06:00:00Z",
    "2026-04-05T09:00:00+03:00",
    "2026-04-05T10:00:00+04:00",
    "2026-04-05T06:00:00.000+0000",
    "2026-04-05 09:00:00",
])
try:
    parse_datetime_before(zones)
    before_error = None
except ValueError as exc:
    before_error = exc
check("прежний разбор падал на смеси поясов", before_error is not None, "исключения не было")
try:
    got_zones = parse_datetime(zones)
    zone_error = None
except Exception as exc:  # noqa: BLE001 — исключение и есть результат проверки
    got_zones, zone_error = None, exc
check("смесь поясов Z и +03:00 — без исключения", zone_error is None, repr(zone_error))
check("время с поясом переведено в московское (09:00 у всех)",
      got_zones is not None and got_zones.tolist() == [pd.Timestamp("2026-04-05 09:00")] * 5,
      str(None if got_zones is None else got_zones.tolist()))
naive = parse_datetime(pd.Series(["2026-04-05 23:30:00", "2026-04-30 00:15:00"]))
check("время без пояса не сдвигается (это московское время из выгрузки)",
      naive.tolist() == [pd.Timestamp("2026-04-05 23:30"), pd.Timestamp("2026-04-30 00:15")], str(naive.tolist()))
check("колонка без пояса — дата без пояса", naive.dt.tz is None and got_zones is not None and got_zones.dt.tz is None)
dashed = parse_datetime(pd.Series(["05-04-2026", "13-04-2026", "5-4-2026"]))
check("«дд-мм-гггг» без времени — полночь, не 03:00 (хвост «-2026» — не пояс)",
      dashed.tolist() == [pd.Timestamp("2026-04-05"), pd.Timestamp("2026-04-13"), pd.Timestamp("2026-04-05")],
      str(dashed.tolist()))
# Строки с года, которые не подошли под ISO (запятая, пояс словом), идут в
# остаток — и там без dayfirst: иначе 5 апреля снова стало бы 4 мая.
year_rest = parse_datetime(pd.Series(["2026-04-05, 09:00", "2026-04-13, 10:00", "2026-04-05T06:00:00 UTC"]))
check("остаток строк с года — без перестановки дня и месяца, «UTC» — в московское",
      year_rest.tolist() == [pd.Timestamp("2026-04-05 09:00"), pd.Timestamp("2026-04-13 10:00"),
                             pd.Timestamp("2026-04-05 09:00")],
      str(year_rest.tolist()))
# «дд.мм.гггг» с разными поясами в остатке: строки с поясом разбираются
# отдельно, иначе смесь поясов сделала бы их все пустыми.
ru_zones = parse_datetime(pd.Series(["05.04.2026 09:00 +03:00", "05.04.2026 06:00 +00:00", "13.04.2026 10:00"]))
check("«дд.мм.гггг» со смесью поясов — московское время у всех",
      ru_zones.tolist() == [pd.Timestamp("2026-04-05 09:00"), pd.Timestamp("2026-04-05 09:00"),
                            pd.Timestamp("2026-04-13 10:00")],
      str(ru_zones.tolist()))
# Время без даты прежний разбор превращал в сегодняшнее число: сообщение
# уезжало в день загрузки, период растягивался до него.
time_only = pd.Series(["09:15", "9:15:00", "09:15 PM", "09:15 +03:00", "today", "now", "T09:15"])
today = pd.Timestamp.now().normalize()
old_time_only = parse_datetime_before(time_only.iloc[:2]).dt.normalize()
check("прежний разбор делал из времени без даты сегодняшнее число",
      old_time_only.isin([today, pd.Timestamp.now().normalize()]).all(), str(old_time_only.tolist()))
check("время без даты и «today»/«now» — не дата (NaT)",
      parse_datetime(time_only).isna().all(), str(parse_datetime(time_only).tolist()))

serials = parse_datetime(pd.Series(["46113", "46113.375", "46142.99999999"]))
check("номер дня Excel — дата (46113 = 01.04.2026, доля — время)",
      serials.tolist() == [pd.Timestamp("2026-04-01"), pd.Timestamp("2026-04-01 09:00"), pd.Timestamp("2026-05-01")],
      str(serials.tolist()))
junk = pd.Series(["", None, float("nan"), "не дата", "31.02.2026", "   ", "\xa0", "12345678"])
junk_parsed = parse_datetime(junk)
check("пустое и мусор — NaT, без исключения", junk_parsed.isna().all(), str(junk_parsed.tolist()))
odd_index = pd.Series(["2026-04-05", "05.04.2026"], index=[7, 7])
check("индекс входа сохраняется, даже повторяющийся",
      parse_datetime(odd_index).index.tolist() == [7, 7]
      and parse_datetime(odd_index).tolist() == [pd.Timestamp("2026-04-05")] * 2)
check("parse_datetime и есть общий разбор (services.date_parsing)",
      not same_series(parse_datetime(iso), parse_source_datetimes(iso)))

print("2. Склейка «Дата» + «Время»")
glued = join_date_and_time(
    pd.Series(["2026-04-01 00:00:00", "2026-04-13 00:00:00", "2026-04-30 00:00:00", "46113", "2026-04-05 09:00:00", ""]),
    pd.Series(["09:30:00", "1900-01-01 10:15:05", "0.395833333333", "09:00:00", "10:00:00", "11:00"]),
).tolist()
check("полночь ячейки-даты убрана перед временем", glued[0] == "2026-04-01 09:30:00", glued[0])
check("время-ячейка «1900-01-01 чч:мм:сс» — только время", glued[1] == "2026-04-13 10:15:05", glued[1])
check("время долей суток — чч:мм:сс", glued[2] == "2026-04-30 09:30:00", glued[2])
check("номер дня Excel в дате — дата", glued[3] == "2026-04-01 09:00:00", glued[3])
check("своё время в ячейке даты не портится колонкой «Время»", glued[4] == "2026-04-05 09:00:00", glued[4])
check("без даты склейка как раньше — одно время, но датой оно не становится",
      glued[5] == "11:00" and parse_datetime(pd.Series([glued[5]])).isna().all(), glued[5])
# Время больше суток openpyxl отдаёт по-разному: промежутком «1 day, 9:15:00»
# (формат «[ч]:мм»), числом без формата, записью с порядком.
long_times = join_date_and_time(
    pd.Series(["2026-04-01 00:00:00"] * 4),
    pd.Series(["1 day, 9:15:00", "2 days, 9:15:00.500000", "1.385416666666667", "1.1574074074074073e-05"]),
).tolist()
check("время больше суток в любом виде — время суток",
      long_times == ["2026-04-01 9:15:00", "2026-04-01 9:15:00.500000", "2026-04-01 09:15:00", "2026-04-01 00:00:01"],
      str(long_times))
long_parsed = parse_datetime(pd.Series(long_times))
check("время больше суток разбирается как дата",
      long_parsed.dt.strftime("%Y-%m-%d %H:%M:%S").tolist()
      == ["2026-04-01 09:15:00", "2026-04-01 09:15:00", "2026-04-01 09:15:00", "2026-04-01 00:00:01"],
      str(long_parsed.tolist()))
# Ячейка «дата и время» без формата — дробный номер дня: своё время в нём есть.
serial_with_time = join_date_and_time(pd.Series(["46113.38541666666", "46113.0"]), pd.Series(["10:00:00", "10:00:00"])).tolist()
check("дробный номер дня Excel — дата со своим временем, целый — дата плюс «Время»",
      serial_with_time == ["2026-04-01 09:15:00", "2026-04-01 10:00:00"], str(serial_with_time))
# «9.30» — время, написанное человеком через точку, а не доля суток.
text_dates = pd.Series(["24.04.2026", "25.04.2026", " 26.04.2026 ", "", "27.04.2026", "28.04.2026", "29.04.2026"])
text_times = pd.Series(["09:30", "", "10:00", "11:00", " 12:15 ", "13:00:00", "9.30"])
check("текстовые дата и время склеиваются как раньше",
      join_date_and_time(text_dates, text_times).tolist() == glue_before(text_dates, text_times),
      f"{join_date_and_time(text_dates, text_times).tolist()} против {glue_before(text_dates, text_times)}")

print("3. Сквозная загрузка через поддельный Supabase")
PROJECT = "dates"
CLIENT.db["platform_projects"] = [{"project_id": PROJECT, "project_name": "Даты", "status": "active", "settings": {}}]
WORK_DIR = tempfile.mkdtemp(prefix="dates_")
WORDS = ["доставка", "цена", "качество", "магазин", "кровля", "утеплитель", "монтаж", "отзыв"]


def text_of(i):
    return f"Сообщение {i}: {WORDS[i % len(WORDS)]} и {WORDS[(i * 3) % len(WORDS)]}, вопрос номер {i}"


def xlsx_bytes(header, rows, formats=None, cell_formats=None):
    """Книга Excel с настоящими ячейками: дата — datetime, время — time.

    formats — формат числа по колонке (номер с 1), cell_formats — по ячейке
    {(номер строки данных с 0, номер колонки с 1): формат}.
    """
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Сообщения"
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    for column, number_format in (formats or {}).items():
        for line in range(2, len(rows) + 2):
            sheet.cell(line, column).number_format = number_format
    for (line, column), number_format in (cell_formats or {}).items():
        sheet.cell(line + 2, column).number_format = number_format
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def upload(file_bytes, filename, source_system="auto"):
    CLIENT.db["platform_periods"] = []
    CLIENT.db["platform_table_rows"] = []
    canonical = read_canonical_bytes(file_bytes, filename, source_system)
    result = process_canonical(
        canonical, project_id=PROJECT, source_filename=filename, file_bytes=file_bytes,
        source_system=source_system, work_dir=WORK_DIR,
    )
    stamps = pd.Series(
        [r["payload"].get("datetime") for r in CLIENT.db["platform_table_rows"] if r["table_name"] == "messages"],
        dtype="object",
    )
    manifest = (CLIENT.db["platform_periods"][0] if CLIENT.db["platform_periods"] else {}).get("manifest") or {}
    return canonical, result, pd.to_datetime(stamps, errors="coerce"), manifest


def check_april(where, result, stamps):
    name = "01.04.2026–30.04.2026"
    check(f"[{where}] название периода «{name}»", result["period_name"] == name, result["period_name"])
    check(f"[{where}] даты периода 2026-04-01 … 2026-04-30",
          (result["date_from"], result["date_to"]) == ("2026-04-01", "2026-04-30"),
          f"{result['date_from']} … {result['date_to']}")
    check(f"[{where}] period_id — от названия по датам",
          result["period_id"] == store.make_period_id(PROJECT, name, ""), result["period_id"])
    check(f"[{where}] у всех сообщений есть дата, все в апреле, все 30 дней",
          len(stamps) == 60 and stamps.notna().all() and (stamps.dt.month == 4).all()
          and sorted(set(stamps.dt.day)) == list(range(1, 31)),
          f"{len(stamps)} сообщений, без даты {stamps.isna().sum()}, месяцы {sorted(set(stamps.dropna().dt.month))}")
    check(f"[{where}] без даты — ноль в итоге загрузки", result.get(UNDATED_MESSAGES) == 0, str(result.get(UNDATED_MESSAGES)))


# Две публикации в день: утром и вечером, всего 60 строк.
MOMENTS = [datetime(d.year, d.month, d.day, hour, 5 * (d.day % 12)) for d in APRIL for hour in (9, 18)]

# (а) Обычная таблица, дата — настоящая ячейка «дата и время».
generic_rows = [
    [moment, text_of(i), f"https://t.me/news/{i}", f"Автор {i % 7}", "telegram"]
    for i, moment in enumerate(MOMENTS)
]
generic_xlsx = xlsx_bytes(["Дата", "Текст", "Ссылка", "Автор", "Площадка"], generic_rows, {1: "dd.mm.yyyy hh:mm"})
canonical, result, stamps, manifest = upload(generic_xlsx, "апрель.xlsx")
check("xlsx с ячейками-датами: «Дата» приезжает строкой с года",
      str(canonical.loc[0, "Дата"]) == "2026-04-01 09:05:00", str(canonical.loc[0, "Дата"]))
check("прежний разбор на этом файле ошибался",
      parse_datetime_before(canonical["Дата"]).isna().sum() >= 30, str(parse_datetime_before(canonical["Дата"]).isna().sum()))
check_april("xlsx с ячейками-датами", result, stamps)
check("[xlsx с ячейками-датами] часы и минуты сохранены",
      sorted(stamps.dt.strftime("%Y-%m-%d %H:%M")) == sorted(m.strftime("%Y-%m-%d %H:%M") for m in MOMENTS))

# (б) Раскладка Brand Analytics: дата и время — отдельные ячейки. Время одной
# строки — доля суток числом, других — время больше суток: в ячейке «чч:мм»
# openpyxl отдаёт его датой 1900-01-01, в «[ч]:мм» — промежутком «1 day, …»,
# без формата — числом. У одной строки в ячейке даты своё время, а формата
# нет: она приезжает дробным номером дня.
BA_HEADER = ["Дата", "Время", "ID сообщения", "Url", "Источник", "Тип источника", "Текст", "Автор", "Сюжет"]
ba_rows = []
for i, moment in enumerate(MOMENTS):
    if i == 3:
        clock = 0.4375  # 10:30 числом
    elif i in (4, 5, 6):
        clock = 1.4375  # 10:30 в ячейке «больше суток»
    else:
        clock = moment.time()
    ba_rows.append([
        moment if i == 8 else moment.date(), clock, str(1000 + i), f"https://vk.com/wall-1_{i}", "vk.com",
        "Соцсети", text_of(i), f"Автор {i % 9}", f"Сюжет {i % 4}",
    ])
BA_FORMATS = {1: "dd.mm.yyyy", 2: "hh:mm"}
# Доля суток без формата времени — просто число.
BA_CELL_FORMATS = {(3, 2): "General", (5, 2): "[h]:mm", (6, 2): "General", (8, 1): "General"}
ba_xlsx = xlsx_bytes(BA_HEADER, ba_rows, BA_FORMATS, BA_CELL_FORMATS)
raw_ba = pd.read_excel(io.BytesIO(ba_xlsx), dtype=str, keep_default_na=False)
check("раскладка Brand Analytics: ячейки читаются как «2026-04-01 00:00:00», «1900-01-01 …», «0.4375», "
      "«1 day, …», «1.4375», дробный номер дня",
      raw_ba.loc[0, "Дата"] == "2026-04-01 00:00:00" and raw_ba.loc[4, "Время"].startswith("1900-01-01")
      and raw_ba.loc[3, "Время"] == "0.4375" and raw_ba.loc[5, "Время"] == "1 day, 10:30:00"
      and raw_ba.loc[6, "Время"] == "1.4375" and raw_ba.loc[8, "Дата"].startswith("46117."),
      str(raw_ba.loc[:8, ["Дата", "Время"]].values.tolist()))
before_ba = parse_datetime_before(pd.Series(glue_before(raw_ba["Дата"], raw_ba["Время"])))
check("прежняя склейка и разбор читали 1 апреля как 4 января",
      before_ba.iloc[0] == pd.Timestamp("2026-01-04 09:05"), str(before_ba.iloc[0]))
canonical, result, stamps, manifest = upload(ba_xlsx, "апрель BA.xlsx")
check("раскладка Brand Analytics распознана", canonical.loc[0, "source_system"] == "brand_analytics",
      str(canonical.loc[0, "source_system"]))
check("раскладка Brand Analytics: дата и время склеены без полуночи",
      canonical.loc[0, "Дата"] == "2026-04-01 09:05:00" and canonical.loc[3, "Дата"] == "2026-04-02 10:30:00"
      and canonical.loc[4, "Дата"] == "2026-04-03 10:30:00",
      str(canonical.loc[:4, "Дата"].tolist()))
check("раскладка Brand Analytics: время больше суток и дробный номер дня — верное время",
      canonical.loc[5, "Дата"] == "2026-04-03 10:30:00" and canonical.loc[6, "Дата"] == "2026-04-04 10:30:00"
      and canonical.loc[8, "Дата"] == "2026-04-05 09:25:00",
      str(canonical.loc[5:8, "Дата"].tolist()))
check_april("раскладка Brand Analytics", result, stamps)

# (б2) Та же раскладка, у одной строки пустая ячейка даты при заполненном
# времени. Прежде такая строка становилась сегодняшним числом, и период
# тянулся до дня загрузки.
gap_rows = [list(row) for row in ba_rows]
gap_rows[3][0], gap_rows[3][1] = None, datetime(2026, 4, 2, 9, 15).time()
canonical, result, stamps, manifest = upload(xlsx_bytes(BA_HEADER, gap_rows, BA_FORMATS, BA_CELL_FORMATS),
                                             "апрель BA пустая дата.xlsx")
check("раскладка Brand Analytics, пустая дата: не сегодняшнее число, а без даты",
      result["period_name"] == "01.04.2026–30.04.2026" and result.get(UNDATED_MESSAGES) == 1
      and manifest.get(UNDATED_MESSAGES) == 1 and stamps.isna().sum() == 1
      and (stamps.dropna().dt.month == 4).all(),
      f"{result['period_name']}, без даты {result.get(UNDATED_MESSAGES)} / {manifest.get(UNDATED_MESSAGES)}, "
      f"строка: {canonical.loc[3, 'Дата']!r}")

# (в) CSV с датами ISO, первая строка — 1 апреля.
csv_frame = pd.DataFrame({
    "Дата": [m.strftime("%Y-%m-%d %H:%M:%S") for m in MOMENTS],
    "Текст": [text_of(i) for i in range(len(MOMENTS))],
    "Ссылка": [f"https://t.me/csv/{i}" for i in range(len(MOMENTS))],
    "Автор": [f"Автор {i % 5}" for i in range(len(MOMENTS))],
    "Площадка": ["telegram"] * len(MOMENTS),
})
canonical, result, stamps, manifest = upload(csv_frame.to_csv(index=False).encode("utf-8"), "апрель.csv")
check_april("CSV с ISO-датами", result, stamps)

# (г) Пояса Z и +03:00 в одной колонке: загрузка не обрывается.
zoned_frame = csv_frame.copy()
zoned_frame["Дата"] = [
    (m - pd.Timedelta(hours=3)).strftime("%Y-%m-%dT%H:%M:%SZ") if i % 2 else m.strftime("%Y-%m-%dT%H:%M:%S+03:00")
    for i, m in enumerate(MOMENTS)
]
try:
    canonical, result, stamps, manifest = upload(zoned_frame.to_csv(index=False).encode("utf-8"), "пояса.csv")
    zone_upload_error = None
except Exception as exc:  # noqa: BLE001 — исключение и есть результат проверки
    zone_upload_error = exc
check("CSV со смесью поясов загружается", zone_upload_error is None, repr(zone_upload_error))
if zone_upload_error is None:
    check_april("CSV со смесью поясов", result, stamps)
    check("[CSV со смесью поясов] время — московское, как в выгрузке",
          sorted(stamps.dt.strftime("%Y-%m-%d %H:%M")) == sorted(m.strftime("%Y-%m-%d %H:%M") for m in MOMENTS))

# (д) Пустые и нераспознанные даты: файл загружается, число — в итоге и манифесте.
undated_frame = csv_frame.copy()
undated_frame.loc[[5, 17, 33], "Дата"] = ["не дата", "вчера", "31.02.2026"]
undated_frame.loc[40, "Дата"] = ""
canonical, result, stamps, manifest = upload(undated_frame.to_csv(index=False).encode("utf-8"), "с пропусками.csv")
check("без даты — в итоге и манифесте (4 сообщения)",
      result.get(UNDATED_MESSAGES) == 4 and manifest.get(UNDATED_MESSAGES) == 4,
      f"{result.get(UNDATED_MESSAGES)}, {manifest.get(UNDATED_MESSAGES)}")
check("сообщения без даты сохранены — в ленте их не меньше",
      len(stamps) == 60 and stamps.isna().sum() == 4, f"{len(stamps)}, {stamps.isna().sum()}")
check("период назван по распознанным датам", result["period_name"] == "01.04.2026–30.04.2026", result["period_name"])

# (е) Колонки даты в файле нет совсем: сообщать о «нераспознанных» датах не о
# чем. Каноническая таблица всё равно содержит пустую «Дата» — счётчик
# смотрит на содержимое, а не на имя колонки.
no_date_frame = csv_frame.drop(columns=["Дата"])
canonical, result, stamps, manifest = upload(no_date_frame.to_csv(index=False).encode("utf-8"), "без даты.csv")
check("файл без колонки даты: в канонической таблице «Дата» пустая",
      "Дата" in canonical.columns and not canonical["Дата"].str.strip().ne("").any())
check("файл без колонки даты: сообщений без даты — 0 в итоге и манифесте, строки нет",
      result.get(UNDATED_MESSAGES) == 0 and manifest.get(UNDATED_MESSAGES) == 0 and undated_line(result) == "",
      f"{result.get(UNDATED_MESSAGES)}, {manifest.get(UNDATED_MESSAGES)}")

print("4. Отчёт о загрузке")
check("строка: одно сообщение",
      undated_line({UNDATED_MESSAGES: 1})
      == "Дата не распознана у 1 сообщения — оно не попадёт в динамику и сравнение периодов",
      undated_line({UNDATED_MESSAGES: 1}))
check("строка: несколько сообщений",
      undated_line({UNDATED_MESSAGES: 4})
      == "Дата не распознана у 4 сообщений — они не попадут в динамику и сравнение периодов",
      undated_line({UNDATED_MESSAGES: 4}))
check("строка: 21 и 1 500 — падеж и разряды",
      "у 21 сообщения — они" in undated_line({UNDATED_MESSAGES: 21})
      and "у 1 500 сообщений" in undated_line({UNDATED_MESSAGES: 1500}),
      f"{undated_line({UNDATED_MESSAGES: 21})} / {undated_line({UNDATED_MESSAGES: 1500})}")
check("ноль и старый манифест без счётчика — строки нет",
      undated_line({UNDATED_MESSAGES: 0}) == "" and undated_line({"rows_messages": 10}) == "" and undated_line({}) == "")
check("предупреждение — только больше 5 %",
      undated_needs_warning({UNDATED_MESSAGES: 6, "messages": 100})
      and not undated_needs_warning({UNDATED_MESSAGES: 5, "messages": 100})
      and undated_needs_warning({UNDATED_MESSAGES: 6, "rows_messages": 100})
      and not undated_needs_warning({UNDATED_MESSAGES: 0, "messages": 100}))

import streamlit as st  # noqa: E402
from streamlit.testing.v1 import AppTest  # noqa: E402

import ingest_admin_ui  # noqa: E402
import upload_history_ui as ui  # noqa: E402


def submit(file_bytes, filename="выгрузка.csv"):
    CLIENT.db["platform_periods"] = []
    CLIENT.db["platform_table_rows"] = []
    st.cache_data.clear()
    at = AppTest.from_file(str(REPO / "tests" / "pages_app.py"), default_timeout=180)
    at.session_state["page_name"] = "upload"
    at.session_state["page_role"] = "editor"
    at.session_state["page_project_id"] = PROJECT
    at.session_state["page_work_dir"] = WORK_DIR
    at.run()
    uploaders = [u for u in at.file_uploader if "CSV или Excel" in str(u.label)]
    buttons = [b for b in at.button if "Обработать и сохранить" in str(b.label)]
    if not uploaders or not buttons:
        return at, "нет поля загрузки или кнопки"
    uploaders[0].set_value((filename, file_bytes, "text/csv"))
    buttons[0].click().run()
    return at, ""


def saved(at):
    return any("сохранен" in str(s.value) for s in at.success)


def dated_lines(elements):
    return [str(e.value) for e in elements if "Дата не распознана" in str(e.value)]


many_undated = csv_frame.copy()
many_undated.loc[[3, 9, 15, 21, 27], "Дата"] = "неизвестно"
at, problem = submit(many_undated.to_csv(index=False).encode("utf-8"))
warnings_shown = dated_lines(at.warning)
check("страница: 5 из 60 без даты — файл загружен", not problem and not at.exception and saved(at),
      problem or str([e.value for e in at.error] or at.exception))
check("страница: 5 из 60 (8 %) — предупреждение с числом",
      len(warnings_shown) == 1 and "у 5 сообщений" in warnings_shown[0], str(warnings_shown + dated_lines(at.info)))
few_undated = pd.concat([csv_frame] * 2, ignore_index=True)
few_undated["Ссылка"] = [f"https://t.me/two/{i}" for i in range(len(few_undated))]
few_undated.loc[7, "Дата"] = ""
at, problem = submit(few_undated.to_csv(index=False).encode("utf-8"))
check("страница: 1 из 120 — справка, не предупреждение",
      not problem and saved(at) and not dated_lines(at.warning)
      and dated_lines(at.info) == ["Дата не распознана у 1 сообщения — оно не попадёт в динамику и сравнение периодов."],
      str(dated_lines(at.warning) + dated_lines(at.info)))
at, problem = submit(csv_frame.to_csv(index=False).encode("utf-8"))
check("страница: все даты распознаны — строки нет",
      not problem and saved(at) and not dated_lines(at.warning) and not dated_lines(at.info),
      str(dated_lines(at.warning) + dated_lines(at.info)))
at, problem = submit(no_date_frame.to_csv(index=False).encode("utf-8"))
check("страница: в файле нет колонки даты — ни предупреждения, ни справки о датах",
      not problem and saved(at) and not dated_lines(at.warning) and not dated_lines(at.info),
      problem or str(dated_lines(at.warning) + dated_lines(at.info)))

LONG_ERROR = "Traceback (most recent call last):\n" + "  File \"x.py\", line 1, in <module>\n" * 20
TASKS = pd.DataFrame([
    {"status": "done", "original_filename": "a.xlsx", "error_message": "",
     "result": {"messages": 200, UNDATED_MESSAGES: 30, "duplicate_rows_removed": 2, "shared_id_messages": 0}},
    {"status": "done", "original_filename": "b.xlsx", "error_message": "", "result": {"messages": 10, UNDATED_MESSAGES: 0}},
    {"status": "done", "original_filename": "c.xlsx", "error_message": "",
     "result": {"messages": 2000, UNDATED_MESSAGES: 1200, "duplicate_rows_removed": 37, "shared_id_messages": 12}},
    {"status": "error", "original_filename": "d.xlsx", "error_message": LONG_ERROR, "result": None},
])
comments = ingest_admin_ui._tasks_view(TASKS, is_admin=True)["Комментарий"].tolist()
check("очередь автозагрузки: строка о датах после строк о повторах",
      comments[:2] == ["Повторы одного и того же сообщения убраны: 2; "
                       "Дата не распознана у 30 сообщений — они не попадут в динамику и сравнение периодов", ""],
      str(comments[:2]))
check("очередь автозагрузки: при всех трёх счётчиках строка о датах видна целиком",
      comments[2].endswith("Дата не распознана у 1 200 сообщений — они не попадут в динамику и сравнение периодов"),
      comments[2])
check("очередь автозагрузки: длинный текст ошибки по-прежнему обрезан до 160 знаков",
      comments[3] == LONG_ERROR[:160], f"{len(comments[3])} знаков")
summary = ui._rebuild_summary({"period_name": "Апрель", "messages": 60, "events_before": 1, "events": 1, UNDATED_MESSAGES: 3})
check("сводка пересборки называет сообщения без даты", "Дата не распознана у 3 сообщений" in summary, summary)

print("5. Проверка уже загруженных периодов")
BROKEN = {"period_id": "p_bad", "period_name": "04.01.2026–04.12.2026", "date_from": "2026-01-04", "date_to": "2026-12-04",
          "manifest": {"rows_messages": 2000}}
WEEK_BROKEN = {"period_id": "p_week", "date_from": "2026-01-04", "date_to": "2026-07-04", "manifest": {}}
GOOD = {"period_id": "p_ok", "period_name": "01.04.2026–30.04.2026", "date_from": "2026-04-01", "date_to": "2026-04-30",
        "manifest": {"rows_messages": 2000, UNDATED_MESSAGES: 3}}
NEW_UNDATED = {"period_id": "p_new", "date_from": "2026-04-01", "date_to": "2026-04-30",
               "manifest": {"rows_messages": 100, UNDATED_MESSAGES: 40}}
problems = period_date_problems(BROKEN)
check("«04.01.2026–04.12.2026»: перепутаны день и месяц, слишком длинный",
      any("перепутаны" in p for p in problems) and any("длиннее" in p for p in problems), str(problems))
check("неделя 1–7 апреля, ставшая «04.01–04.07»: перепутаны",
      any("перепутаны" in p for p in period_date_problems(WEEK_BROKEN)), str(period_date_problems(WEEK_BROKEN)))
check("обычный апрель и 3 сообщения без даты из 2 000 — признаков нет", period_date_problems(GOOD) == [],
      str(period_date_problems(GOOD)))
check("новый период: 40 % без даты — признак", any("без даты 40 из 100" in p for p in period_date_problems(NEW_UNDATED)),
      str(period_date_problems(NEW_UNDATED)))
swapped_messages = pd.DataFrame({"datetime": [f"2026-{m:02d}-04T10:00:00" for m in range(1, 13)] + [None] * 8})
check("сообщения: все — 4-е число разных месяцев, много без даты",
      len(messages_date_problems(swapped_messages)) == 2, str(messages_date_problems(swapped_messages)))
check("сообщения обычного апреля — признаков нет",
      messages_date_problems(pd.DataFrame({"datetime": [m.isoformat() for m in MOMENTS]})) == [])
# Обычные периоды, у которых одно число на границах или у сообщений.
MONTH_FROM_FIFTH = {"period_id": "p_month", "date_from": "2026-03-05", "date_to": "2026-04-05", "manifest": {}}
check("месячный период 05.03–05.04 — не «перепутаны»", period_date_problems(MONTH_FROM_FIFTH) == [],
      str(period_date_problems(MONTH_FROM_FIFTH)))
check("три дня 01–03.04, ставшие «04.01–04.03», — перепутаны",
      any("перепутаны" in p for p in period_date_problems({"date_from": "2026-01-04", "date_to": "2026-03-04"})))
one_day = pd.DataFrame({"datetime": [f"2026-03-31T23:4{i}:00" for i in range(5)]
                                    + [f"2026-04-01T{h:02d}:{m:02d}:00" for h in range(24) for m in (0, 15, 30, 45)]})
check("выгрузка за сутки на стыке месяцев (почти все — 1-е число) — признаков нет",
      messages_date_problems(one_day) == [], str(messages_date_problems(one_day)))

import check_period_dates  # noqa: E402

CLIENT.db["platform_projects"] = [{"project_id": PROJECT, "project_name": "Даты", "status": "active", "settings": {}}]
CLIENT.db["platform_periods"] = [
    {"project_id": PROJECT, "status": "active", "uploaded_at": "2026-05-01T10:00:00+00:00", **BROKEN},
    {"project_id": PROJECT, "status": "active", "uploaded_at": "2026-05-02T10:00:00+00:00", **GOOD},
]
CLIENT.db["platform_table_rows"] = [
    {"project_id": PROJECT, "period_id": "p_bad", "table_name": "messages", "row_id": f"m{i}",
     "payload": {"message_id": f"m{i}", "datetime": value}}
    for i, value in enumerate(swapped_messages["datetime"].tolist())
]
WRITES = []


class WatchedTable(Table):
    """Таблица поддельной базы, которая запоминает любую запись в неё."""

    def update(self, payload):
        WRITES.append(("update", self.name))
        return super().update(payload)

    def insert(self, payload):
        WRITES.append(("insert", self.name))
        return super().insert(payload)

    def upsert(self, payload, on_conflict=None, ignore_duplicates=False):
        WRITES.append(("upsert", self.name))
        return super().upsert(payload, on_conflict=on_conflict, ignore_duplicates=ignore_duplicates)

    def delete(self):
        WRITES.append(("delete", self.name))
        return super().delete()


CLIENT.table = lambda name: WatchedTable(CLIENT.db, name)
periods_before = copy.deepcopy(CLIENT.db["platform_periods"])
files_before = dict(CLIENT.files)
found = check_period_dates.find_suspicious_periods(with_messages=True)
check("скрипт находит только период с перепутанными датами",
      [item["period_id"] for item in found] == ["p_bad"] and len(found[0]["problems"]) == 4,
      str(found))
buffer = io.StringIO()
real_stdout, sys.stdout = sys.stdout, buffer
try:
    code = check_period_dates.main([])
finally:
    sys.stdout = real_stdout
check("скрипт печатает период и признаки", code == 0 and "04.01.2026–04.12.2026" in buffer.getvalue(), buffer.getvalue()[:300])
check("скрипт ничего не пишет в базу",
      not WRITES and CLIENT.db["platform_periods"] == periods_before and CLIENT.files == files_before, str(WRITES))

print()
if failures:
    print(f"ПРОВАЛЕНО: {len(failures)}")
    for f in failures:
        print(f"  - {f}")
    raise SystemExit(1)
print("Даты выгрузки читаются верно: ISO и ячейки Excel — без перестановки дня и месяца, «дд.мм.гггг» — как раньше.")
